"""El motor de documentos de la escuela: el 911, constancias, listas de grupo y los formatos de la supervisión
(validacion/formatos_oficiales.py), en Excel, Word y PDF.

Vive en la PC 4: el cerebro los arma con la base de todos los alumnos (POST /documento, validacion/cerebro_red.py). Si la
PC 4 no contesta, el servidor (SIDP.exe) arma los mismos con sus listas de grupo, sin lo que sacó el escaneo. Solo usa
los datos del estándar (validacion/estandar.py) y los de la escuela (validacion/escuela.py), que llegan con cada pedido.

Los documentos llevan nombres y CURP: se descargan en la PC de quien los pide y no se suben a internet.
Excel con openpyxl, PDF con reportlab y Word escrito aquí mismo (un .docx es un zip con XML).
"""
import datetime
import io
import re
import unicodedata
import zipfile
from xml.sax.saxutils import escape

from validacion import escuela as escuela_mod
from validacion import estandar, formatos_oficiales, indicadores

try:  # Arriba para que PyInstaller lo meta en SIDP.exe.
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except ImportError:  # Sin reportlab salen Excel y Word; el PDF avisa qué falta.
    colors = None

MIME = {'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'pdf': 'application/pdf'}
MESES = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')
ORDINAL = {1: 'primer', 2: 'segundo', 3: 'tercer', 4: 'cuarto', 5: 'quinto', 6: 'sexto'}
EN_BLANCO = '________________'
CUIDADO = 'Tiene datos personales de los alumnos: guárdalo en la escuela, no lo mandes por correo ni lo subas a internet.'

CATALOGO = (
    {'clave': '911', 'nombre': 'Estadística 911 de inicio de curso', 'boton': 'Bajar el 911 en Excel', 'alcance': 'escuela', 'formatos': ['xlsx'],
     'descripcion': 'Los cuadros del cuestionario 911.3 ya contados: alumnos por grado, sexo y edad, repetidores, grupos, discapacidad, '
                    'lengua indígena, lugar de nacimiento y becas. Trae una hoja «Revisión» que dice si todo cuadra y otra con los alumnos a los que les falta un dato.'},
    {'clave': 'constancia', 'nombre': 'Constancia de inscripción', 'boton': 'Hacer constancias', 'alcance': 'alumno', 'formatos': ['docx', 'pdf'],
     'descripcion': 'Para la beca Rita Cetina, el IMSS o cualquier trámite de los padres. De un alumno o de todo el grupo '
                    '(una hoja por alumno), lista para firmar y sellar.'},
    {'clave': 'lista', 'nombre': 'Lista de grupo', 'alcance': 'grupo_o_escuela', 'formatos': ['pdf', 'xlsx', 'docx'],
     'descripcion': 'Número, nombre, CURP, sexo y edad de cada alumno, con cuántos son hombres y mujeres. '
                    'De un grupo o de toda la escuela (una hoja por grupo).'},
    {'clave': 'emergencia', 'nombre': 'Lo que la maestra debe saber', 'alcance': 'grupo_o_escuela', 'formatos': ['pdf', 'xlsx'],
     'descripcion': 'Por grupo: alergias, padecimientos, medicamentos, discapacidad, tutor, teléfono y contacto de emergencia, '
                    'para tenerlo a la mano en el salón.'},
    {'clave': 'metadato', 'nombre': 'Base de datos Metadato (71 columnas)', 'boton': 'Bajar el Metadato', 'alcance': 'escuela', 'formatos': ['xlsx'],
     'descripcion': 'El formato de la supervisión con los datos del alumno (columnas 1 a 40) y del tutor (41 a 71), ya llenos con lo que '
                    'tiene SIDP. La edad se calcula sola. Lo que quede en blanco se completa antes de entregarlo.'},
    {'clave': 'base_911', 'nombre': 'Estadística 911: base que se calcula sola', 'boton': 'Bajar la base del 911', 'alcance': 'escuela',
     'formatos': ['xlsx'],
     'descripcion': 'La hoja BASE con la CURP de cada alumno: el sexo, la fecha de nacimiento, la edad al 1 de septiembre y los cuadros '
                    'por grado, sexo y edad salen solos. Si se agrega a un alumno a mano, los totales cambian.'},
    {'clave': 'condicion', 'nombre': 'Censo de condición y discapacidad', 'alcance': 'escuela', 'formatos': ['xlsx'],
     'descripcion': 'Una hoja por grado con las columnas de condición y de discapacidad, marcadas con X según el certificado médico y lo '
                    'que capturó la escuela. Los totales se cuentan solos.'},
    {'clave': 'calificaciones', 'nombre': 'Concentrado de calificaciones y dispersión', 'alcance': 'grupo_o_escuela', 'formatos': ['xlsx'],
     'descripcion': 'Una hoja por grupo con los alumnos ya puestos. El maestro escribe las calificaciones de los cuatro campos formativos; '
                    'el promedio, la dispersión por nivel y el concentrado escolar se calculan solos.'},
    {'clave': 'boleta', 'nombre': 'Base de datos de la boleta', 'alcance': 'grupo_o_escuela', 'formatos': ['xlsx'],
     'descripcion': 'Los cuatro campos formativos y sus observaciones de los tres periodos, con grado, grupo, CURP y nombre de cada alumno '
                    'ya escritos.'},
    {'clave': 'asistencias', 'nombre': 'Concentrado de asistencias y faltas', 'alcance': 'grupo_o_escuela', 'formatos': ['xlsx'],
     'descripcion': 'Una hoja por grupo, de septiembre a julio: asistencias, faltas justificadas y sin justificar. Los días del mes y el '
                    'total del ciclo se suman solos.'},
    {'clave': 'vive_saludable', 'nombre': 'Lista para Vive Saludable', 'boton': 'Bajar la lista de Vive Saludable', 'alcance': 'grupo_o_escuela',
     'formatos': ['xlsx'],
     'descripcion': 'Los alumnos por grupo con CURP, sexo, fecha de nacimiento y edad, para cuando el centro de salud avise de la jornada. '
                    'Las medidas las toma y las guarda el personal de salud.'},
)
POR_CLAVE = {d['clave']: d for d in CATALOGO}


class ErrorDocumento(Exception):
    """No se pudo armar el documento; el mensaje es para la persona."""


def catalogo():
    return [dict(d) for d in CATALOGO]


# ---------------------------------------------------------------- Los alumnos

def de_lista(filas, ciclo, hoy=None):
    """Registros estándar con lo que tiene el servidor (grupos.exportar_alumnos), para cuando la PC 4 no contesta."""
    salida = []
    for fila in filas:
        grado, letra, primer, segundo, nombres, curp, _, sexo, nacimiento = fila[:9]
        alumno = {'curp': curp, 'nombres': nombres, 'primer_apellido': primer, 'segundo_apellido': segundo, 'grado': grado, 'letra': letra,
                  'sexo': sexo if sexo in ('H', 'M') else _sexo_curp(curp), 'fecha_nacimiento': indicadores.fecha(nacimiento) or _fecha_curp(curp)}
        registro = estandar.registro(alumno, {}, ciclo, hoy)
        registro.update(tutor=fila[9] or registro['tutor'], telefono=fila[11] or registro['telefono'])
        salida.append(registro)
    return salida


def _sexo_curp(curp):
    curp = (curp or '').upper()
    return curp[10] if len(curp) == 18 and curp[10] in 'HM' else None


def _fecha_curp(curp):
    curp = (curp or '').upper()
    if len(curp) != 18:
        return None
    try:
        return datetime.date((2000 if curp[16].isalpha() else 1900) + int(curp[4:6]), int(curp[6:8]), int(curp[8:10]))
    except ValueError:
        return None


def _sexo(r):
    return r.get('sexo') or r.get('sexo_curp') or None


def _fecha_texto(valor):
    fecha = indicadores.fecha(valor)
    return f'{fecha.day:02d}/{fecha.month:02d}/{fecha.year}' if fecha else ''


def _fecha_larga(hoy):
    return f'{hoy.day} de {MESES[hoy.month - 1]} de {hoy.year}'


def _grupos(registros):
    """{(grado, letra): [registros]} en orden."""
    salida = {}
    for r in sorted(registros, key=lambda x: (x.get('grado') or 9, x.get('letra') or '', x.get('nombre') or '')):
        salida.setdefault((r.get('grado'), r.get('letra') or ''), []).append(r)
    return salida


def _nombre_grupo(grado, letra):
    return f'{grado}° {letra}'.strip() if grado else 'Sin grupo'


def _archivo(*partes):
    texto = '_'.join(str(p) for p in partes if p not in (None, ''))
    texto = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode()
    return re.sub(r'[^A-Za-z0-9_.-]+', '_', texto).strip('_')


def generar(clave, formato, registros, escuela=None, ciclo='', hoy=None, grado=None, letra=None, curp=None):
    """(nombre de archivo, contenido, tipo MIME) del documento `clave` en `formato`, de toda la escuela, un grupo o un alumno."""
    hoy = hoy or datetime.date.today()
    documento = POR_CLAVE.get(clave)
    if not documento:
        raise ErrorDocumento('Ese documento no existe.')
    if formato not in documento['formatos']:
        raise ErrorDocumento(f'«{documento["nombre"]}» no sale en ese formato.')
    escuela = {**dict.fromkeys(escuela_mod.CAMPOS), **(escuela or {})}
    escuela['turno_texto'] = escuela.get('turno_texto') or escuela_mod.TURNOS.get(escuela.get('turno'))
    elegidos = list(registros)
    if grado and documento['alcance'] != 'escuela':
        elegidos = [r for r in elegidos if r.get('grado') == int(grado) and (not letra or (r.get('letra') or '') == letra)]
    if curp:
        elegidos = [r for r in elegidos if r.get('curp') == curp]
    if not elegidos:
        raise ErrorDocumento('No hay alumnos para ese documento. Revisa el grupo o carga primero las listas de grupo.')
    sufijo = _archivo(_nombre_grupo(grado, letra).replace('° ', '')) if grado and documento['alcance'] != 'escuela' else 'escuela'
    if clave == '911':
        return _archivo('911_inicio', ciclo) + '.xlsx', libro_911(registros, escuela, ciclo, hoy), MIME['xlsx']
    if clave == 'constancia':
        r = elegidos[0]
        nombre = _archivo('constancia', r['primer_apellido'], r['nombres']) if curp else _archivo('constancias', sufijo, hoy.isoformat())
        return f'{nombre}.{formato}', _escribir(formato, constancias(elegidos, escuela, ciclo, hoy)), MIME[formato]
    nombre = _archivo(clave, sufijo, hoy.isoformat()) + f'.{formato}'
    if clave == 'vive_saludable':
        return nombre, _excel(vive_saludable(elegidos, escuela, ciclo)), MIME['xlsx']
    if clave in formatos_oficiales.GENERADORES:
        return nombre, formatos_oficiales.GENERADORES[clave](elegidos, escuela, ciclo, hoy), MIME['xlsx']
    hojas = listas(elegidos, escuela, ciclo, salud=clave == 'emergencia')
    if formato == 'xlsx':
        return nombre, _excel(hojas), MIME['xlsx']
    return nombre, _escribir(formato, _hojas_a_bloques(hojas, escuela), horizontal=clave == 'emergencia'), MIME[formato]


# ---------------------------------------------------------------- Cada documento

def _encabezado(escuela, ciclo=''):
    lineas = [escuela.get('nombre') or 'Escuela primaria ' + EN_BLANCO]
    datos = [f'CCT: {escuela.get("cct") or EN_BLANCO}', f'Turno: {escuela.get("turno_texto") or EN_BLANCO}']
    if escuela.get('zona_escolar'):
        datos.append(f'Zona escolar: {escuela["zona_escolar"]}')
    if ciclo:
        datos.append(f'Ciclo escolar {ciclo}')
    lineas.append(' · '.join(datos))
    lugar = ', '.join(filter(None, (escuela.get('domicilio'), escuela.get('municipio'), 'Estado de México')))
    lineas.append(lugar)
    return ('encabezado', lineas)


def constancias(registros, escuela, ciclo, hoy):
    """Una hoja por alumno, para firmar y sellar."""
    bloques = []
    for numero, r in enumerate(registros):
        if numero:
            bloques.append(('salto',))
        sexo = _sexo(r)
        alumno, inscrito = {'H': ('el alumno', 'inscrito'), 'M': ('la alumna', 'inscrita')}.get(sexo, ('el alumno o la alumna', 'inscrito(a)'))
        grado = r.get('grado')
        nombre = escuela.get('nombre') or EN_BLANCO
        de_la = f'de la {nombre}' if nombre.lower().startswith('escuela') else f'de la escuela {nombre}'
        bloques += [
            _encabezado(escuela),
            ('espacio', 2), ('parrafo', [('ASUNTO: CONSTANCIA DE INSCRIPCIÓN', True)], 'derecha'), ('espacio', 1),
            ('parrafo', [('A QUIEN CORRESPONDA:', True)]), ('espacio', 1),
            ('parrafo', [(f'El que suscribe, director {de_la}, con clave de centro de trabajo '
                          f'{escuela.get("cct") or EN_BLANCO}, hace constar que {alumno} ', False),
                         (' '.join(filter(None, (r.get('nombres'), r.get('primer_apellido'), r.get('segundo_apellido')))), True),
                         (', con CURP ', False), (r.get('curp') or EN_BLANCO, True),
                         (f', está {inscrito} en este plantel en el {ORDINAL.get(grado, EN_BLANCO)} grado, grupo {r.get("letra") or EN_BLANCO}, '
                          f'del ciclo escolar {ciclo or EN_BLANCO}, en el turno {(escuela.get("turno_texto") or EN_BLANCO).lower()}.', False)],
             'justificado'),
            ('espacio', 1),
            ('parrafo', [('Se extiende la presente a petición del interesado, para los fines que le convengan, en '
                          f'{escuela.get("municipio") or EN_BLANCO}, Estado de México, a {_fecha_larga(hoy)}.', False)], 'justificado'),
            ('espacio', 3), ('parrafo', [('ATENTAMENTE', True)], 'centro'), ('espacio', 3),
            ('firma', [escuela.get('director') or EN_BLANCO, 'Director de la escuela']),
            ('espacio', 2), ('parrafo', [('Sello de la escuela', False)], 'chico'),
        ]
    return bloques


def listas(registros, escuela, ciclo, salud=False):
    """Una hoja por grupo: la lista (con CURP, sexo y edad) o lo que la maestra debe saber (salud y contacto)."""
    hojas = []
    for (grado, letra), alumnos in _grupos(registros).items():
        grupo = _nombre_grupo(grado, letra)
        if salud:
            encabezado = ['No.', 'Nombre', 'Alergias', 'Padecimientos', 'Medicamentos', 'Discapacidad o condición', 'Tutor', 'Teléfono',
                          'Contacto de emergencia']
            filas = [[i, r['nombre'], ', '.join(r.get('alergias') or []), ', '.join(r.get('padecimientos') or []),
                      ', '.join(r.get('medicamentos') or []), ', '.join(r.get('discapacidades') or []), r.get('tutor') or '',
                      r.get('telefono') or '', r.get('emergencia') or ''] for i, r in enumerate(alumnos, 1)]
            titulo = f'Lo que la maestra debe saber · Grupo {grupo}'
            nota = 'Datos de salud de los alumnos: tenlo a la mano en el salón y no lo compartas.'
            anchos = (5, 26, 16, 16, 15, 15, 20, 11, 18)
        else:
            encabezado = ['No.', 'Nombre (apellidos y nombres)', 'CURP', 'Sexo', 'Fecha de nacimiento', 'Edad']
            filas = [[i, r['nombre'], r.get('curp') or '', {'H': 'H', 'M': 'M'}.get(_sexo(r), ''), _fecha_texto(r.get('fecha_nacimiento')),
                      r.get('edad_911') if r.get('edad_911') is not None else ''] for i, r in enumerate(alumnos, 1)]
            titulo = f'Lista del grupo {grupo}'
            hombres, mujeres = sum(_sexo(r) == 'H' for r in alumnos), sum(_sexo(r) == 'M' for r in alumnos)
            nota = f'{len(alumnos)} alumnos: {hombres} hombres y {mujeres} mujeres. Edad cumplida al 1 de septiembre (como la pide el 911).'
            anchos = (5, 42, 22, 6, 12, 6)
        hojas.append({'nombre': grupo, 'titulo': titulo, 'nota': nota, 'tablas': [(None, encabezado, filas)], 'anchos': anchos,
                      'escuela': escuela, 'ciclo': ciclo})
    return hojas


def vive_saludable(registros, escuela, ciclo):
    encabezado = ['Grupo', 'No.', 'Primer apellido', 'Segundo apellido', 'Nombre(s)', 'CURP', 'Sexo', 'Fecha de nacimiento', 'Edad']
    filas = []
    for (grado, letra), alumnos in _grupos(registros).items():
        filas += [[_nombre_grupo(grado, letra), i, r.get('primer_apellido'), r.get('segundo_apellido'), r.get('nombres'), r.get('curp'),
                   {'H': 'H', 'M': 'M'}.get(_sexo(r), ''), _fecha_texto(r.get('fecha_nacimiento')), r.get('edad_911')]
                  for i, r in enumerate(alumnos, 1)]
    return [{'nombre': 'Alumnos', 'titulo': f'Alumnos para la jornada de Vive Saludable · {escuela.get("nombre") or ""}'.strip(' ·'),
             'nota': f'CCT {escuela.get("cct") or EN_BLANCO} · Ciclo {ciclo}. {CUIDADO}', 'tablas': [(None, encabezado, filas)],
             'anchos': (8, 5, 16, 16, 20, 21, 6, 12, 6)}]


# ---------------------------------------------------------------- El 911

def _hm(alumnos):
    hombres = sum(_sexo(r) == 'H' for r in alumnos)
    mujeres = sum(_sexo(r) == 'M' for r in alumnos)
    return [hombres, mujeres, len(alumnos)]


def _por_grado(registros):
    return {g: [r for r in registros if r.get('grado') == g] for g in range(1, 7)}


def cuadros_911(registros):
    """Los cuadros del 911.3 como listas de filas (también los usa la prueba)."""
    grados = _por_grado(registros)
    edades = list(indicadores.EDADES_911) + ['Sin fecha']
    edad = []
    for g, alumnos in grados.items():
        for sexo, nombre in (('H', 'Hombres'), ('M', 'Mujeres'), (None, 'Subtotal')):
            parte = [r for r in alumnos if sexo is None or _sexo(r) == sexo]
            cuenta = [sum((r.get('columna_911') or 'Sin fecha') == e for r in parte) for e in edades]
            edad.append([f'{g}°', nombre, *cuenta, len(parte)])
    for sexo, nombre in (('H', 'Hombres'), ('M', 'Mujeres'), (None, 'Total')):
        parte = [r for r in registros if sexo is None or _sexo(r) == sexo]
        edad.append(['Total', nombre, *[sum((r.get('columna_911') or 'Sin fecha') == e for r in parte) for e in edades], len(parte)])
    repetidores = []
    for g, alumnos in grados.items():
        repiten = [r for r in alumnos if r.get('repite') == 'sí']
        repetidores.append([f'{g}°', *_hm(alumnos), *_hm(repiten), sum(r.get('repite') == 'no' for r in alumnos),
                            sum(r.get('repite') not in ('sí', 'no') for r in alumnos)])
    repetidores.append(['Total', *[sum(f[i] for f in repetidores) for i in range(1, 9)]])
    grupos = [[f'{g}°', _nombre_grupo(g, letra), *_hm(alumnos)] for (g, letra), alumnos in _grupos(registros).items() if g]
    por_grado = [[f'{g}°', len({r.get('letra') for r in alumnos}), *_hm(alumnos)] for g, alumnos in grados.items()]
    tipos = [t for t in estandar.CONDICIONES_911 if any(t in (r.get('discapacidades') or []) for r in registros)]
    discapacidad = []
    for tipo in tipos:
        fila = [tipo]
        for alumnos in grados.values():
            fila += _hm([r for r in alumnos if tipo in (r.get('discapacidades') or [])])[:2]
        discapacidad.append(fila + [sum(fila[1:])])
    lengua = [[f'{g}°', *_hm([r for r in a if r.get('lengua_indigena') == 'sí']), *_hm([r for r in a if r.get('afro') == 'sí']),
               sum(r.get('lengua_indigena') not in ('sí', 'no') for r in a), sum(r.get('afro') not in ('sí', 'no') for r in a)]
              for g, a in grados.items()]
    lengua.append(['Total', *[sum(f[i] for f in lengua) for i in range(1, 9)]])
    nacidos = [[f'{g}°', *_hm([r for r in a if r.get('nacido') == 'Otra entidad']), *_hm([r for r in a if r.get('nacido') == 'Otro país']),
                sum(r.get('nacido') is None for r in a)] for g, a in grados.items()]
    nacidos.append(['Total', *[sum(f[i] for f in nacidos) for i in range(1, 8)]])
    entidades = sorted(_contar(r.get('entidad') for r in registros if r.get('nacido') == 'Otra entidad').items(), key=lambda x: -x[1])
    becas = [[texto, *_hm([r for r in registros if r.get('beca') == clave])] for clave, texto in estandar.BECA.items()]
    becas.append(['Sin el dato', *_hm([r for r in registros if r.get('beca') not in estandar.BECA])])
    procedencia = [[texto, *_hm([r for r in registros if r.get('procedencia') == clave])] for clave, texto in estandar.PROCEDENCIA.items()]
    procedencia.append(['Sin el dato', *_hm([r for r in registros if r.get('procedencia') not in estandar.PROCEDENCIA])])
    lenguas = sorted(_contar(r.get('lengua') for r in registros if r.get('lengua_indigena') == 'sí' and r.get('lengua')).items(),
                     key=lambda x: -x[1])
    return {'edad': edad, 'edades': edades, 'repetidores': repetidores, 'grupos': grupos, 'grupos_por_grado': por_grado,
            'discapacidad': discapacidad, 'lengua': lengua, 'lenguas': lenguas, 'nacidos': nacidos, 'entidades': entidades,
            'becas': becas, 'procedencia': procedencia}


def _contar(valores):
    salida = {}
    for v in valores:
        if v:
            salida[v] = salida.get(v, 0) + 1
    return salida


def revisar_911(registros, escuela, cuadros):
    """La hoja «Revisión»: lo que hay que revisar antes de capturar en la plataforma de la SEP."""
    total = len(registros)
    suma_edades = cuadros['edad'][-1][-1]
    suma_grupos = sum(f[-1] for f in cuadros['grupos'])
    sin = {clave: sum(1 for r in registros if r.get(clave) in (None, '')) for clave, *_ in estandar.PARA_911}
    vacios = [g for g in cuadros['grupos'] if not g[-1]]
    filas = [
        ['La tabla de edades suma lo mismo que la matrícula', 'Sí' if suma_edades == total else 'No', f'{suma_edades} de {total} alumnos'],
        ['Todos los alumnos están en un grupo', 'Sí' if suma_grupos == total else 'No', f'{suma_grupos} de {total} en los grupos'],
        ['Todos tienen sexo', 'Sí' if not sin['sexo_curp'] else 'No', f'{sin["sexo_curp"]} sin sexo' if sin['sexo_curp'] else 'Sale de la CURP'],
        ['Todos tienen fecha de nacimiento', 'Sí' if not sin['fecha_nacimiento'] else 'No',
         f'{sin["fecha_nacimiento"]} sin fecha: cuentan en «Sin fecha»' if sin['fecha_nacimiento'] else 'Edad al 1 de septiembre'],
        ['Todos tienen «¿Repite el grado?»', 'Sí' if not sin['repite'] else 'No',
         f'{sin["repite"]} sin el dato: complétalo en «Reportes para la SEP»' if sin['repite'] else 'Completo'],
        ['Todos tienen el dato de beca', 'Sí' if not sin['beca'] else 'No', f'{sin["beca"]} sin el dato' if sin['beca'] else 'Completo'],
        ['Todos tienen el dato de lengua indígena', 'Sí' if not sin['lengua_indigena'] else 'No',
         f'{sin["lengua_indigena"]} sin el dato' if sin['lengua_indigena'] else 'Completo'],
        ['Ningún grupo está vacío', 'Sí' if not vacios else 'No', ', '.join(g[1] for g in vacios) or 'Bien'],
        ['Datos de la escuela completos', 'Sí' if not escuela_mod.faltan(escuela) else 'No',
         'Faltan ' + ', '.join(escuela_mod.faltan(escuela)) if escuela_mod.faltan(escuela) else 'CCT, turno y director'],
    ]
    return filas


def libro_911(registros, escuela, ciclo, hoy):
    c = cuadros_911(registros)
    grados = [f'{g}°' for g in range(1, 7)]
    hm = [f'{g} {s}' for g in grados for s in ('H', 'M')]
    copia = 'Copia estos números en el cuestionario 911.3 de la plataforma de la SEP: el que vale es el que se captura allá.'
    faltan = []
    for r in sorted(registros, key=lambda x: (x.get('grado') or 9, x.get('letra') or '', x.get('nombre') or '')):
        que = [nombre for clave, nombre, *_ in estandar.PARA_911 if r.get(clave) in (None, '')]
        if que:
            faltan.append([r.get('grupo'), r.get('nombre'), r.get('curp'), ', '.join(que)])
    hojas = [
        {'nombre': 'Escuela', 'titulo': 'Estadística 911 de inicio de curso', 'nota': copia,
         'tablas': [(None, ['Dato', 'Valor'], [['Escuela', escuela.get('nombre') or ''], ['Clave del centro de trabajo (CCT)', escuela.get('cct') or ''],
                                                ['Turno', escuela.get('turno_texto') or ''], ['Zona escolar', escuela.get('zona_escolar') or ''],
                                                ['Municipio', escuela.get('municipio') or ''], ['Director', escuela.get('director') or ''],
                                                ['Ciclo escolar', ciclo], ['Alumnos inscritos', len(registros)],
                                                ['Contado el', _fecha_texto(hoy)]])],
         'anchos': (34, 40)},
        {'nombre': 'Alumnos por edad', 'titulo': 'Alumnos por grado, sexo y edad cumplida al 1 de septiembre', 'nota': copia,
         'tablas': [(None, ['Grado', 'Sexo', *c['edades'], 'Total'], c['edad'])], 'anchos': (7, 10) + (7,) * (len(c['edades']) + 1)},
        {'nombre': 'Repetidores', 'titulo': 'Inscritos y repetidores por grado y sexo', 'nota': '«No repiten» y «Sin el dato» cuentan a todos los del grado.',
         'tablas': [(None, ['Grado', 'Inscritos H', 'Inscritos M', 'Inscritos total', 'Repiten H', 'Repiten M', 'Repiten total', 'No repiten',
                            'Sin el dato'], c['repetidores'])], 'anchos': (7,) + (11,) * 8},
        {'nombre': 'Grupos', 'titulo': 'Grupos y alumnos por grupo', 'nota': copia,
         'tablas': [('Por grado', ['Grado', 'Grupos', 'Hombres', 'Mujeres', 'Total'], c['grupos_por_grado']),
                    ('Por grupo', ['Grado', 'Grupo', 'Hombres', 'Mujeres', 'Total'], c['grupos'])], 'anchos': (8, 10, 10, 10, 10)},
        {'nombre': 'Discapacidad', 'titulo': 'Alumnos con discapacidad o condición, por tipo, grado y sexo',
         'nota': 'Un alumno con dos condiciones cuenta en las dos. Sale del certificado médico o de lo que escribió la escuela.',
         'tablas': [(None, ['Tipo', *hm, 'Total'], c['discapacidad'] or [['Ningún alumno con discapacidad registrada', *[0] * 13]])],
         'anchos': (34,) + (6,) * 13},
        {'nombre': 'Lengua y afro', 'titulo': 'Lengua indígena y afrodescendientes, por grado y sexo', 'nota': 'Lo responde la familia; la escuela lo anota.',
         'tablas': [(None, ['Grado', 'Lengua indígena H', 'Lengua indígena M', 'Lengua indígena total', 'Afro H', 'Afro M', 'Afro total',
                            'Sin el dato de lengua', 'Sin el dato de afro'], c['lengua']),
                    ('Lenguas que hablan', ['Lengua', 'Alumnos'], [list(x) for x in c['lenguas']] or [['Ninguna registrada', 0]])],
         'anchos': (16,) + (12,) * 8},
        {'nombre': 'Lugar de nacimiento', 'titulo': 'Alumnos nacidos en otra entidad o en otro país', 'nota': 'Sale de la CURP (posiciones 12 y 13).',
         'tablas': [(None, ['Grado', 'Otra entidad H', 'Otra entidad M', 'Otra entidad total', 'Otro país H', 'Otro país M', 'Otro país total',
                            'Sin el dato'], c['nacidos']),
                    ('De qué entidad', ['Entidad', 'Alumnos'], [list(x) for x in c['entidades']] or [['Ninguno', 0]])],
         'anchos': (22,) + (12,) * 7},
        {'nombre': 'Becas', 'titulo': 'Alumnos con beca en dinero, por tipo y sexo', 'nota': 'Si un alumno tiene varias becas, cuenta la de mayor monto.',
         'tablas': [(None, ['Beca', 'Hombres', 'Mujeres', 'Total'], c['becas'])], 'anchos': (36, 10, 10, 10)},
        {'nombre': 'Procedencia', 'titulo': 'De dónde vienen los alumnos', 'nota': copia,
         'tablas': [(None, ['Procedencia', 'Hombres', 'Mujeres', 'Total'], c['procedencia'])], 'anchos': (44, 10, 10, 10)},
        {'nombre': 'Revisión', 'titulo': 'Antes de capturar: ¿cuadra todo?', 'nota': 'Si algo dice «No», corrígelo en SIDP y vuelve a bajar el 911.',
         'tablas': [(None, ['Revisión', '¿Bien?', 'Detalle'], revisar_911(registros, escuela, c))], 'anchos': (44, 8, 50)},
        {'nombre': 'Faltan', 'titulo': 'Alumnos a los que les falta un dato del 911', 'nota': CUIDADO,
         'tablas': [(None, ['Grupo', 'Nombre', 'CURP', 'Le falta'], faltan or [['', 'A nadie le falta nada', '', '']])], 'anchos': (8, 36, 21, 60)},
    ]
    return _excel(hojas)


# ---------------------------------------------------------------- Excel

def _excel(hojas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    libro = Workbook()
    libro.remove(libro.active)
    borde = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
    encabezado_relleno = PatternFill('solid', fgColor='1F4E79')
    total_relleno = PatternFill('solid', fgColor='DDEBF7')
    for hoja_datos in hojas:
        hoja = libro.create_sheet(re.sub(r'[\[\]:*?/\\]', '', hoja_datos['nombre'])[:31])
        hoja['A1'] = hoja_datos['titulo']
        hoja['A1'].font = Font(bold=True, size=14, color='1F4E79')
        renglon = 2
        if hoja_datos.get('nota'):
            hoja.cell(renglon, 1, hoja_datos['nota']).font = Font(italic=True, color='595959')
            renglon += 1
        renglon += 1
        primera = None
        for subtitulo, encabezado, filas in hoja_datos['tablas']:
            if subtitulo:
                hoja.cell(renglon, 1, subtitulo).font = Font(bold=True, size=12)
                renglon += 1
            for columna, texto in enumerate(encabezado, 1):
                celda = hoja.cell(renglon, columna, texto)
                celda.font, celda.fill, celda.border = Font(bold=True, color='FFFFFF'), encabezado_relleno, borde
                celda.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            primera = primera or renglon
            renglon += 1
            for fila in filas:
                total = str(fila[0]).startswith('Total') or (len(fila) > 1 and str(fila[1]) in ('Subtotal', 'Total'))
                for columna, valor in enumerate(fila, 1):
                    celda = hoja.cell(renglon, columna, valor)
                    celda.border = borde
                    celda.alignment = Alignment(wrap_text=isinstance(valor, str) and len(valor) > 30, vertical='top')
                    if total:
                        celda.font, celda.fill = Font(bold=True), total_relleno
                renglon += 1
            renglon += 1
        for columna, ancho in enumerate(hoja_datos.get('anchos') or (), 1):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        if len(hoja_datos['tablas']) == 1 and primera:
            hoja.freeze_panes = hoja.cell(primera + 1, 1)
        hoja.page_setup.orientation = 'landscape'
        hoja.page_setup.fitToWidth = 1
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_setup.fitToHeight = 0
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------- Word y PDF (los mismos bloques)

def _hojas_a_bloques(hojas, escuela):
    bloques = []
    for numero, hoja in enumerate(hojas):
        if numero:
            bloques.append(('salto',))
        bloques += [_encabezado(escuela, hoja.get('ciclo', '')), ('espacio', 1), ('titulo', hoja['titulo'])]
        for subtitulo, encabezado, filas in hoja['tablas']:
            if subtitulo:
                bloques.append(('parrafo', [(subtitulo, True)]))
            bloques.append(('tabla', encabezado, filas, hoja.get('anchos')))
        if hoja.get('nota'):
            bloques += [('espacio', 1), ('parrafo', [(hoja['nota'], False)], 'chico')]
    return bloques


def _escribir(formato, bloques, horizontal=False):
    if formato == 'docx':
        return docx(bloques, horizontal)
    if formato == 'pdf':
        return pdf(bloques, horizontal)
    raise ErrorDocumento('Ese formato no sale de este documento.')


def _run(texto, negrita=False, tamano=22, color=None):
    propiedades = '<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>' + ('<w:b/>' if negrita else '') + \
        (f'<w:color w:val="{color}"/>' if color else '') + f'<w:sz w:val="{tamano}"/>'
    return f'<w:r><w:rPr>{propiedades}</w:rPr><w:t xml:space="preserve">{escape(str(texto))}</w:t></w:r>'


def _parrafo(partes, alinear=None, tamano=22, despues=120, color=None):
    jc = {'centro': 'center', 'derecha': 'right', 'justificado': 'both'}.get(alinear)
    propiedades = (f'<w:jc w:val="{jc}"/>' if jc else '') + f'<w:spacing w:after="{despues}" w:line="276" w:lineRule="auto"/>'
    return f'<w:p><w:pPr>{propiedades}</w:pPr>{"".join(_run(t, n, tamano, color) for t, n in partes)}</w:p>'


def docx(bloques, horizontal=False):
    """Un .docx que abren Word, LibreOffice y Google Docs (carta, márgenes de 2 cm)."""
    ancho, alto = (15840, 12240) if horizontal else (12240, 15840)
    util = ancho - 2 * 1134
    cuerpo = []
    for bloque in bloques:
        tipo = bloque[0]
        if tipo == 'encabezado':
            lineas = bloque[1]
            cuerpo.append(_parrafo([(lineas[0], True)], 'centro', 28, 40, '1F4E79'))
            cuerpo += [_parrafo([(linea, False)], 'centro', 18, 20, '595959') for linea in lineas[1:]]
        elif tipo == 'titulo':
            cuerpo.append(_parrafo([(bloque[1], True)], 'centro', 26, 160))
        elif tipo == 'parrafo':
            alinear = bloque[2] if len(bloque) > 2 else None
            chico = alinear == 'chico'
            cuerpo.append(_parrafo(bloque[1], None if chico else alinear, 16 if chico else 22, 120, '595959' if chico else None))
        elif tipo == 'espacio':
            cuerpo += [_parrafo([('', False)], despues=0)] * bloque[1]
        elif tipo == 'salto':
            cuerpo.append('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
        elif tipo == 'firma':
            cuerpo.append(_parrafo([('_' * 36, False)], 'centro', 22, 0))
            cuerpo += [_parrafo([(linea, i == 0)], 'centro', 22, 0) for i, linea in enumerate(bloque[1])]
        elif tipo == 'tabla':
            _, encabezado, filas, anchos = bloque
            pesos = list(anchos or [1] * len(encabezado))[:len(encabezado)]
            columnas = [int(util * p / sum(pesos)) for p in pesos]
            rejilla = ''.join(f'<w:gridCol w:w="{c}"/>' for c in columnas)
            bordes = ''.join(f'<w:{lado} w:val="single" w:sz="4" w:space="0" w:color="A6A6A6"/>'
                             for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
            renglones = []
            for numero, fila in enumerate([encabezado] + [list(f) for f in filas]):
                celdas = []
                for c, valor in zip(columnas, fila):
                    sombra = '<w:shd w:val="clear" w:color="auto" w:fill="1F4E79"/>' if numero == 0 else ''
                    celdas.append(f'<w:tc><w:tcPr><w:tcW w:w="{c}" w:type="dxa"/>{sombra}</w:tcPr>'
                                  f'{_parrafo([("" if valor is None else valor, numero == 0)], None, 18, 0, "FFFFFF" if numero == 0 else None)}</w:tc>')
                repetir = '<w:trPr><w:tblHeader/></w:trPr>' if numero == 0 else ''
                renglones.append(f'<w:tr>{repetir}{"".join(celdas)}</w:tr>')
            cuerpo.append(f'<w:tbl><w:tblPr><w:tblW w:w="{sum(columnas)}" w:type="dxa"/><w:tblBorders>{bordes}</w:tblBorders>'
                          f'<w:tblCellMar><w:left w:w="60" w:type="dxa"/><w:right w:w="60" w:type="dxa"/></w:tblCellMar></w:tblPr>'
                          f'<w:tblGrid>{rejilla}</w:tblGrid>{"".join(renglones)}</w:tbl>')
            cuerpo.append(_parrafo([('', False)], despues=0))
    orientacion = ' w:orient="landscape"' if horizontal else ''
    seccion = (f'<w:sectPr><w:pgSz w:w="{ancho}" w:h="{alto}"{orientacion}/>'
               '<w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134" w:header="567" w:footer="567" w:gutter="0"/></w:sectPr>')
    documento = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                 f'<w:body>{"".join(cuerpo)}{seccion}</w:body></w:document>')
    tipos = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
             '<Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/word/document.xml" '
             'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    relaciones = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                  '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                  'Target="word/document.xml"/></Relationships>')
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, 'w', zipfile.ZIP_DEFLATED) as archivo:
        archivo.writestr('[Content_Types].xml', tipos)
        archivo.writestr('_rels/.rels', relaciones)
        archivo.writestr('word/document.xml', documento)
    return salida.getvalue()


def pdf(bloques, horizontal=False):
    if colors is None:
        raise ErrorDocumento('Para el PDF falta reportlab en esta computadora (python -m pip install reportlab). Mientras, bájalo en Word.')
    pagina = landscape(letter) if horizontal else letter
    salida = io.BytesIO()
    hoja = SimpleDocTemplate(salida, pagesize=pagina, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm,
                             title='SIDP', author='SIDP')
    azul = colors.HexColor('#1F4E79')
    estilos = {
        'escuela': ParagraphStyle('escuela', fontName='Helvetica-Bold', fontSize=14, leading=17, alignment=TA_CENTER, textColor=azul),
        'datos': ParagraphStyle('datos', fontName='Helvetica', fontSize=9, leading=11, alignment=TA_CENTER, textColor=colors.HexColor('#595959')),
        'titulo': ParagraphStyle('titulo', fontName='Helvetica-Bold', fontSize=13, leading=16, alignment=TA_CENTER, spaceAfter=8),
        'normal': ParagraphStyle('normal', fontName='Helvetica', fontSize=11, leading=16, spaceAfter=6),
        'chico': ParagraphStyle('chico', fontName='Helvetica', fontSize=8, leading=10, textColor=colors.HexColor('#595959')),
        'celda': ParagraphStyle('celda', fontName='Helvetica', fontSize=8, leading=9.5),
        'cabeza': ParagraphStyle('cabeza', fontName='Helvetica-Bold', fontSize=8, leading=9.5, textColor=colors.white),
    }
    alinear = {'centro': TA_CENTER, 'derecha': TA_RIGHT, 'justificado': 4}
    util = pagina[0] - 4 * cm
    historia = []
    for bloque in bloques:
        tipo = bloque[0]
        if tipo == 'encabezado':
            historia.append(Paragraph(escape(bloque[1][0]), estilos['escuela']))
            historia += [Paragraph(escape(linea), estilos['datos']) for linea in bloque[1][1:]]
        elif tipo == 'titulo':
            historia.append(Paragraph(escape(bloque[1]), estilos['titulo']))
        elif tipo == 'parrafo':
            como = bloque[2] if len(bloque) > 2 else None
            estilo = estilos['chico'] if como == 'chico' else ParagraphStyle(f'p{len(historia)}', parent=estilos['normal'],
                                                                              alignment=alinear.get(como, 0))
            texto = ''.join(f'<b>{escape(str(t))}</b>' if n else escape(str(t)) for t, n in bloque[1])
            historia.append(Paragraph(texto, estilo))
        elif tipo == 'espacio':
            historia.append(Spacer(1, 0.45 * cm * bloque[1]))
        elif tipo == 'salto':
            historia.append(PageBreak())
        elif tipo == 'firma':
            centro = ParagraphStyle(f'f{len(historia)}', parent=estilos['normal'], alignment=TA_CENTER, spaceAfter=0)
            historia.append(Paragraph('_' * 36, centro))
            historia += [Paragraph(f'<b>{escape(linea)}</b>' if i == 0 else escape(linea), centro) for i, linea in enumerate(bloque[1])]
        elif tipo == 'tabla':
            _, encabezado, filas, anchos = bloque
            pesos = list(anchos or [1] * len(encabezado))[:len(encabezado)]
            columnas = [util * p / sum(pesos) for p in pesos]
            datos = [[Paragraph(escape(str(t)), estilos['cabeza']) for t in encabezado]]
            datos += [[Paragraph(escape('' if v is None else str(v)), estilos['celda']) for v in fila] for fila in filas]
            tabla = Table(datos, colWidths=columnas, repeatRows=1)
            tabla.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), azul), ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#A6A6A6')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F6FA')]),
                ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]))
            historia += [tabla, Spacer(1, 0.3 * cm)]
    hoja.build(historia)
    return salida.getvalue()
