"""El Excel de la escuela por grado: lo que hay que cuidar de cada alumno, dónde viven, sus boletas y las tablas del 911.

Lo arma el servidor (SIDP.exe) con las listas de grupo de aquí y, si la PC 4 contesta, los datos que el escaneo sacó de
cada alumno (certificado médico, domicilio, boletas de años anteriores). Trae nombres y CURP: es un archivo de la
escuela, se descarga en la PC del administrador y no se sube a internet. No trae nada de documentos.

Hojas:
- Resumen: un renglón por grado y las gráficas que comparan los grados.
- Salud: un renglón por alumno con lo que la maestra debe saber (alergias, padecimientos, medicamentos, discapacidad), en rojo.
- Salud por grado: cuántos tienen alergias, padecimientos, medicamentos o discapacidad, y las alergias más comunes.
  Solo lo necesario (validacion/estandar.py): sin peso, talla, vista, dientes, sangre ni vacunas.
- Dónde viven: municipios, colonias y códigos postales.
- Boletas: el promedio de cada año anterior de cada alumno y el promedio por grado cursado.
- 911: alumnos por grado, sexo y edad al 1 de septiembre; grupos; discapacidad por tipo; nacidos en otra entidad o país.
- Todos y una hoja por grado (sus alumnos, sus grupos y una gráfica por alumno).
Cada alumno se cuenta con validacion/indicadores.py, igual que el análisis de la PC 4. Usa openpyxl, que va dentro de SIDP.exe.
"""
import datetime
import io

from validacion import indicadores

GRADOS = range(1, 7)
# (clave, encabezado, ancho): las columnas de la lista de alumnos, en las hojas "Todos" y de cada grado.
COLUMNAS = [
    ('grupo', 'Grupo', 8), ('primer_apellido', 'Primer apellido', 16), ('segundo_apellido', 'Segundo apellido', 16),
    ('nombres', 'Nombre(s)', 20), ('curp', 'CURP', 21), ('renapo', 'CURP en RENAPO', 10), ('sexo', 'Sexo', 7),
    ('fecha_nacimiento', 'Fecha de nacimiento', 12), ('edad', 'Edad al 1 de sept. (911)', 9), ('situacion', 'Edad para el grado', 13),
    ('repite', 'Repite el grado', 9), ('alergias', 'Alergias', 22), ('padecimientos', 'Padecimientos', 22),
    ('medicamentos', 'Medicamentos', 18), ('discapacidad', 'Discapacidad o condición (911)', 20), ('municipio', 'Municipio', 18),
    ('colonia', 'Colonia', 18), ('promedio', 'Promedio de la última boleta', 10), ('tutor', 'Tutor', 26), ('parentesco', 'Parentesco', 11),
    ('telefono', 'Teléfono', 13), ('por_confirmar', 'Datos por confirmar', 11),
]
SALUD = [
    ('grupo', 'Grupo', 8), ('primer_apellido', 'Primer apellido', 15), ('segundo_apellido', 'Segundo apellido', 15), ('nombres', 'Nombre(s)', 18),
    ('edad', 'Edad', 6), ('sexo', 'Sexo', 7), ('alergias', 'Alergias', 24),
    ('padecimientos', 'Enfermedades o padecimientos', 24), ('medicamentos', 'Medicamentos', 20), ('discapacidad', 'Discapacidad o condición', 20),
    ('tutor', 'Tutor', 24), ('telefono', 'Teléfono', 13), ('emergencia', 'Contacto de emergencia', 24),
]
IMPORTANTES = ('alergias', 'padecimientos', 'medicamentos', 'discapacidad')  # Lo que la maestra debe saber: en rojo.
SITUACIONES = ('En su edad', 'Extraedad', 'Menor edad', 'Sin fecha')
AZUL, VERDE, NARANJA, MORADO, ROJO, GRIS = '2F5597', '548235', 'C55A11', '7030A0', 'C00000', 'A6A6A6'


class ErrorLibro(Exception):
    """Falta algo para armar el Excel; el mensaje es para la persona."""


# ---------------------------------------------------------------- Los datos de cada alumno

def de_curp(curp):
    """(fecha de nacimiento, sexo) que trae la CURP: posiciones 5 a 10 (AAMMDD) y 11 (H o M)."""
    curp = (curp or '').upper()
    if len(curp) != 18:
        return None, None
    siglo = 2000 if curp[16].isalpha() else 1900  # El carácter 17 es letra para los nacidos desde el 2000.
    try:
        fecha = datetime.date(siglo + int(curp[4:6]), int(curp[6:8]), int(curp[8:10]))
    except ValueError:
        fecha = None
    return fecha, curp[10] if curp[10] in 'HM' else None


edad_al = indicadores.edad_al


def alumno(fila, anio, expediente=None, hoy=None):
    """Un renglón de la lista de aquí (grupos.exportar_alumnos) más el expediente de la PC 4, si lo hay."""
    grado, letra, primer, segundo, nombres, curp, renapo, sexo, nacimiento, tutor, parentesco, telefono = fila[:12]
    expediente = expediente or {}
    datos = expediente.get('datos') or {}
    ficha = expediente.get('alumno') or {}
    curp_fecha, curp_sexo = de_curp(curp)
    nacimiento = indicadores.fecha(nacimiento) or indicadores.fecha(ficha.get('fecha_nacimiento')) or curp_fecha
    sexo = sexo if sexo in ('H', 'M') else ficha.get('sexo') if ficha.get('sexo') in ('H', 'M') else curp_sexo
    ind = indicadores.de_alumno(datos, grado, sexo, nacimiento, curp, f'{anio}-{anio + 1}', hoy)
    edad_dic = edad_al(nacimiento, datetime.date(anio, 12, 31)) if nacimiento else None  # La regla de inscripción de la SEP.
    if edad_dic is None:
        situacion = 'Sin fecha'
    elif edad_dic - 5 - int(grado) >= 2:
        situacion = 'Extraedad'
    elif edad_dic - 5 < int(grado):
        situacion = 'Menor edad'
    else:
        situacion = 'En su edad'
    principal = (expediente.get('tutores') or [{}])[0]
    if not tutor and principal.get('nombres'):
        tutor = ' '.join(p for p in (principal.get('nombres'), principal.get('primer_apellido'), principal.get('segundo_apellido')) if p)
        parentesco = parentesco or principal.get('parentesco') or ''
        telefono = telefono or principal.get('telefono') or ''
    telefono = telefono or ', '.join(indicadores.valores_de(datos, 'contacto', 'telefono')[:2])
    unir = lambda valores: ', '.join(valores) or None  # noqa: E731
    return {
        'grado': int(grado), 'letra': letra, 'grupo': f'{grado}° {letra}', 'primer_apellido': primer, 'segundo_apellido': segundo,
        'nombres': nombres, 'curp': curp, 'renapo': renapo, 'sexo': {'H': 'Hombre', 'M': 'Mujer'}.get(sexo, ''),
        'fecha_nacimiento': nacimiento, 'edad': ind['edad_911'], 'situacion': situacion,
        'repite': {'sí': 'Sí', 'no': 'No'}.get(ind['repite']),
        'alergias': unir(ind['alergias']), 'padecimientos': unir(ind['padecimientos']), 'medicamentos': unir(ind['medicamentos']),
        'discapacidad': unir(ind['discapacidades']), 'emergencia': indicadores._uno(datos, 'contacto', 'nombre_emergencia'),
        'municipio': indicadores.titulo(ind['municipio']) if ind['municipio'] else None,
        'colonia': indicadores.titulo(ind['colonia']) if ind['colonia'] else None, 'cp': ind['cp'], 'nacido': ind['nacido'],
        'entidad': ind['entidad'], 'promedio': ind['ultimo_promedio'], **{f'promedio_{g}': ind['promedios'].get(g) for g in GRADOS},
        'tutor': tutor, 'parentesco': parentesco, 'telefono': telefono,
        'por_confirmar': sum(1 for lista in datos.values() for d in lista if d.get('estado') == 'por_confirmar'),
        'corto': f'{primer} {nombres.split()[0] if nombres else ""}'.strip(), 'ind': ind,
    }


def _promedio(valores):
    valores = [v for v in valores if v is not None]
    return round(sum(valores) / len(valores), 2) if valores else None


def _por_ciento(parte, todo):
    return round(100 * parte / todo, 1) if todo else None


def resumen_grado(alumnos):
    """Un renglón de la hoja Resumen: los números de un grado (o de la escuela)."""
    total = len(alumnos)
    return {
        'grupos': len({a['letra'] for a in alumnos}), 'alumnos': total,
        'hombres': sum(a['sexo'] == 'Hombre' for a in alumnos), 'mujeres': sum(a['sexo'] == 'Mujer' for a in alumnos),
        'edad': _promedio([a['edad'] for a in alumnos]),
        **{s: sum(a['situacion'] == s for a in alumnos) for s in SITUACIONES},
        'con_salud': sum(1 for a in alumnos if a['ind']['con_salud']),
        'con_alergias': sum(1 for a in alumnos if a['alergias']), 'con_padecimientos': sum(1 for a in alumnos if a['padecimientos']),
        'con_medicamento': sum(1 for a in alumnos if a['medicamentos']), 'con_discapacidad': sum(1 for a in alumnos if a['discapacidad']),
        'repiten': sum(1 for a in alumnos if a['repite'] == 'Sí'),
        'otra_entidad': sum(a['nacido'] in ('Otra entidad', 'Otro país') for a in alumnos),
        'promedio': _promedio([a['promedio'] for a in alumnos]),
        'bajos': sum(1 for a in alumnos if a['promedio'] is not None and a['promedio'] < indicadores.PROMEDIO_BAJO),
        'por_confirmar': sum(a['por_confirmar'] for a in alumnos),
    }


RESUMEN = [('grupos', 'Grupos'), ('alumnos', 'Alumnos'), ('hombres', 'Hombres'), ('mujeres', 'Mujeres'), ('edad', 'Edad promedio (1 sept.)'),
           ('En su edad', 'En su edad'), ('Extraedad', 'Extraedad'), ('Menor edad', 'Menor edad'), ('Sin fecha', 'Sin fecha de nacimiento'),
           ('con_salud', 'Con datos de salud'), ('con_alergias', 'Con alergias'), ('con_padecimientos', 'Con padecimientos'),
           ('con_medicamento', 'Toman medicamento'), ('con_discapacidad', 'Con discapacidad o condición'), ('repiten', 'Repiten el grado'),
           ('otra_entidad', 'Nacidos en otra entidad o país'), ('promedio', 'Promedio de la última boleta'), ('bajos', 'Promedio menor a 7'),
           ('por_confirmar', 'Datos por confirmar')]


# ---------------------------------------------------------------- El libro

def _estilos():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    linea = Side(style='thin', color='BFBFBF')
    return {
        'titulo': Font(bold=True, size=16, color=AZUL), 'sub': Font(italic=True, size=10, color='595959'),
        'encabezado': Font(bold=True, color='FFFFFF'), 'relleno': PatternFill('solid', fgColor=AZUL),
        'total': Font(bold=True), 'relleno_total': PatternFill('solid', fgColor='DDEBF7'),
        'importante': PatternFill('solid', fgColor='F8CBAD'), 'negrita': Font(bold=True, color=ROJO),
        'borde': Border(left=linea, right=linea, top=linea, bottom=linea),
        'centro': Alignment(horizontal='center', vertical='center', wrap_text=True),
    }


def _tabla(hoja, fila, columna, encabezado, renglones, estilos, total=False):
    """Escribe una tabla con encabezado; devuelve (primera fila de datos, última fila)."""
    for j, texto in enumerate(encabezado):
        celda = hoja.cell(fila, columna + j, texto)
        celda.font, celda.fill, celda.alignment, celda.border = estilos['encabezado'], estilos['relleno'], estilos['centro'], estilos['borde']
    for i, renglon in enumerate(renglones, 1):
        for j, valor in enumerate(renglon):
            celda = hoja.cell(fila + i, columna + j, valor)
            celda.border = estilos['borde']
            if isinstance(valor, datetime.date):
                celda.number_format = 'dd/mm/yyyy'
            if total and i == len(renglones):
                celda.font, celda.fill = estilos['total'], estilos['relleno_total']
    return fila + 1, fila + len(renglones)


def _titulo(hoja, texto, sub, estilos):
    hoja['A1'] = texto
    hoja['A1'].font = estilos['titulo']
    if sub:
        hoja['A2'] = sub
        hoja['A2'].font = estilos['sub']


def _subtitulo(hoja, fila, texto, estilos):
    hoja.cell(fila, 1, texto).font = estilos['total']


def _grafica(tipo, titulo, eje_x='', eje_y='', apilada=False, horizontal=False, ancho=16, alto=8):
    from openpyxl.chart import BarChart, PieChart, ScatterChart
    grafica = {'barras': BarChart, 'pastel': PieChart, 'puntos': ScatterChart}[tipo]()
    grafica.title, grafica.width, grafica.height = titulo, ancho, alto
    grafica.visible_cells_only = False  # Los nombres cortos de las gráficas por alumno van en una columna escondida.
    if tipo == 'barras':
        grafica.type = 'bar' if horizontal else 'col'
        if horizontal:  # El primero de la lista arriba, como en la tabla.
            grafica.x_axis.scaling.orientation = 'maxMin'
        if apilada:
            grafica.grouping, grafica.overlap = 'stacked', 100
    if tipo != 'pastel':
        grafica.x_axis.title, grafica.y_axis.title = eje_x, eje_y
        grafica.x_axis.delete = grafica.y_axis.delete = False  # openpyxl 3.1 esconde los ejes si no se dice.
        grafica.legend.position = 'b'
    return grafica


def _series(grafica, hoja, columnas, fila_titulo, primera, ultima, col_categorias, colores=()):
    from openpyxl.chart import Reference
    for indice, columna in enumerate(columnas):
        grafica.add_data(Reference(hoja, min_col=columna, min_row=fila_titulo, max_row=ultima), titles_from_data=True)
        if indice < len(colores):
            grafica.series[-1].graphicalProperties.solidFill = colores[indice]
            grafica.series[-1].graphicalProperties.line.solidFill = colores[indice]
    grafica.set_categories(Reference(hoja, min_col=col_categorias, min_row=primera, max_row=ultima))
    return grafica


def _anchos(hoja, anchos, desde=1):
    from openpyxl.utils import get_column_letter
    for j, ancho in enumerate(anchos):
        hoja.column_dimensions[get_column_letter(desde + j)].width = ancho


def _poner(hoja, graficas, fila, columnas=('A', 'J'), alto=17):
    for i, grafica in enumerate(graficas):
        hoja.add_chart(grafica, f'{columnas[i % len(columnas)]}{fila + (i // len(columnas)) * alto}')


def _lista(hoja, fila, alumnos, estilos, columnas=COLUMNAS, congelar=5):
    """La lista de alumnos con sus columnas; devuelve (primera fila de datos, última)."""
    renglones = [[a[clave] for clave, _, _ in columnas] for a in alumnos]
    primera, ultima = _tabla(hoja, fila, 1, [titulo for _, titulo, _ in columnas], renglones, estilos)
    _anchos(hoja, [ancho for _, _, ancho in columnas])
    hoja.row_dimensions[fila].height = 45
    if renglones:
        from openpyxl.utils import get_column_letter
        hoja.auto_filter.ref = f'A{fila}:{get_column_letter(len(columnas))}{ultima}'
        for i, a in enumerate(alumnos):  # Lo que la maestra debe saber, en rojo.
            for j, (clave, _, _) in enumerate(columnas, 1):
                if clave in IMPORTANTES and a[clave]:
                    celda = hoja.cell(primera + i, j)
                    celda.fill, celda.font = estilos['importante'], estilos['negrita']
    hoja.freeze_panes = hoja.cell(fila + 1, congelar)
    return primera, ultima


def _col(clave, columnas=COLUMNAS):
    return next(i for i, (c, _, _) in enumerate(columnas, 1) if c == clave)


def _conteo(valores, cuantos=None):
    cuenta = {}
    for valor in valores:
        if valor:
            cuenta[valor] = cuenta.get(valor, 0) + 1
    salida = sorted(cuenta.items(), key=lambda x: (-x[1], x[0]))
    return [list(x) for x in salida[:cuantos] if True] if cuantos else [list(x) for x in salida]


def _hoja_resumen(libro, alumnos, ciclo, hoy, estilos, escuela):
    hoja = libro.active
    hoja.title = 'Resumen'
    _titulo(hoja, f'{escuela} · ciclo {ciclo}' if escuela else f'Ciclo escolar {ciclo}',
            f'Hecho por SIDP el {hoy:%d/%m/%Y}. Trae datos personales de alumnos: no lo mandes por correo ni lo subas a internet. '
            'Lo que dice "por confirmar" salió del escaneo y todavía no lo revisa una persona.', estilos)
    renglones = [[f'{grado}°'] + [resumen_grado([a for a in alumnos if a['grado'] == grado])[c] for c, _ in RESUMEN] for grado in GRADOS]
    total = resumen_grado(alumnos)
    renglones.append(['Escuela'] + [total[c] for c, _ in RESUMEN])
    primera, ultima = _tabla(hoja, 4, 1, ['Grado'] + [t for _, t in RESUMEN], renglones, estilos, total=True)
    hoja.row_dimensions[4].height = 60
    _anchos(hoja, [9] + [11] * len(RESUMEN))
    ultima_grado = ultima - 1  # Las gráficas comparan los grados, sin el renglón de la escuela.
    columna = {c: i for i, (c, _) in enumerate(RESUMEN, 2)}
    serie = lambda titulo, eje, claves, colores, apilada=False: _series(  # noqa: E731
        _grafica('barras', titulo, 'Grado', eje, apilada=apilada), hoja, [columna[c] for c in claves], 4, primera, ultima_grado, 1, colores)
    graficas = [
        serie('Alumnos por grado', 'Alumnos', ['hombres', 'mujeres'], (AZUL, NARANJA), True),
        serie('Edad para el grado', 'Alumnos', list(SITUACIONES), (VERDE, NARANJA, AZUL, GRIS), True),
        serie('Alergias, padecimientos y medicamentos', 'Alumnos', ['con_alergias', 'con_padecimientos', 'con_medicamento'], (ROJO, MORADO, NARANJA)),
        serie('Promedio de la última boleta', 'Promedio', ['promedio'], (AZUL,)),
        serie('Discapacidad y nacidos fuera del estado', 'Alumnos', ['con_discapacidad', 'otra_entidad'], (MORADO, GRIS)),
    ]
    _poner(hoja, graficas, ultima + 3)


def _hoja_salud(libro, alumnos, estilos):
    hoja = libro.create_sheet('Salud')
    _titulo(hoja, 'La salud de cada alumno',
            'Del certificado médico que se escaneó. En rojo: alergias, padecimientos, medicamentos y discapacidad, lo que su maestra debe '
            'saber. Usa el filtro de cada columna para ver, por ejemplo, solo los que tienen alergia.',
            estilos)
    _lista(hoja, 4, alumnos, estilos, SALUD, congelar=5)


def _hoja_salud_por_grado(libro, alumnos, estilos):
    hoja = libro.create_sheet('Salud por grado')
    _titulo(hoja, 'Salud por grado', 'Cuántos alumnos de cada grado tienen algo que su maestra debe cuidar.', estilos)
    claves = (('con_alergias', 'Con alergias'), ('con_padecimientos', 'Con padecimientos'), ('con_medicamento', 'Toman medicamento'),
              ('con_discapacidad', 'Con discapacidad o condición'))
    renglones = [[f'{g}°'] + [resumen_grado([a for a in alumnos if a['grado'] == g])[c] for c, _ in claves] for g in GRADOS]
    renglones.append(['Escuela'] + [sum(r[j] for r in renglones) for j in range(1, len(claves) + 1)])
    p, u = _tabla(hoja, 5, 1, ['Grado'] + [n for _, n in claves], renglones, estilos, total=True)
    graficas = [_series(_grafica('barras', 'Lo que hay que cuidar, por grado', 'Grado', 'Alumnos'), hoja, list(range(2, len(claves) + 2)), 5, p,
                        u - 1, 1, (ROJO, MORADO, NARANJA, AZUL))]
    fila = u + 3
    for titulo, clave in (('Alergias más comunes', 'alergias'), ('Enfermedades o padecimientos más comunes', 'padecimientos'),
                          ('Medicamentos más comunes', 'medicamentos')):
        _subtitulo(hoja, fila - 1, titulo, estilos)
        cuenta = _conteo([v for a in alumnos for v in a['ind'][clave]], 10) or [['Ninguno anotado', 0]]
        p, u = _tabla(hoja, fila, 1, ['Cuál', 'Alumnos'], cuenta, estilos)
        graficas.append(_series(_grafica('barras', titulo, '', 'Alumnos', horizontal=True), hoja, [2], fila, p, u, 1, (ROJO,)))
        fila = u + 3
    _anchos(hoja, [26] + [14] * 4)
    _poner(hoja, graficas, 5, ('H', 'Q'))


def _hoja_domicilio(libro, alumnos, estilos):
    hoja = libro.create_sheet('Dónde viven')
    con_dato = [a for a in alumnos if a['ind']['con_domicilio']]
    _titulo(hoja, 'Dónde viven los alumnos', f'Del comprobante de domicilio o de la INE del tutor: {len(con_dato)} de {len(alumnos)} alumnos tienen el dato.',
            estilos)
    graficas, fila = [], 5
    for titulo, clave, cuantos in (('Por municipio', 'municipio', None), ('Colonias con más alumnos', 'colonia', 20), ('Por código postal', 'cp', 15)):
        _subtitulo(hoja, fila - 1, titulo, estilos)
        cuenta = _conteo([a[clave] for a in alumnos], cuantos) or [['Sin dato', 0]]
        total = sum(n for _, n in cuenta)
        p, u = _tabla(hoja, fila, 1, [titulo.split(' ')[-1].capitalize(), 'Alumnos', '%'], [[c, n, _por_ciento(n, total)] for c, n in cuenta], estilos)
        graficas.append(_series(_grafica('barras', titulo, '', 'Alumnos', horizontal=True, alto=max(8, 0.5 * len(cuenta) + 3)), hoja, [2], fila,
                                p, u, 1, (AZUL,)))
        fila = u + 3
    _anchos(hoja, [30, 10, 8])
    _poner(hoja, graficas, 5, ('F',), 20)


def _hoja_boletas(libro, alumnos, estilos):
    hoja = libro.create_sheet('Boletas')
    _titulo(hoja, 'Boletas de años anteriores', 'El promedio final de cada boleta que se escaneó (un alumno de 4.° trae las de 1.°, 2.° y 3.°). '
            f'En rojo, menos de {indicadores.PROMEDIO_BAJO:g}.', estilos)
    columnas = ['Grupo', 'Primer apellido', 'Segundo apellido', 'Nombre(s)'] + [f'Boleta de {g}°' for g in range(1, 6)] + ['Última']
    renglones = [[a['grupo'], a['primer_apellido'], a['segundo_apellido'], a['nombres']] + [a[f'promedio_{g}'] for g in range(1, 6)] + [a['promedio']]
                 for a in alumnos if a['grado'] > 1]
    p, u = _tabla(hoja, 4, 1, columnas, renglones, estilos)
    for fila in range(p, u + 1):
        for columna in range(5, 11):
            celda = hoja.cell(fila, columna)
            if isinstance(celda.value, (int, float)) and celda.value < indicadores.PROMEDIO_BAJO:
                celda.fill, celda.font = estilos['importante'], estilos['negrita']
    hoja.freeze_panes = 'E5'
    _anchos(hoja, [8, 16, 16, 20] + [10] * 6)
    fila = max(u, 4) + 3
    _subtitulo(hoja, fila - 1, 'Promedio por grado cursado', estilos)
    por_grado = [[f'{g}°', _promedio([a[f'promedio_{g}'] for a in alumnos]), sum(1 for a in alumnos if a[f'promedio_{g}'] is not None)]
                 for g in range(1, 6)]
    gp, gu = _tabla(hoja, fila, 1, ['Grado cursado', 'Promedio', 'Boletas'], por_grado, estilos)
    fila = gu + 3
    _subtitulo(hoja, fila - 1, 'Por campo formativo (última boleta)', estilos)
    formativos = [[nombre, _promedio([a['ind']['formativos'].get(clave) for a in alumnos])] for clave, nombre in indicadores.FORMATIVOS]
    fp, fu = _tabla(hoja, fila, 1, ['Campo formativo', 'Promedio'], formativos, estilos)
    _poner(hoja, [_series(_grafica('barras', 'Promedio por grado cursado', 'Grado', 'Promedio'), hoja, [2], gp - 1, gp, gu, 1, (AZUL,)),
                  _series(_grafica('barras', 'Promedio por campo formativo', '', 'Promedio', horizontal=True), hoja, [2], fp - 1, fp, fu, 1, (VERDE,))],
           gp - 1, ('L',), 17)


def _hoja_911(libro, alumnos, estilos):
    hoja = libro.create_sheet('911')
    _titulo(hoja, 'Para la estadística 911 (cuestionario 911.3, primaria)',
            'Los números ya contados para capturarlos en f911.sep.gob.mx. La edad es la cumplida al 1 de septiembre. El 911 completo (repetidores, '
            'becas, lengua indígena, de dónde vienen) sale en «Formatos oficiales» > Estadística 911.', estilos)
    edades = list(indicadores.EDADES_911)
    renglones = []
    for grado in GRADOS:
        del_grado = [a for a in alumnos if a['grado'] == grado]
        for sexo, nombre in (('Hombre', 'Hombres'), ('Mujer', 'Mujeres'), (None, 'Subtotal')):
            parte = [a for a in del_grado if sexo is None or a['sexo'] == sexo]
            renglones.append([f'{grado}°', nombre] + [sum(a['ind']['columna_911'] == e for a in parte) for e in edades]
                             + [sum(a['ind']['columna_911'] is None for a in parte), len(parte)])
    renglones.append(['Total', ''] + [sum(r[j] for r in renglones if r[1] == 'Subtotal') for j in range(2, len(edades) + 4)])
    p, u = _tabla(hoja, 4, 1, ['Grado', 'Sexo'] + [f'{e} años' if e[0].isdigit() and ' ' not in e else e for e in edades] + ['Sin fecha', 'Total'],
                  renglones, estilos, total=True)
    for fila in range(p, u):
        if hoja.cell(fila, 2).value == 'Subtotal':
            for columna in range(1, len(edades) + 5):
                hoja.cell(fila, columna).font, hoja.cell(fila, columna).fill = estilos['total'], estilos['relleno_total']
    fila = u + 3
    _subtitulo(hoja, fila - 1, 'Grupos por grado', estilos)
    _, u = _tabla(hoja, fila, 1, ['Grado', 'Grupos'], [[f'{g}°', len({a['letra'] for a in alumnos if a['grado'] == g})] for g in GRADOS], estilos)
    fila = u + 3
    por_grado_sexo = [(g, s) for g in GRADOS for s in ('Hombre', 'Mujer')]
    encabezado = ['Tipo'] + [f'{g}° {s[0]}' for g, s in por_grado_sexo] + ['Total']
    _subtitulo(hoja, fila - 1, 'Discapacidad, condiciones y aptitudes sobresalientes (H hombres, M mujeres)', estilos)
    tipos = [t for t, _ in indicadores.DISCAPACIDAD_911] + [indicadores.OTRAS_CONDICIONES]
    renglones = [[t] + [sum(t in a['ind']['discapacidades'] for a in alumnos if a['grado'] == g and a['sexo'] == s) for g, s in por_grado_sexo]
                 + [sum(t in a['ind']['discapacidades'] for a in alumnos)] for t in tipos]
    _, u = _tabla(hoja, fila, 1, encabezado, renglones, estilos)
    fila = u + 3
    _subtitulo(hoja, fila - 1, 'Lugar de nacimiento (sale de la CURP)', estilos)
    renglones = [[lugar] + [sum(a['nacido'] == lugar for a in alumnos if a['grado'] == g and a['sexo'] == s) for g, s in por_grado_sexo]
                 + [sum(a['nacido'] == lugar for a in alumnos)] for lugar in ('Estado de México', 'Otra entidad', 'Otro país')]
    _, u = _tabla(hoja, fila, 1, encabezado, renglones, estilos)
    fila = u + 3
    _subtitulo(hoja, fila - 1, 'Nacidos en otra entidad, por entidad', estilos)
    _tabla(hoja, fila, 1, ['Entidad', 'Alumnos'], _conteo([a['entidad'] for a in alumnos if a['nacido'] == 'Otra entidad']) or [['Ninguno', 0]],
           estilos)
    _anchos(hoja, [34, 10] + [9] * 13)
    hoja.row_dimensions[4].height = 30


def _hoja_todos(libro, alumnos, estilos):
    hoja = libro.create_sheet('Todos')
    _titulo(hoja, 'Todos los alumnos', None, estilos)
    _lista(hoja, 3, alumnos, estilos)


def _hoja_grado(libro, grado, alumnos, estilos):
    hoja = libro.create_sheet(f'{grado}°')
    hoja['A1'] = f'{grado}° grado: {len(alumnos)} alumnos'
    hoja['A1'].font = estilos['titulo']
    if not alumnos:
        hoja['A2'] = 'No hay alumnos en este grado.'
        return
    primera, ultima = _lista(hoja, 3, alumnos, estilos)
    # Por grupo (A, B...): para comparar los grupos del grado.
    fila = ultima + 3
    _subtitulo(hoja, fila - 1, 'Por grupo', estilos)
    renglones = []
    for letra in sorted({a['letra'] for a in alumnos}):
        n = resumen_grado([a for a in alumnos if a['letra'] == letra])
        renglones.append([f'{grado}° {letra}', n['alumnos'], n['hombres'], n['mujeres'], n['edad'], n['Extraedad'], n['con_alergias'],
                          n['con_padecimientos'], n['repiten'], n['promedio']])
    g_primera, g_ultima = _tabla(hoja, fila, 1, ['Grupo', 'Alumnos', 'Hombres', 'Mujeres', 'Edad promedio', 'Extraedad', 'Con alergias',
                                                 'Con padecimientos', 'Repiten el grado', 'Promedio de la boleta'],
                                 renglones, estilos)
    # Edades del grado (tabla chica para su gráfica).
    edades = sorted({a['edad'] for a in alumnos if a['edad'] is not None})
    fila_edad = g_ultima + 3
    _subtitulo(hoja, fila_edad - 1, 'Edades (al 1 de septiembre)', estilos)
    e_primera, e_ultima = _tabla(hoja, fila_edad, 1, ['Edad', 'Hombres', 'Mujeres'],
                                 [[f'{e} años', sum(a['edad'] == e and a['sexo'] == 'Hombre' for a in alumnos),
                                   sum(a['edad'] == e and a['sexo'] == 'Mujer' for a in alumnos)] for e in edades] or [['Sin dato', 0, 0]], estilos)
    alto = max(8, 0.6 * len(alumnos) + 3)  # Una barra por alumno, con su nombre.
    nombre = len(COLUMNAS) + 2  # Columna oculta con el nombre corto para las gráficas por alumno.
    hoja.cell(3, nombre, 'Alumno')
    for i, a in enumerate(alumnos):
        hoja.cell(primera + i, nombre, a['corto'])
    from openpyxl.utils import get_column_letter
    hoja.column_dimensions[get_column_letter(nombre)].hidden = True
    por_alumno = lambda titulo, eje, clave, color: _series(  # noqa: E731
        _grafica('barras', titulo, '', eje, horizontal=True, alto=alto), hoja, [_col(clave)], 3, primera, ultima, nombre, (color,))
    graficas = []
    if any(a['promedio'] for a in alumnos):
        graficas.append(por_alumno('Promedio de la última boleta de cada alumno', 'Promedio', 'promedio', AZUL))
    fila_graficas = e_ultima + 3
    hoja.add_chart(_series(_grafica('barras', 'Hombres y mujeres por grupo', 'Grupo', 'Alumnos', apilada=True), hoja, [3, 4], fila, g_primera,
                           g_ultima, 1, (AZUL, NARANJA)), f'A{fila_graficas}')
    hoja.add_chart(_series(_grafica('barras', 'Edades del grado', 'Edad', 'Alumnos', apilada=True), hoja, [2, 3], fila_edad, e_primera, e_ultima, 1,
                           (AZUL, NARANJA)), f'J{fila_graficas}')
    _poner(hoja, graficas, fila_graficas + 17, alto=int(alto * 2 + 2))  # Dos por renglón; cada renglón mide lo que la gráfica.


def libro(alumnos, ciclo, hoy=None, escuela=None):
    """Los bytes del .xlsx. alumnos: lo que devuelve alumno() para cada uno."""
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise ErrorLibro('Falta openpyxl para hacer el Excel con gráficas (python -m pip install openpyxl).') from error
    hoy = hoy or datetime.date.today()
    alumnos = sorted(alumnos, key=lambda a: (a['grado'], a['letra'], a['primer_apellido'] or '', a['segundo_apellido'] or '', a['nombres'] or ''))
    estilos = _estilos()
    trabajo = Workbook()
    _hoja_resumen(trabajo, alumnos, ciclo, hoy, estilos, escuela)
    _hoja_salud(trabajo, alumnos, estilos)
    _hoja_salud_por_grado(trabajo, alumnos, estilos)
    _hoja_domicilio(trabajo, alumnos, estilos)
    _hoja_boletas(trabajo, alumnos, estilos)
    _hoja_911(trabajo, alumnos, estilos)
    _hoja_todos(trabajo, alumnos, estilos)
    for grado in GRADOS:
        _hoja_grado(trabajo, grado, [a for a in alumnos if a['grado'] == grado], estilos)
    trabajo.properties.creator = 'SIDP'
    salida = io.BytesIO()
    trabajo.save(salida)
    return salida.getvalue()
