"""Los formatos que la escuela entrega a la supervisión, en Excel y con la forma que ya usan.

Son los libros que hoy se llenan a mano, uno por uno, con el mismo padrón (lo vio el agente en la computadora de la
dirección: cerebro/laboratorio/resultados/2026-10-25_analisis_formatos_directivos_profesores.md):

- El Metadato de 71 columnas: alumno (1 a 40) y tutor (41 a 71).
- La base del 911 que se calcula sola: se escribe o se pega la CURP y salen el sexo, la fecha, la edad al corte y los
  cuadros por grado, sexo y edad. Si se agrega un alumno a mano, los totales cambian solos.
- El censo de condición y discapacidad, una hoja por grado.
- Las plantillas de calificaciones (con la dispersión que se calcula sola), de la boleta con observaciones y de
  asistencias, con los alumnos ya puestos: el maestro solo escribe calificaciones y faltas.

Solo se llena lo que SIDP tiene del estándar (validacion/estandar.py). Peso, talla, tipo de sangre y vacunas quedan en
blanco por la decisión «Solo lo necesario» del 23 de octubre de 2026; las demás columnas vacías las llena la escuela.
"""
import datetime
import io
import re

from validacion import indicadores

GRADOS = ('PRIMERO', 'SEGUNDO', 'TERCERO', 'CUARTO', 'QUINTO', 'SEXTO')
FILAS_LIBRES = 40  # Renglones con fórmulas de sobra en la base del 911, para alumnos que lleguen después.
FINAL = 600  # Hasta qué renglón cuentan las fórmulas del 911.

METADATO_ALUMNO = (
    'NP', 'GRADO Y GRUPO', 'APELLIDO PATERNO', 'APELLIDO MATERNO', 'NOMBRE (S)', 'GÉNERO', 'CURP', 'FECHA DE NACIMIENTO', 'EDAD',
    'LUGAR DE NACIMIENTO', 'NACIONALIDAD', 'TIPO DE SANGRE', 'PESO', 'TALLA', 'ENFERMEDAD CRÓNICA O PADECIMIENTO', 'ALERGIAS',
    'PIE PLANO, USA LENTES, APARATO AUDITIVO O PRÓTESIS', 'CUENTA CON SERVICIO MÉDICO', 'INSTITUCIÓN', 'CALLE', 'NO. EXTERIOR',
    'NO. INTERIOR', 'COLONIA', 'C. P.', 'LOCALIDAD', 'MUNICIPIO', 'ESTADO', 'ENTRE QUÉ CALLES', 'REFERENCIA DEL DOMICILIO',
    'PRESENTA ALGUNA CONDICIÓN', 'PRESENTA ALGUNA DISCAPACIDAD', 'ENTREGÓ CERTIFICADO MÉDICO', 'CUENTA CON ESQUEMA DE VACUNACIÓN COMPLETO',
    'HABLA ALGUNA LENGUA INDÍGENA', 'PERTENECE A ALGÚN GRUPO ÉTNICO O INDÍGENA', 'CUENTA CON EQUIPO DE CÓMPUTO O DISPOSITIVO INTELIGENTE',
    'SERVICIOS CON LOS QUE CUENTA EN CASA', 'VIVE CON AMBOS PADRES, SOLO MADRE, SOLO PADRE O TUTOR', 'SITUACIÓN DEL ALUMNO',
    'OBSERVACIONES DEL ALUMNO')
METADATO_TUTOR = (
    'APELLIDO PATERNO', 'APELLIDO MATERNO', 'NOMBRE (S)', 'PARENTESCO', 'CURP', 'FECHA DE NACIMIENTO', 'EDAD', 'GÉNERO',
    'LUGAR DE NACIMIENTO', 'NACIONALIDAD', 'ESTADO CIVIL', 'GRADO MÁXIMO DE ESTUDIOS', 'OCUPACIÓN', 'RAMA DE ACTIVIDAD ECONÓMICA',
    'HABLA ALGUNA LENGUA INDÍGENA', 'CALLE', 'NO. EXTERIOR', 'NO. INTERIOR', 'COLONIA', 'C. P.', 'LOCALIDAD', 'MUNICIPIO', 'ESTADO',
    'TELÉFONO FIJO (CASA)', 'TELÉFONO CELULAR DEL TUTOR', 'TELÉFONO CELULAR DE ALGÚN FAMILIAR', 'CORREO ELECTRÓNICO',
    'NOMBRE COMPLETO DE CONTACTO EN CASO DE EMERGENCIA (Y PARENTESCO)', 'TELÉFONO DE EMERGENCIA', 'ENTREGÓ DOCUMENTACIÓN COMPLETA',
    'OBSERVACIONES DEL TUTOR')
# Columnas del censo de condición y discapacidad, y el tipo del 911 (indicadores.DISCAPACIDAD_911) que marca cada una.
CONDICION = (('COMUNICACIÓN', None), ('LENGUAJE', None), ('CONDUCTA', None), ('APRENDIZAJE', None), ('TDA', 'TDAH'),
             ('TRASTORNOS DEL ESPECTRO AUTISTA', 'Trastorno del espectro autista'), ('CRÓNICAS DEGENERATIVAS', None))
DISCAPACIDAD = (('MOTRIZ', 'Discapacidad motriz'), ('BAJA VISIÓN', 'Baja visión'), ('CEGUERA', 'Ceguera'), ('SORDERA', 'Sordera'),
                ('SORDOCEGUERA', 'Sordoceguera'), ('HIPOACUSIA', 'Hipoacusia'), ('INTELECTUAL', 'Discapacidad intelectual'),
                ('PSICOSOCIAL', 'Discapacidad psicosocial'))
CONDICIONES = {'TDAH', 'Trastorno del espectro autista', 'Aptitudes sobresalientes', indicadores.OTRAS_CONDICIONES}
CAMPOS_NEM = ('LENGUAJES', 'SABERES Y PENSAMIENTO CIENTÍFICO', 'ÉTICA, NATURALEZA Y SOCIEDADES', 'DE LO HUMANO Y LO COMUNITARIO')
DISPERSION = (('ESPERADO', '10 a 8.0', 8, 10.01), ('EN DESARROLLO', '7.9 a 7.0', 7, 8), ('REQUIERE APOYO', '6.9 a 6.0', 6, 7),
              ('REQUIERE APOYO (RIESGO)', '5.9 a 5.0', 5, 6))
MESES_CICLO = ('SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE', 'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO')
PARTICULAS = {'DE', 'DEL', 'LA', 'LAS', 'LOS', 'Y', 'MC', 'MAC', 'VAN', 'VON', 'SAN', 'SANTA'}


# ---------------------------------------------------------------- Lo que sale de la CURP y del registro

def _fecha_curp(curp):
    curp = (curp or '').upper()
    if len(curp) != 18:
        return None
    try:
        return datetime.date((2000 if curp[16].isalpha() else 1900) + int(curp[4:6]), int(curp[6:8]), int(curp[8:10]))
    except ValueError:
        return None


def _genero(sexo):
    return {'H': 'HOMBRE', 'M': 'MUJER'}.get(sexo or '', '')


def _entidad(curp):
    clave = (curp or '').upper()[11:13]
    return indicadores.ENTIDADES.get(clave, '')


def _nacionalidad(curp):
    clave = (curp or '').upper()[11:13]
    return '' if clave not in indicadores.ENTIDADES else 'EXTRANJERA' if clave == 'NE' else 'MEXICANA'


def partir_nombre(completo, curp=None):
    """(primer apellido, segundo apellido, nombres) de un nombre completo. Con la CURP, sus letras dicen cuál es cuál;
    sin ella, los nombres van primero (como en la INE) y "DE LA" se queda con su apellido."""
    texto = indicadores.normal(completo)
    if not texto:
        return '', '', ''
    if curp:
        from validacion import extraer
        hallado = extraer.nombre_por_curp(texto, curp)
        if hallado:
            return hallado['primer_apellido'], hallado['segundo_apellido'], hallado['nombres']
    palabras = texto.split()
    if len(palabras) < 3:
        return (palabras[-1], '', palabras[0]) if len(palabras) == 2 else ('', '', palabras[0])

    def apellido(lista):
        inicio = len(lista) - 1
        while inicio > 1 and lista[inicio - 1] in PARTICULAS:
            inicio -= 1
        return ' '.join(lista[inicio:]), lista[:inicio]
    segundo, resto = apellido(palabras)
    primero, resto = apellido(resto) if len(resto) > 1 else (' '.join(resto), [])
    return primero, segundo, ' '.join(resto)


def _ninguna(valores, con_salud):
    return ', '.join(v.upper() for v in valores) if valores else ('NINGUNA' if con_salud else '')


def _situacion(r):
    if r.get('repite') == 'sí':
        return 'REPETIDOR'
    if r.get('procedencia') in ('otra_edomex', 'otra_entidad', 'otro_pais'):
        return 'ALTA POR TRASLADO'
    if r.get('repite') == 'no':
        return 'NUEVO INGRESO' if r.get('grado') == 1 else 'PROMOVIDO'
    return ''


def _domicilio(r):
    from validacion import campos
    texto, partes = campos.domicilio(r.get('domicilio') or '')
    # Si no se distingue la calle, va todo el domicilio en «CALLE» para que no se pierda.
    return {'calle': partes.get('calle') or texto or '', 'numero': partes.get('numero', '') if partes.get('calle') else '', 'colonia': (r.get('colonia') or partes.get('colonia') or '').upper(),
            'cp': r.get('cp') or partes.get('cp', ''), 'municipio': (r.get('municipio') or partes.get('municipio') or '').upper()}


def _ordenados(registros):
    return sorted(registros, key=lambda x: (x.get('grado') or 9, x.get('letra') or '', x.get('primer_apellido') or '',
                                            x.get('segundo_apellido') or '', x.get('nombres') or ''))


def _grupos(registros):
    salida = {}
    for r in _ordenados(registros):
        if r.get('grado'):
            salida.setdefault((r['grado'], r.get('letra') or ''), []).append(r)
    return salida


def _corte(ciclo, hoy):
    """El 1 de septiembre del año en que empieza el ciclo (la edad del 911)."""
    hallado = re.match(r'(\d{4})', ciclo or '')
    return datetime.date(int(hallado[1]) if hallado else (hoy.year if hoy.month >= 8 else hoy.year - 1), 9, 1)


# ---------------------------------------------------------------- El estilo: sobrio, como un formato oficial

class _Estilo:
    def __init__(self):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        linea = Side(style='thin', color='000000')
        self.borde = Border(left=linea, right=linea, top=linea, bottom=linea)
        self.normal = Font(name='Arial', size=9)
        self.negrita = Font(name='Arial', size=9, bold=True)
        self.titulo = Font(name='Arial', size=12, bold=True)
        self.membrete = Font(name='Arial', size=10, bold=True)
        self.nota = Font(name='Arial', size=8, italic=True, color='404040')
        self.gris = PatternFill('solid', fgColor='D9D9D9')
        self.claro = PatternFill('solid', fgColor='F2F2F2')
        self.centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.izquierda = Alignment(horizontal='left', vertical='center')
        self.vertical = Alignment(horizontal='center', vertical='center', wrap_text=True, text_rotation=90)


def _libro():
    from openpyxl import Workbook
    libro = Workbook()
    libro.remove(libro.active)
    libro.calculation.fullCalcOnLoad = True  # Excel calcula las fórmulas al abrir.
    return libro, _Estilo()


def _guardar(libro):
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _hoja(libro, nombre, horizontal=True):
    hoja = libro.create_sheet(re.sub(r'[\[\]:*?/\\]', '', nombre)[:31])
    hoja.sheet_view.showGridLines = False
    hoja.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    hoja.page_setup.paperSize = hoja.PAPERSIZE_LETTER
    hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_options.horizontalCentered = True
    return hoja


def _membrete(hoja, e, escuela, ciclo, titulo, columnas):
    """Renglones 1 a 3: escuela, datos de la escuela y título. Devuelve el renglón donde sigue."""
    from openpyxl.utils import get_column_letter
    ultima = get_column_letter(max(columnas, 2))
    datos = [f'CCT: {escuela.get("cct") or "__________"}', f'TURNO: {(escuela.get("turno_texto") or "__________").upper()}']
    if escuela.get('zona_escolar'):
        datos.append(f'ZONA ESCOLAR: {escuela["zona_escolar"]}')
    if escuela.get('municipio'):
        datos.append(f'MUNICIPIO: {escuela["municipio"].upper()}')
    datos.append(f'CICLO ESCOLAR: {ciclo or "__________"}')
    for renglon, (texto, fuente) in enumerate(((f'ESCUELA PRIMARIA {(escuela.get("nombre") or "__________").upper()}'.replace(
            'ESCUELA PRIMARIA ESCUELA PRIMARIA', 'ESCUELA PRIMARIA'), e.membrete), ('   ·   '.join(datos), e.normal), (titulo.upper(), e.titulo)), 1):
        hoja.merge_cells(f'A{renglon}:{ultima}{renglon}')
        celda = hoja.cell(renglon, 1, texto)
        celda.font, celda.alignment = fuente, e.centro
    hoja.row_dimensions[3].height = 20
    return 5


def _encabezados(hoja, e, renglon, textos, alto=30, desde=1):
    for columna, texto in enumerate(textos, desde):
        celda = hoja.cell(renglon, columna, texto)
        celda.font, celda.fill, celda.border, celda.alignment = e.negrita, e.gris, e.borde, e.centro
    hoja.row_dimensions[renglon].height = alto


def _bloque(hoja, e, renglon, desde, hasta, texto):
    from openpyxl.utils import get_column_letter
    hoja.merge_cells(f'{get_column_letter(desde)}{renglon}:{get_column_letter(hasta)}{renglon}')
    for columna in range(desde, hasta + 1):
        hoja.cell(renglon, columna).border = e.borde
    celda = hoja.cell(renglon, desde, texto)
    celda.font, celda.fill, celda.alignment = e.negrita, e.gris, e.centro


def _celda(hoja, e, renglon, columna, valor, formato=None, centro=False, negrita=False):
    celda = hoja.cell(renglon, columna, valor)
    celda.font, celda.border = e.negrita if negrita else e.normal, e.borde
    celda.alignment = e.centro if centro else e.izquierda
    if formato:
        celda.number_format = formato
    return celda


def _anchos(hoja, anchos):
    from openpyxl.utils import get_column_letter
    for columna, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho


def _nota(hoja, e, renglon, texto):
    hoja.cell(renglon, 1, texto).font = e.nota


AVISO = 'Contiene datos personales de los alumnos: guárdelo en la escuela y no lo envíe por correo ni lo suba a internet.'


# ---------------------------------------------------------------- Metadato de 71 columnas

def metadato(registros, escuela, ciclo, hoy):
    from openpyxl.utils import get_column_letter
    libro, e = _libro()
    hoja = _hoja(libro, 'METADATO')
    total = len(METADATO_ALUMNO) + len(METADATO_TUTOR)
    renglon = _membrete(hoja, e, escuela, ciclo, 'Base de datos Metadato', 12)
    corte = _corte(ciclo, hoy)
    hoja.cell(4, 1, 'Edad cumplida al:').font = e.negrita
    _celda(hoja, e, 4, 3, corte, 'DD/MM/YYYY', centro=True, negrita=True)
    _bloque(hoja, e, renglon, 1, len(METADATO_ALUMNO), 'DATOS DEL ALUMNO(A)')
    _bloque(hoja, e, renglon, len(METADATO_ALUMNO) + 1, total, 'DATOS DE LA MADRE, PADRE O TUTOR(A)')
    _encabezados(hoja, e, renglon + 1, METADATO_ALUMNO + METADATO_TUTOR, alto=48)
    col = {texto: n for n, texto in enumerate(METADATO_ALUMNO, 1)}
    tut = {texto: n for n, texto in enumerate(METADATO_TUTOR, len(METADATO_ALUMNO) + 1)}
    for numero, r in enumerate(_ordenados(registros), 1):
        fila = renglon + 1 + numero
        nacimiento = indicadores.fecha(r.get('fecha_nacimiento')) or _fecha_curp(r.get('curp'))
        domicilio = _domicilio(r)
        tipos = r.get('discapacidades') or []
        con_salud = r.get('con_salud')
        tutor_curp = (r.get('tutor_curp') or '').upper()
        primero, segundo, nombres = partir_nombre(r.get('tutor') or '', tutor_curp or None)
        valores = {
            col['NP']: numero, col['GRADO Y GRUPO']: f'{r["grado"]}° {r.get("letra") or ""}'.strip() if r.get('grado') else '',
            col['APELLIDO PATERNO']: (r.get('primer_apellido') or '').upper(), col['APELLIDO MATERNO']: (r.get('segundo_apellido') or '').upper(),
            col['NOMBRE (S)']: (r.get('nombres') or '').upper(), col['GÉNERO']: _genero(r.get('sexo') or r.get('sexo_curp')),
            col['CURP']: (r.get('curp') or '').upper(), col['FECHA DE NACIMIENTO']: nacimiento,
            col['LUGAR DE NACIMIENTO']: _entidad(r.get('curp')).upper(), col['NACIONALIDAD']: _nacionalidad(r.get('curp')),
            col['ENFERMEDAD CRÓNICA O PADECIMIENTO']: _ninguna(r.get('padecimientos') or [], con_salud),
            col['ALERGIAS']: _ninguna(r.get('alergias') or [], con_salud),
            col['CALLE']: domicilio['calle'], col['NO. EXTERIOR']: domicilio['numero'], col['COLONIA']: domicilio['colonia'],
            col['C. P.']: domicilio['cp'], col['MUNICIPIO']: domicilio['municipio'],
            col['PRESENTA ALGUNA CONDICIÓN']: _ninguna([t for t in tipos if t in CONDICIONES], con_salud),
            col['PRESENTA ALGUNA DISCAPACIDAD']: _ninguna([t for t in tipos if t not in CONDICIONES], con_salud),
            col['HABLA ALGUNA LENGUA INDÍGENA']: {'sí': f'SÍ ({(r.get("lengua") or "").upper()})'.replace(' ()', ''), 'no': 'NO'}.get(r.get('lengua_indigena'), ''),
            col['SITUACIÓN DEL ALUMNO']: _situacion(r),
            tut['APELLIDO PATERNO']: primero, tut['APELLIDO MATERNO']: segundo, tut['NOMBRE (S)']: nombres,
            tut['PARENTESCO']: (r.get('parentesco') or '').upper(), tut['CURP']: tutor_curp, tut['FECHA DE NACIMIENTO']: _fecha_curp(tutor_curp),
            tut['GÉNERO']: _genero(tutor_curp[10:11]), tut['LUGAR DE NACIMIENTO']: _entidad(tutor_curp).upper(),
            tut['NACIONALIDAD']: _nacionalidad(tutor_curp), tut['TELÉFONO CELULAR DEL TUTOR']: r.get('telefono') or '',
            tut['CORREO ELECTRÓNICO']: r.get('correo') or '', tut['NOMBRE COMPLETO DE CONTACTO EN CASO DE EMERGENCIA (Y PARENTESCO)']: (r.get('emergencia') or '').upper(),
        }
        for columna in range(1, total + 1):
            valor = valores.get(columna)
            _celda(hoja, e, fila, columna, valor if valor not in (None, '') else None, 'DD/MM/YYYY' if isinstance(valor, datetime.date) else None,
                   centro=columna in (col['NP'], col['GRADO Y GRUPO'], col['GÉNERO']))
        # La edad se calcula sola con la fecha de nacimiento y la fecha de corte de arriba.
        for edad, nace in ((col['EDAD'], col['FECHA DE NACIMIENTO']), (tut['EDAD'], tut['FECHA DE NACIMIENTO'])):
            letra = get_column_letter(nace)
            _celda(hoja, e, fila, edad, f'=IF({letra}{fila}="","",DATEDIF({letra}{fila},$C$4,"Y"))', centro=True)
    anchos = [5, 9, 16, 16, 18, 9, 21, 12, 6] + [14] * (total - 9)
    anchos[col['CURP'] - 1] = anchos[tut['CURP'] - 1] = 21
    _anchos(hoja, anchos)
    hoja.freeze_panes = hoja.cell(renglon + 2, col['CURP'] + 1)
    hoja.auto_filter.ref = f'A{renglon + 1}:{get_column_letter(total)}{renglon + 1 + len(registros)}'
    _nota(hoja, e, renglon + 3 + len(registros), 'Lo que está en blanco no lo tiene SIDP: complételo antes de entregar. ' + AVISO)
    return _guardar(libro)


# ---------------------------------------------------------------- Base del 911 que se calcula sola

def base_911(registros, escuela, ciclo, hoy):
    from openpyxl.utils import get_column_letter
    libro, e = _libro()
    base = _hoja(libro, 'BASE')
    renglon = _membrete(base, e, escuela, ciclo, 'Estadística 911: base de alumnos', 12)
    base.cell(4, 1, 'Edad cumplida al:').font = e.negrita
    _celda(base, e, 4, 3, _corte(ciclo, hoy), 'DD/MM/YYYY', centro=True, negrita=True)
    base.cell(4, 5, 'Escriba o pegue la CURP: el sexo, la fecha de nacimiento, la edad y los cuadros de la hoja «911» se calculan solos.').font = e.nota
    encabezados = ('NP', 'CURP', 'PRIMER APELLIDO', 'SEGUNDO APELLIDO', 'NOMBRE(S)', 'GRADO', 'GRUPO', 'ESTATUS', 'SEXO',
                   'FECHA DE NACIMIENTO', 'EDAD', 'RANGO 911')
    _encabezados(base, e, renglon, encabezados)
    primera = renglon + 1
    alumnos = _ordenados(registros)
    for numero in range(len(alumnos) + FILAS_LIBRES):
        fila = primera + numero
        r = alumnos[numero] if numero < len(alumnos) else {}
        fijos = (numero + 1 if r else None, (r.get('curp') or '').upper() or None, (r.get('primer_apellido') or '').upper() or None,
                 (r.get('segundo_apellido') or '').upper() or None, (r.get('nombres') or '').upper() or None, r.get('grado'),
                 (r.get('letra') or '') or None, ('REPETIDOR' if r.get('repite') == 'sí' else 'INSCRITO') if r else None)
        for columna, valor in enumerate(fijos, 1):
            _celda(base, e, fila, columna, valor, centro=columna in (1, 6, 7, 8))
        b = f'B{fila}'
        formulas = (f'=IF(LEN({b})<18,"",MID({b},11,1))',
                    f'=IF(LEN({b})<18,"",DATE(IF(ISNUMBER(VALUE(MID({b},17,1))),1900,2000)+VALUE(MID({b},5,2)),VALUE(MID({b},7,2)),VALUE(MID({b},9,2))))',
                    f'=IF(J{fila}="","",DATEDIF(J{fila},$C$4,"Y"))',
                    f'=IF(K{fila}="","",IF(K{fila}<6,"Menos de 6",IF(K{fila}>=15,"15 y más",TEXT(K{fila},"0"))))')
        for columna, formula in enumerate(formulas, 9):
            _celda(base, e, fila, columna, formula, 'DD/MM/YYYY' if columna == 10 else None, centro=True)
    _anchos(base, (5, 21, 16, 16, 20, 7, 7, 11, 6, 12, 6, 11))
    base.freeze_panes = base.cell(primera, 3)
    from openpyxl.worksheet.datavalidation import DataValidation
    validar = DataValidation(type='list', formula1='"INSCRITO,REPETIDOR"', allow_blank=True)
    base.add_data_validation(validar)
    validar.add(f'H{primera}:H{FINAL}')

    cuadro = _hoja(libro, '911')
    edades = indicadores.EDADES_911
    columnas = 3 + len(edades) + 1
    renglon = _membrete(cuadro, e, escuela, ciclo, 'Estadística 911 de inicio de curso: alumnos por grado, sexo y edad', columnas)
    _nota(cuadro, e, 4, 'Se calcula con la hoja «BASE». Copie estos números en el cuestionario 911.3 de la plataforma de la SEP.')
    _encabezados(cuadro, e, renglon, ('GRADO', 'SEXO', 'ESTATUS', *edades, 'TOTAL'))
    rango = lambda letra: f'BASE!${letra}${primera}:${letra}${FINAL}'  # noqa: E731
    fila = renglon + 1
    subtotales = []
    for g in range(1, 7):
        for sexo, texto in (('H', 'HOMBRES'), ('M', 'MUJERES')):
            for estatus in ('INSCRITO', 'REPETIDOR'):
                _celda(cuadro, e, fila, 1, f'{g}°', centro=True)
                _celda(cuadro, e, fila, 2, texto, centro=True)
                _celda(cuadro, e, fila, 3, 'NUEVO INGRESO' if estatus == 'INSCRITO' else 'REPETIDORES', centro=True)
                for n, edad in enumerate(edades, 4):
                    _celda(cuadro, e, fila, n, f'=COUNTIFS({rango("F")},{g},{rango("I")},"{sexo}",{rango("H")},"{estatus}",{rango("L")},"{edad}")', centro=True)
                _celda(cuadro, e, fila, columnas, f'=SUM(D{fila}:{get_column_letter(columnas - 1)}{fila})', centro=True, negrita=True)
                fila += 1
        _celda(cuadro, e, fila, 1, f'{g}°', centro=True, negrita=True)
        _celda(cuadro, e, fila, 2, 'TOTAL', centro=True, negrita=True)
        _celda(cuadro, e, fila, 3, '', centro=True)
        for n in range(4, columnas + 1):
            letra = get_column_letter(n)
            _celda(cuadro, e, fila, n, f'=SUM({letra}{fila - 4}:{letra}{fila - 1})', centro=True, negrita=True).fill = e.claro
        subtotales.append(fila)
        fila += 1
    _celda(cuadro, e, fila, 1, 'ESCUELA', centro=True, negrita=True)
    _celda(cuadro, e, fila, 2, 'TOTAL', centro=True, negrita=True)
    _celda(cuadro, e, fila, 3, '', centro=True)
    for n in range(4, columnas + 1):
        letra = get_column_letter(n)
        _celda(cuadro, e, fila, n, '=' + '+'.join(f'{letra}{s}' for s in subtotales), centro=True, negrita=True).fill = e.gris
    fila += 2
    cuadro.cell(fila, 1, 'GRUPOS Y ALUMNOS POR GRADO').font = e.negrita
    fila += 1
    _encabezados(cuadro, e, fila, ('GRADO', 'GRUPOS', 'HOMBRES', 'MUJERES', 'TOTAL'))
    grupos = _grupos(registros)
    for g in range(1, 7):
        fila += 1
        _celda(cuadro, e, fila, 1, f'{g}°', centro=True)
        _celda(cuadro, e, fila, 2, sum(1 for (grado, _) in grupos if grado == g), centro=True)
        _celda(cuadro, e, fila, 3, f'=COUNTIFS({rango("F")},{g},{rango("I")},"H")', centro=True)
        _celda(cuadro, e, fila, 4, f'=COUNTIFS({rango("F")},{g},{rango("I")},"M")', centro=True)
        _celda(cuadro, e, fila, 5, f'=C{fila}+D{fila}', centro=True, negrita=True)
    _anchos(cuadro, (8, 10, 15) + (8,) * len(edades) + (8,))
    cuadro.freeze_panes = cuadro.cell(renglon + 1, 4)
    libro.move_sheet('911', offset=-1)
    libro.active = 0
    return _guardar(libro)


# ---------------------------------------------------------------- Censo de condición y discapacidad

def condicion_discapacidad(registros, escuela, ciclo, hoy):
    from openpyxl.utils import get_column_letter
    libro, e = _libro()
    por_grado = {g: [r for r in _ordenados(registros) if r.get('grado') == g] for g in range(1, 7)}
    columnas = 2 + len(CONDICION) + len(DISCAPACIDAD) + 1
    for g, alumnos in por_grado.items():
        hoja = _hoja(libro, GRADOS[g - 1])
        renglon = _membrete(hoja, e, escuela, ciclo, f'Alumnos con condición o discapacidad: {GRADOS[g - 1].lower()} grado', columnas)
        _encabezados(hoja, e, renglon, ('NO.', 'NOMBRE DEL ALUMNO'))
        hoja.merge_cells(start_row=renglon, start_column=1, end_row=renglon + 1, end_column=1)
        hoja.merge_cells(start_row=renglon, start_column=2, end_row=renglon + 1, end_column=2)
        _bloque(hoja, e, renglon, 3, 2 + len(CONDICION), 'CONDICIÓN')
        _bloque(hoja, e, renglon, 3 + len(CONDICION), 2 + len(CONDICION) + len(DISCAPACIDAD), 'DISCAPACIDAD')
        _encabezados(hoja, e, renglon, ('OBSERVACIONES',), desde=columnas)
        hoja.merge_cells(start_row=renglon, start_column=columnas, end_row=renglon + 1, end_column=columnas)
        _encabezados(hoja, e, renglon + 1, [t for t, _ in CONDICION + DISCAPACIDAD], alto=60, desde=3)
        for columna in range(3, columnas):
            hoja.cell(renglon + 1, columna).alignment = e.vertical
        fila = renglon + 1
        for numero, r in enumerate(alumnos, 1):
            fila += 1
            tipos = r.get('discapacidades') or []
            _celda(hoja, e, fila, 1, numero, centro=True)
            _celda(hoja, e, fila, 2, ' '.join(filter(None, (r.get('primer_apellido'), r.get('segundo_apellido'), r.get('nombres')))).upper())
            for columna, (_, tipo) in enumerate(CONDICION + DISCAPACIDAD, 3):
                _celda(hoja, e, fila, columna, 'X' if tipo and tipo in tipos else None, centro=True)
            otras = [t for t in tipos if t not in {tipo for _, tipo in CONDICION + DISCAPACIDAD}]
            _celda(hoja, e, fila, columnas, ', '.join(otras).upper() or None)
        fila += 1
        _celda(hoja, e, fila, 2, 'TOTAL', negrita=True).fill = e.claro
        _celda(hoja, e, fila, 1, None)
        for columna in range(3, columnas):
            letra = get_column_letter(columna)
            _celda(hoja, e, fila, columna, f'=COUNTIF({letra}{renglon + 2}:{letra}{fila - 1},"X")', centro=True, negrita=True).fill = e.claro
        _celda(hoja, e, fila, columnas, None)
        _anchos(hoja, (5, 38) + (5,) * (len(CONDICION) + len(DISCAPACIDAD)) + (24,))
        hoja.freeze_panes = hoja.cell(renglon + 2, 3)
        _nota(hoja, e, fila + 2, 'Sale del certificado médico y de lo que capturó la escuela. Marque con X lo que falte. ' + AVISO)
    return _guardar(libro)


# ---------------------------------------------------------------- Plantillas para el maestro, con los alumnos ya puestos

def _alumno(hoja, e, fila, numero, r):
    _celda(hoja, e, fila, 1, numero, centro=True)
    _celda(hoja, e, fila, 2, (r.get('primer_apellido') or '').upper())
    _celda(hoja, e, fila, 3, (r.get('segundo_apellido') or '').upper())
    _celda(hoja, e, fila, 4, (r.get('nombres') or '').upper())


def calificaciones(registros, escuela, ciclo, hoy):
    """Concentrado trimestral por grupo: el maestro escribe las calificaciones y el promedio y la dispersión salen solos."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    libro, e = _libro()
    grupos = _grupos(registros)
    concentrado = _hoja(libro, 'CONCENTRADO')
    resumen = []
    encabezados = ('NO.', 'APELLIDO PATERNO', 'APELLIDO MATERNO', 'NOMBRE(S)', *CAMPOS_NEM, 'EDUCACIÓN FÍSICA', 'INGLÉS', 'PROMEDIO', 'FALTAS')
    for (grado, letra), alumnos in grupos.items():
        nombre = f'{grado}° {letra}'.strip()
        hoja = _hoja(libro, nombre)
        renglon = _membrete(hoja, e, escuela, ciclo, f'Concentrado de calificaciones del trimestre: {nombre}', len(encabezados))
        hoja.cell(4, 1, 'Trimestre:').font = e.negrita
        _celda(hoja, e, 4, 2, None, centro=True)
        hoja.cell(4, 4, 'Escriba las calificaciones (5 a 10): el promedio y la dispersión se calculan solos.').font = e.nota
        _encabezados(hoja, e, renglon, encabezados, alto=42)
        primera = renglon + 1
        for numero, r in enumerate(alumnos, 1):
            fila = primera + numero - 1
            _alumno(hoja, e, fila, numero, r)
            for columna in range(5, 11):
                _celda(hoja, e, fila, columna, None, '0', centro=True)
            _celda(hoja, e, fila, 11, f'=IF(COUNT(E{fila}:H{fila})<4,"",ROUND(AVERAGE(E{fila}:H{fila}),1))', '0.0', centro=True, negrita=True)
            _celda(hoja, e, fila, 12, None, centro=True)
        ultima = primera + len(alumnos) - 1
        validar = DataValidation(type='decimal', operator='between', formula1='5', formula2='10', allow_blank=True,
                                 error='La calificación va de 5 a 10.', errorTitle='Calificación')
        hoja.add_data_validation(validar)
        validar.add(f'E{primera}:J{ultima}')
        fila = ultima + 2
        hoja.cell(fila, 2, 'DISPERSIÓN POR NIVEL DE DESEMPEÑO').font = e.negrita
        _encabezados(hoja, e, fila + 1, ('NIVEL', 'PROMEDIO', 'ALUMNOS', '%'), desde=2)
        promedios = f'$K${primera}:$K${ultima}'
        filas_nivel = []
        for n, (nivel, texto, desde, hasta) in enumerate(DISPERSION):
            f = fila + 2 + n
            _celda(hoja, e, f, 2, nivel)
            _celda(hoja, e, f, 3, texto, centro=True)
            _celda(hoja, e, f, 4, f'=COUNTIFS({promedios},">={desde}",{promedios},"<{hasta}")', centro=True)
            _celda(hoja, e, f, 5, f'=IF(SUM($D${fila + 2}:$D${fila + 1 + len(DISPERSION)})=0,"",D{f}/SUM($D${fila + 2}:$D${fila + 1 + len(DISPERSION)}))',
                   '0.0%', centro=True)
            filas_nivel.append(f)
        f = fila + 2 + len(DISPERSION)
        _celda(hoja, e, f, 2, 'PROMEDIO DEL GRUPO', negrita=True).fill = e.claro
        _celda(hoja, e, f, 3, None)
        _celda(hoja, e, f, 4, f'=IF(COUNT({promedios})=0,"",ROUND(AVERAGE({promedios}),1))', '0.0', centro=True, negrita=True).fill = e.claro
        _celda(hoja, e, f, 5, None)
        _anchos(hoja, (5, 18, 18, 22, 13, 13, 13, 13, 11, 9, 10, 8))
        hoja.freeze_panes = hoja.cell(primera, 5)
        resumen.append((nombre, hoja.title, filas_nivel, f))
    renglon = _membrete(concentrado, e, escuela, ciclo, 'Concentrado escolar de dispersión por nivel de desempeño', 2 + 2 * len(DISPERSION) + 1)
    _nota(concentrado, e, 4, 'Se llena solo con las hojas de cada grupo.')
    _encabezados(concentrado, e, renglon, ('GRUPO', 'ALUMNOS', *[f'{nivel} ({texto})' for nivel, texto, *_ in DISPERSION], 'PROMEDIO'), alto=42)
    for n, (nombre, titulo, filas_nivel, fila_promedio) in enumerate(resumen, 1):
        f = renglon + n
        hoja_ref = f"'{titulo}'"
        _celda(concentrado, e, f, 1, nombre, centro=True, negrita=True)
        _celda(concentrado, e, f, 2, '=' + '+'.join(f'{hoja_ref}!D{x}' for x in filas_nivel), centro=True)
        for m, x in enumerate(filas_nivel, 3):
            _celda(concentrado, e, f, m, f'={hoja_ref}!D{x}', centro=True)
        _celda(concentrado, e, f, 3 + len(filas_nivel), f'={hoja_ref}!D{fila_promedio}', '0.0', centro=True, negrita=True)
    _anchos(concentrado, (10, 10) + (16,) * len(DISPERSION) + (11,))
    libro.active = 0
    return _guardar(libro)


def boleta(registros, escuela, ciclo, hoy):
    """Base de datos de la boleta con observaciones: los 4 campos formativos de los 3 periodos, con los alumnos puestos."""
    from openpyxl.worksheet.datavalidation import DataValidation
    libro, e = _libro()
    hoja = _hoja(libro, 'BASE BOLETA')
    encabezados = ['GRADO', 'GRUPO', 'FOLIO', 'CURP', 'PRIMER APELLIDO', 'SEGUNDO APELLIDO', 'NOMBRE(S)']
    for periodo in ('P1', 'P2', 'P3'):
        encabezados += [f'{c} {periodo}' for c in CAMPOS_NEM] + [f'OBS {c} {periodo}' for c in CAMPOS_NEM]
    renglon = _membrete(hoja, e, escuela, ciclo, 'Base de datos de la boleta y observaciones', 12)
    _encabezados(hoja, e, renglon, encabezados, alto=60)
    alumnos = _ordenados(registros)
    for n, r in enumerate(alumnos, 1):
        fila = renglon + n
        for columna, valor in enumerate((r.get('grado'), r.get('letra') or None, None, (r.get('curp') or '').upper(),
                                         (r.get('primer_apellido') or '').upper(), (r.get('segundo_apellido') or '').upper(),
                                         (r.get('nombres') or '').upper()), 1):
            _celda(hoja, e, fila, columna, valor, centro=columna in (1, 2))
        for columna in range(8, len(encabezados) + 1):
            _celda(hoja, e, fila, columna, None, centro=((columna - 8) % 8) < 4)
    if alumnos:
        validar = DataValidation(type='whole', operator='between', formula1='5', formula2='10', allow_blank=True,
                                 error='La calificación va de 5 a 10.', errorTitle='Calificación')
        hoja.add_data_validation(validar)
        from openpyxl.utils import get_column_letter
        for inicio in (8, 16, 24):
            validar.add(f'{get_column_letter(inicio)}{renglon + 1}:{get_column_letter(inicio + 3)}{renglon + len(alumnos)}')
    _anchos(hoja, (7, 7, 10, 21, 16, 16, 20) + tuple((10,) * 4 + (28,) * 4) * 3)
    hoja.freeze_panes = hoja.cell(renglon + 1, 8)
    return _guardar(libro)


def asistencias(registros, escuela, ciclo, hoy):
    """Concentrado de asistencias y faltas por mes; el total de cada mes y del año se suma solo."""
    from openpyxl.utils import get_column_letter
    libro, e = _libro()
    for (grado, letra), alumnos in _grupos(registros).items():
        nombre = f'{grado}° {letra}'.strip()
        hoja = _hoja(libro, nombre)
        columnas = 4 + 4 * len(MESES_CICLO) + 4
        renglon = _membrete(hoja, e, escuela, ciclo, f'Concentrado de asistencias y faltas: {nombre}', 30)
        _nota(hoja, e, 4, 'ASIS. = asistencias · F. J. = faltas justificadas · F = faltas sin justificar · T = días del mes (se suma solo).')
        _encabezados(hoja, e, renglon, ('NO.', 'APELLIDO PATERNO', 'APELLIDO MATERNO', 'NOMBRE(S)'))
        for columna in range(1, 5):
            hoja.merge_cells(start_row=renglon, start_column=columna, end_row=renglon + 1, end_column=columna)
        for n, mes in enumerate(MESES_CICLO + ('TOTAL DEL CICLO',)):
            desde = 5 + 4 * n
            _bloque(hoja, e, renglon, desde, desde + 3, mes)
            _encabezados(hoja, e, renglon + 1, ('ASIS.', 'F. J.', 'F', 'T'), alto=18, desde=desde)
        for numero, r in enumerate(alumnos, 1):
            fila = renglon + 1 + numero
            _alumno(hoja, e, fila, numero, r)
            for n in range(len(MESES_CICLO)):
                desde = 5 + 4 * n
                for k in range(3):
                    _celda(hoja, e, fila, desde + k, None, centro=True)
                a, c = get_column_letter(desde), get_column_letter(desde + 2)
                _celda(hoja, e, fila, desde + 3, f'=IF(COUNT({a}{fila}:{c}{fila})=0,"",SUM({a}{fila}:{c}{fila}))', centro=True).fill = e.claro
            desde = 5 + 4 * len(MESES_CICLO)
            for k in range(4):
                partes = ','.join(f'{get_column_letter(5 + 4 * n + k)}{fila}' for n in range(len(MESES_CICLO)))
                _celda(hoja, e, fila, desde + k, f'=SUM({partes})', centro=True, negrita=True).fill = e.claro
        _anchos(hoja, (5, 16, 16, 20) + (5,) * (4 * len(MESES_CICLO)) + (7,) * 4)
        hoja.freeze_panes = hoja.cell(renglon + 2, 5)
    return _guardar(libro)


GENERADORES = {'metadato': metadato, 'base_911': base_911, 'condicion': condicion_discapacidad, 'calificaciones': calificaciones,
               'boleta': boleta, 'asistencias': asistencias}
