"""El Excel por grado: un resumen que compara los grados, todos los alumnos y una hoja por grado con gráficas."""
import datetime
import io
import re
import zipfile

import openpyxl

from pruebas.test_entorno_pruebas import descargar, entorno, pedir  # noqa: F401 (entorno es la prueba con el panel)
from validacion import libro_grados

HOY = datetime.date(2026, 10, 16)


def salud(alergias=(), estado='vigente', municipio=None, boletas=None, repite=None, emergencia=None):
    datos = [{'campo': 'alergia', 'valor': a, 'estado': estado} for a in alergias]
    escolar = [{'campo': f'promedio_{g}', 'valor': v, 'estado': 'vigente'} for g, v in (boletas or {}).items()]
    escolar += [{'campo': 'repite_grado', 'valor': repite, 'estado': 'vigente'}] if repite else []
    domicilio = [{'campo': 'municipio', 'valor': municipio, 'estado': 'vigente'}] if municipio else []
    contacto = [{'campo': 'nombre_emergencia', 'valor': emergencia, 'estado': 'vigente'}] if emergencia else []
    return {'datos': {'salud': datos, 'escolar': escolar, 'domicilio': domicilio, 'contacto': contacto}}


def fila(grado, letra, apellido, nombre, curp='', sexo='', nacimiento=''):
    return [grado, letra, apellido, 'RUIZ', nombre, curp, 'No', sexo, nacimiento, '', '', '', 0, 0]


def test_cada_alumno_junta_la_lista_de_aqui_con_lo_que_saco_el_escaneo():
    # Sin fecha ni sexo en la lista: salen de la CURP (inventada; el carácter 17 es letra porque nació en el 2000 o después).
    nino = libro_grados.alumno(fila(1, 'A', 'LOPEZ', 'ANA', 'LORA190305MMCPZNA1'), 2026,
                               salud(['POLVO', 'POLVO', 'NUECES'], estado='por_confirmar', emergencia='TIA ROSA'), HOY)
    assert (nino['fecha_nacimiento'], nino['sexo'], nino['edad'], nino['situacion']) == (datetime.date(2019, 3, 5), 'Mujer', 7, 'En su edad')
    assert nino['alergias'] == 'POLVO, NUECES' and nino['por_confirmar'] == 3 and nino['nacido'] == 'Estado de México'
    assert (nino['emergencia'], nino['repite']) == ('TIA ROSA', None)
    assert not {'peso', 'talla', 'imc', 'tipo_sangre', 'nutricion'} & set(nino)  # Solo lo necesario.
    grande = libro_grados.alumno(fila(4, 'A', 'PEREZ', 'LUIS', nacimiento='2015-01-10'), 2026,
                                 salud(['Negadas'], municipio='CHALCO, EDO. MEX.', boletas={1: '9', 2: '8,5', 3: '6.8'}, repite='si'), HOY)
    assert (grande['edad'], grande['situacion'], grande['alergias'], grande['municipio']) == (11, 'Extraedad', None, 'Chalco')
    assert (grande['promedio_1'], grande['promedio_3'], grande['promedio'], grande['repite']) == (9.0, 6.8, 6.8, 'Sí')  # La última es la de 3.°


def test_el_libro_trae_salud_domicilio_boletas_911_y_una_hoja_por_grado_sin_documentos():
    alumnos = [libro_grados.alumno(fila(1, 'A', 'LOPEZ', 'ANA', sexo='M', nacimiento='2019-03-05'), 2026,
                                   salud(['POLVO'], municipio='Chalco', repite='no'), HOY),
               libro_grados.alumno(fila(1, 'B', 'GARCIA', 'JUAN', sexo='H', nacimiento='2019-07-01'), 2026,
                                   salud(['NINGUNA'], municipio='Valle de Chalco Solidaridad', repite='sí'), HOY),
               libro_grados.alumno(fila(3, 'A', 'CRUZ', 'SOFIA', 'CURS170202MDFRZFA1', 'M', '2017-02-02'), 2026,
                                   salud(boletas={1: '9.5', 2: '6.5'}), HOY)]
    contenido = libro_grados.libro(alumnos, '2026-2027', HOY)
    libro = openpyxl.load_workbook(io.BytesIO(contenido))
    assert libro.sheetnames == ['Resumen', 'Salud', 'Salud por grado', 'Dónde viven', 'Boletas', '911', 'Todos', '1°', '2°', '3°', '4°', '5°', '6°']
    resumen = libro['Resumen']
    encabezado = [c.value for c in resumen[4]]
    assert not any('ocumento' in str(e) for e in encabezado)  # Nada de documentos: solo de los alumnos.
    primero = dict(zip(encabezado, [c.value for c in resumen[5]]))
    escuela = dict(zip(encabezado, [c.value for c in resumen[11]]))
    assert (primero['Grado'], primero['Grupos'], primero['Alumnos'], primero['Hombres'], primero['Mujeres']) == ('1°', 2, 2, 1, 1)
    assert (primero['Con alergias'], primero['Repiten el grado']) == (1, 1)
    assert not any('eso' in str(e) or 'alla' in str(e) for e in encabezado)  # Sin peso ni talla.
    assert (escuela['Grado'], escuela['Alumnos'], escuela['Nacidos en otra entidad o país']) == ('Escuela', 3, 1)
    # Salud: lo importante en rojo.
    salud_ = libro['Salud']
    columnas = [c.value for c in salud_[4]]
    ana = [c for c in salud_[5]]
    assert ana[3].value == 'ANA' and ana[columnas.index('Alergias')].value == 'POLVO'
    assert ana[columnas.index('Alergias')].fill.fgColor.rgb.endswith('F8CBAD')
    # Dónde viven, boletas y 911.
    donde = libro['Dónde viven']
    assert [donde.cell(6, c).value for c in (1, 2)] in (['Chalco', 1], ['Valle de Chalco Solidaridad', 1])
    boletas = libro['Boletas']
    assert [c.value for c in boletas[5]][:7] == ['3° A', 'CRUZ', 'RUIZ', 'SOFIA', 9.5, 6.5, None]
    hoja_911 = libro['911']
    total = [c.value for c in hoja_911[5 + 6 * 3]]  # 6 grados x (hombres, mujeres, subtotal) y luego el total.
    assert total[0] == 'Total' and total[-1] == 3 and sum(total[2:-1]) == 3
    assert [c.value for c in libro['1°'][4]][:4] == ['1° A', 'LOPEZ', 'RUIZ', 'ANA']
    assert libro['2°']['A2'].value == 'No hay alumnos en este grado.'
    with zipfile.ZipFile(io.BytesIO(contenido)) as archivo:
        graficas = [n for n in archivo.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml', n)]
        # Resumen 5; Salud por grado 4 (por grado, alergias, padecimientos, medicamentos); Dónde viven 3; Boletas 2;
        # 1.° 2 (sin boleta) y 3.° 3 (con boleta).
        assert len(graficas) == 5 + 4 + 3 + 2 + 2 + 3
        assert all('<delete val="0"/>' in archivo.read(n).decode() for n in graficas if 'barChart' in archivo.read(n).decode())


def test_el_panel_descarga_el_excel_por_grado_aunque_la_pc_4_no_tenga_base(entorno):  # noqa: F811
    app, url = entorno
    token = {'X-SIDP-Token': app.token}
    pedir(url, '/api/listas/guardar', {'grado': 2, 'letra': 'A', 'filas': [
        {'nombres': 'LUIS', 'primer_apellido': 'HERNANDEZ', 'segundo_apellido': 'RAMOS', 'curp': 'HERL180314HMCRMSA9'}]}, token)
    cabeceras, contenido = descargar(url, '/api/exportar?que=por_grado&formato=xlsx', app.token)
    assert 'sidp_por_grado_' in cabeceras['Content-Disposition']
    libro = openpyxl.load_workbook(io.BytesIO(contenido))
    assert libro['2°']['D4'].value == 'LUIS' and libro['2°']['G4'].value == 'Hombre'  # El sexo sale de la CURP.
