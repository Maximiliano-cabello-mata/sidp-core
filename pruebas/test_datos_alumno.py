"""Sacar los datos de cada campo de lo escaneado y guardarlos en el expediente del alumno (PC 4).

Las primeras pruebas no necesitan nada. Las de la base necesitan el PostgreSQL de prueba (SIDP_PRUEBA_BASE).
El modelo que ve se reemplaza por respuestas fijas: aquí se prueba qué se hace con lo que lee, no el modelo.
Todo es inventado.
"""
import base64
import os
import urllib.request

import pytest

from pruebas.test_entorno_pruebas import HOY, escanear, iniciar, pedir
from pruebas.test_lista_oficial import pdf_de_hojas
from validacion import aplicacion, campos, cerebro_red, comparador, taller, vision
from validacion.ejemplo import generar

pytest.importorskip('pypdf')


# ---------------------------------------------------------------- Reglas de cada tipo de campo (sin base)

def test_cada_tipo_de_campo_se_limpia_a_un_valor_util():
    assert campos.fecha('14/03/2018') == campos.fecha('14 de marzo del 2018') == campos.fecha('2018-03-14') == '2018-03-14'
    assert campos.fecha('14-III-2018') == '2018-03-14' and campos.fecha('31/02/2018') is None and campos.fecha('14-XIV-2018') is None
    assert campos.fecha('02 SEP 26', HOY) == '2026-09-02'
    assert campos.telefono('722 123 45 67') == campos.telefono('+52 722-123-4567') == campos.telefono('722 l23 4S67') == '7221234567'
    assert campos.telefono('7221 23') is None
    assert campos.correo('Mama.Luis @gmial.con') == 'mama.luis@gmail.com' and campos.correo('sin correo') is None
    assert campos.sangre('O positivo') == campos.sangre('0 +') == 'O+' and campos.sangre('AB RH negativo') == 'AB-'
    assert campos.si_no('Sí') == campos.si_no('clínicamente sano') == 'sí' and campos.si_no('No apto') == 'no'
    assert campos.si_no('Se desconoce') is None
    assert campos.lista('Penicilina y polvo; mariscos') == ['PENICILINA', 'POLVO', 'MARISCOS']
    assert campos.lista('Ninguna') == [] and campos.lista('negadas') == []
    texto, partes = campos.domicilio('C. Pino No. 5 Col. Centro, Toluca, C.P. 50000')
    assert texto.startswith('C. PINO') and partes == {'cp': '50000', 'colonia': 'CENTRO', 'calle': 'C. PINO', 'numero': '5'}
    assert campos.limpiar('nombre', 'rosa  hernández 2')[0] == 'ROSA HERNÁNDEZ'
    assert campos.limpiar('curp', 'gahj 2005 12hm crrn a')[0] is None  # Le falta un carácter.


def test_una_plantilla_nueva_propone_a_donde_va_cada_campo():
    telefono = campos.proponer('Teléfono de la mamá')
    assert (telefono['tipo'], telefono['categoria'], telefono['campo'], telefono['de']) == ('telefono', 'contacto', 'telefono', 'madre')
    assert campos.proponer('Nombre del padre')['categoria'] == 'padre'
    assert campos.proponer('Alergias')['categoria'] == 'salud' and campos.proponer('Alergias')['tipo'] == 'lista'
    emergencia = campos.proponer('Nombre de contacto de emergencia')
    assert (emergencia['categoria'], emergencia['campo']) == ('contacto', 'nombre_emergencia')
    assert campos.proponer('Talla de uniforme')['categoria'] == 'salud'  # Se propone; la persona lo cambia si no.
    revisados = campos.revisar_campos([{'etiqueta': 'Teléfono'}, {'etiqueta': 'Teléfono', 'categoria': 'nada', 'tipo': 'raro'}, {'etiqueta': ' '}])
    assert [c['clave'] for c in revisados] == ['telefono', 'telefono_1'] and revisados[1]['categoria'] == 'contacto'
    with pytest.raises(ValueError):
        campos.revisar_campos([{'etiqueta': f'Dato {i}'} for i in range(61)])
    # Lo que se le pide al modelo que ve: todos los campos, con su tipo.
    esquema = campos.esquema(revisados)
    assert esquema['required'] == ['telefono', 'telefono_1'] and 'Teléfono' in campos.pregunta('Ficha', revisados)


def test_en_una_hoja_impresa_el_valor_se_toma_junto_a_su_etiqueta():
    ficha = ['FICHA DE INSCRIPCION', 'Teléfono de la madre: 722 123 4567', 'Correo electrónico:', 'mama.luis@gmail.com',
             'Tipo de sangre: O+', 'Fecha de nacimiento 14/03/2018']
    lista = campos.revisar_campos([{'etiqueta': e} for e in ('Teléfono de la madre', 'Correo electrónico', 'Tipo de sangre',
                                                               'Fecha de nacimiento', 'Alergias')])
    leidos = campos.de_texto('\n'.join(ficha), lista, HOY)
    assert leidos == {'telefono_de_la_madre': '722 123 4567', 'correo_electronico': 'mama.luis@gmail.com', 'tipo_de_sangre': 'O+',
                      'fecha_de_nacimiento': '14/03/2018'}


def test_dos_lecturas_iguales_dan_confianza_alta_y_distintas_quedan_por_confirmar():
    lista = campos.FABRICA['acta_alumno']
    identidad = {'curp': 'GAHJ200512HMCRRNA1', 'nombres': 'JUAN', 'nombre': 'JUAN GARCIA HERNANDEZ'}
    procesos = {'curp': ('GAHJ200512HMCRRNA1', 'texto'), 'nombre_madre': ('ROSA HERNANDEZ DIAZ', 'texto'),
                'nombre_padre': ('PEDRO GARCIA RUIZ', 'texto')}
    vistos = {'nombre_madre': {'valor': 'Rosa Hernández Díaz', 'seguro': True}, 'nombre_padre': {'valor': 'PEDRO GARCIA RUIS', 'seguro': True},
              'lugar_nacimiento': {'valor': 'Toluca, Méx.', 'seguro': False}, 'curp_madre': {'valor': '', 'seguro': True}}
    salida = {r['clave']: r for r in taller.combinar(lista, procesos, vistos, identidad, HOY)}
    assert salida['curp']['confirma_lista']  # Ya está en la lista oficial: no se guarda otra vez.
    assert (salida['nombre_madre']['confianza'], salida['nombre_madre']['metodo']) == ('alta', 'texto+vision')
    assert salida['nombre_padre']['confianza'] == 'baja' and salida['nombre_padre']['detalle']['otra_lectura'] == 'PEDRO GARCIA RUIS'
    assert salida['lugar_nacimiento']['confianza'] == 'baja'  # El modelo no estaba seguro.
    assert 'curp_madre' not in salida
    # Lo que no coincide con la lista oficial no la cambia: queda por confirmar con lo que dice la lista.
    otra = {r['clave']: r for r in taller.combinar(lista, {'nombres': ('JUANA', 'texto')}, {}, identidad, HOY)}
    assert otra['nombres']['confianza'] == 'baja' and otra['nombres']['detalle']['en_lista'] == 'JUAN'


def test_los_procesos_fijos_leen_el_acta_la_ine_y_el_certificado():
    curp = generar.curp('GAHJ200512HMCRRNA')
    acta = generar.acta('JUAN', 'GARCIA', 'HERNANDEZ', '12/05/2020', 'HOMBRE', curp, 'PEDRO GARCIA RUIZ', 'ROSA HERNANDEZ DIAZ')
    leidos = taller.de_procesos('acta_alumno', acta, [], HOY)
    assert leidos['curp'][0] == curp and leidos['nombre_madre'][0] == 'ROSA HERNANDEZ DIAZ' and leidos['nombre_padre'][0] == 'PEDRO GARCIA RUIZ'
    ine = taller.de_procesos('ine_tutor', generar.ine('RAMOS', 'SOTO', 'MARIA ELENA', 'C. PINO NO. 5 COL. CENTRO TOLUCA'), [], HOY)
    assert ine == {'nombre_tutor': ('MARIA ELENA RAMOS SOTO', 'texto'), 'domicilio_ine': ('C. PINO NO. 5 COL. CENTRO TOLUCA', 'texto')}
    # Solo lo necesario: del certificado ya no se saca si es apto ni su fecha; lo que importa lo lee el modelo que ve.
    assert taller.de_procesos('certificado_medico', generar.medico('JUAN GARCIA HERNANDEZ'), [], HOY) == {}


def test_el_telefono_del_padre_no_se_toma_del_renglon_de_la_madre():
    lista = campos.revisar_campos([{'etiqueta': 'Teléfono de la madre'}, {'etiqueta': 'Teléfono del padre'},
                                   {'etiqueta': 'Teléfono de emergencia'}])
    hoja = 'Teléfono de la madre: 722 111 2233\nTeléfono de emergencia: 722 555 6677'
    assert campos.de_texto(hoja, lista, HOY) == {'telefono_de_la_madre': '722 111 2233', 'telefono_de_emergencia': '722 555 6677'}


def test_lo_que_lee_solo_el_ocr_o_dos_lecturas_distintas_queda_por_confirmar(monkeypatch):
    lista = [campos.campo('cp', 'Código postal', 'cp', 'domicilio'), campos.campo('telefono', 'Teléfono', 'telefono', 'contacto', 'telefono'),
             campos.campo('alergias', 'Alergias', 'lista', 'salud', 'alergia')]
    salida = {r['clave']: r for r in taller.combinar(lista, {'cp': ('32079', 'ocr')}, None, {}, HOY)}
    assert salida['cp']['confianza'] == 'baja'  # Del OCR de una foto: un 5 pudo leerse 3.
    lecturas = iter([{'cp': {'valor': '52079', 'seguro': True}, 'telefono': {'valor': '722 123 4567', 'seguro': True},
                      'alergias': {'valor': 'polvo y nuez', 'seguro': True}},
                     {'cp': {'valor': '52079', 'seguro': True}, 'telefono': {'valor': '722 128 4567', 'seguro': True},
                      'alergias': {'valor': 'Nuez, polvo', 'seguro': True}}])
    monkeypatch.setattr(vision, 'preguntar', lambda *a, **k: (next(lecturas), 60.0))
    vistos, segundos = taller.ver([b'hoja'], 'Ficha', lista, 'modelo', lecturas=2, hoy=HOY)
    assert segundos == 120.0
    assert vistos['cp'] == {'valor': '52079', 'seguro': True, 'lecturas': ['52079', '52079']} and vistos['alergias']['seguro'] is True
    assert vistos['telefono'] == {'valor': '722 123 4567', 'seguro': False, 'otra': '722 128 4567', 'lecturas': ['722 123 4567', '722 128 4567']}
    salida = {r['clave']: r for r in taller.combinar(lista, {}, vistos, {}, HOY)}
    assert (salida['cp']['confianza'], salida['telefono']['confianza']) == ('media', 'baja')
    assert salida['telefono']['detalle']['otra_lectura'] == '722 128 4567'
    # Las listas: un resultado por valor, con cuántas lecturas lo vieron. Las dos vieron polvo y nuez: vigentes.
    alergias = [r for r in taller.combinar(lista, {}, vistos, {}, HOY) if r['clave'] == 'alergias']
    assert {(r['valor'], r['confianza'], r['detalle']['lecturas']) for r in alergias} == {('POLVO', 'media', 2), ('NUEZ', 'media', 2)}


def test_la_fecha_de_nacimiento_sale_de_la_curp_y_ninguna_lo_confirma_una_persona():
    curp = generar.curp('GAHJ200512HMCRRNA')
    lista = [campos.campo('curp', 'CURP', 'curp', 'alumno'), campos.campo('fecha_nacimiento', 'Fecha de nacimiento', 'fecha', 'alumno'),
             campos.campo('alergias', 'Alergias', 'lista', 'salud', 'alergia')]
    # El modelo no leyó la fecha: sale de la CURP, que cuadra su dígito.
    salida = {r['clave']: r for r in taller.combinar(lista, {'curp': (curp, 'texto')}, {'alergias': {'valor': 'Negadas', 'seguro': False}}, {}, HOY)}
    assert (salida['fecha_nacimiento']['valor'], salida['fecha_nacimiento']['metodo']) == ('2020-05-12', 'curp')
    assert salida['fecha_nacimiento']['confianza'] == 'media'
    # "Negadas" leído una sola vez queda por confirmar: si era una alergia mal leída, la ve una persona.
    assert salida['alergias']['nada'] and salida['alergias']['confianza'] == 'baja' and salida['alergias']['valor'] == 'NINGUNA'
    # Si la fecha leída no es la de la CURP, se avisa con las dos.
    otra = {r['clave']: r for r in taller.combinar(lista, {'curp': (curp, 'texto')}, {'fecha_nacimiento': {'valor': '12-mayo-2019', 'seguro': True}}, {}, HOY)}
    assert otra['fecha_nacimiento']['confianza'] == 'baja' and otra['fecha_nacimiento']['detalle']['otra_lectura'] == '2020-05-12'
    assert campos.fecha('14-marzo-2018') == '2018-03-14' and campos.fecha('14 de marzo del 2018') == '2018-03-14'


def test_la_medicion_de_letra_a_mano_cuenta_aciertos_sin_sacar_datos(tmp_path, monkeypatch):
    pytest.importorskip('cv2')
    import csv
    from validacion import medir_lectura
    from validacion.ejemplo import manuscritos
    assert manuscritos.crear(tmp_path, cuantas=2) == 4
    respuestas = list(csv.DictReader(open(tmp_path / 'respuestas.csv', encoding='utf-8')))
    assert {r['tipo'] for r in respuestas} == {'ficha_inscripcion', 'certificado_medico'}
    import pypdf  # El segundo certificado trae frente y vuelta: el contacto de emergencia viene atrás.
    assert len(pypdf.PdfReader(str(tmp_path / '02_certificado_medico.pdf')).pages) == 2
    assert {r['clave'] for r in respuestas if r['archivo'] == '02_certificado_medico.pdf'} >= {'contacto_emergencia', 'telefono_emergencia'}
    assert all(comparador.desglose_curp(r['valor']) for r in respuestas if r['clave'] == 'curp')
    # Un modelo que ve que lee bien todo, salvo un teléfono que lee distinto cada vez.
    verdad = {(r['archivo'], r['clave']): r['valor'] for r in respuestas}
    veces = iter(range(100))
    monkeypatch.setattr(vision, 'elegir', lambda url=None, pedido=None: 'modelo-de-prueba')
    monkeypatch.setattr(vision, 'paginas', lambda contenido, nombre='': [nombre.encode()])

    def preguntar(hojas, pregunta, esquema=None, nombre_modelo=None, url=None, espera=900, **opciones):
        archivo, vez = hojas[0].decode(), next(veces)
        if esquema is None:  # La lectura completa no saca nada aquí: cuenta lo de las lecturas con campos.
            return '', 1.0
        if 'respuesta' in esquema['properties']:  # La revisión de cada dato de salud.
            return {'respuesta': 'lo_tiene', 'dice': ''}, 1.0
        extras = {'emergencia_en_la_hoja': verdad.get((archivo, 'contacto_emergencia'), ''),
                  'telefonos_en_la_hoja': f"{verdad[(archivo, 'telefono_emergencia')]} (emergencia)" if (archivo, 'telefono_emergencia') in verdad else ''}
        leido = {c: {'valor': extras.get(c, verdad.get((archivo, c), '')), 'seguro': True} for c in esquema['properties']}
        if 'telefono_padre' in leido:
            leido['telefono_padre']['valor'] = f'72200000{vez:02d}'
        return leido, 1.0

    monkeypatch.setattr(vision, 'preguntar', preguntar)
    monkeypatch.setattr(vision, 'cargar', lambda nombre_modelo, url=None: 0.0)
    detalle, tiempos, modelo_vision = medir_lectura.medir(tmp_path, HOY, lecturas=2)
    cuenta = {c: sum(1 for d in detalle if d['resultado'] == c) for c in medir_lectura.CLASES}
    assert len(detalle) == len(respuestas) and cuenta['peligroso'] == 0 and cuenta['corregir'] == 2  # El teléfono del padre de cada ficha.
    atras = {d['clave']: d['resultado'] for d in detalle if d['archivo'] == '02_certificado_medico.pdf'}
    assert atras['contacto_emergencia'] == atras['telefono_emergencia'] == 'bien_solo'
    texto = medir_lectura.resumen(detalle, tiempos, modelo_vision, 2)
    assert 'Bien (solo o con un clic)' in texto
    assert not any(r['valor'] in texto for r in respuestas if len(r['valor']) > 4)  # Solo números en el resumen.


def test_el_modelo_que_ve_no_piensa_en_voz_alta_y_recibe_la_hoja_achicada(monkeypatch):
    """qwen3-vl "pensaba" 4 minutos antes de contestar y la hoja de 1600 px le costaba 3,000 tokens (PC 4, Tarea 012)."""
    pytest.importorskip('PIL')
    import io
    import json
    import urllib.error
    from PIL import Image
    hoja = io.BytesIO()
    Image.new('RGB', (2550, 3300), 'white').save(hoja, 'PNG')
    pedidos = []

    class Respuesta(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *_):
            self.close()

    def urlopen(solicitud, timeout=None):
        cuerpo = json.loads(solicitud.data)
        pedidos.append(cuerpo)
        if cuerpo.get('think') is not None and len(pedidos) == 1 and viejo:
            raise urllib.error.HTTPError(solicitud.full_url, 400, 'Bad Request', {}, io.BytesIO(b'{"error":"unknown field think"}'))
        return Respuesta(json.dumps({'message': {'content': '{"tipo_sangre": {"valor": "O+", "seguro": true}}'}}).encode())

    monkeypatch.setattr(vision.urllib.request, 'urlopen', urlopen)
    monkeypatch.setattr(vision, 'LADO', 1280)
    viejo = False
    leido, _ = vision.preguntar([hoja.getvalue()], 'Lee', {'type': 'object'}, 'modelo-de-prueba', 'http://127.0.0.1:1/api/chat')
    assert leido['tipo_sangre']['valor'] == 'O+' and pedidos[0]['think'] is False
    assert max(Image.open(io.BytesIO(base64.b64decode(pedidos[0]['messages'][0]['images'][0]))).size) == 1280
    viejo, pedidos[:] = True, []  # Un Ollama que no conoce "think": se vuelve a pedir sin él.
    leido, _ = vision.preguntar([hoja.getvalue()], 'Lee', {'type': 'object'}, 'modelo-de-prueba', 'http://127.0.0.1:1/api/chat')
    assert len(pedidos) == 2 and 'think' not in pedidos[1] and leido['tipo_sangre']['valor'] == 'O+'


def test_el_modelo_del_ajuste_que_no_esta_instalado_no_apaga_la_lectura(monkeypatch):
    monkeypatch.setattr(vision, 'instalados', lambda url=None: ['llama3.1:latest', 'qwen2.5vl:3b'])
    monkeypatch.setenv('SIDP_MODELO_VISION', 'qwen3-vl:8b')
    assert vision.elegir() == 'qwen2.5vl:3b'
    assert vision.elegir(pedido='qwen3-vl:8b') is None  # Pedido a mano: ese o ninguno.
    monkeypatch.setenv('SIDP_MODELO_VISION', 'qwen2.5vl:3b')
    assert vision.elegir() == 'qwen2.5vl:3b'


def test_si_los_procesos_ya_leyeron_todo_sin_ocr_no_se_gasta_el_modelo_que_ve():
    lista = [campos.campo('curp', 'CURP', 'curp', 'alumno'), campos.campo('fecha_nacimiento', 'Fecha', 'fecha', 'alumno')]
    assert not taller.necesita_ver(lista, {'curp': ('X', 'qr'), 'fecha_nacimiento': ('2020-05-12', 'texto')})
    assert taller.necesita_ver(lista, {'curp': ('X', 'qr'), 'fecha_nacimiento': ('2020-05-12', 'ocr')})
    assert taller.necesita_ver(lista, {'curp': ('X', 'qr')})
    # Lo que los procesos sacan del texto de una foto queda marcado como OCR (por confirmar si va solo).
    ine = taller.de_procesos('ine_tutor', generar.ine('RAMOS', 'SOTO', 'MARIA ELENA', 'C. PINO NO. 5 COL. CENTRO TOLUCA'), [], HOY, 'ocr')
    assert ine and {metodo for _, metodo in ine.values()} == {'ocr'}


# ---------------------------------------------------------------- En la base de la PC 4

if os.environ.get('SIDP_PRUEBA_BASE'):
    pytest.importorskip('pg8000')
    from pruebas.test_base_cerebro import CURP_JUAN, alumnos, base, id_de  # noqa: F401  (base es un fixture)
else:
    CURP_JUAN = generar.curp('GAHJ200512HMCRRNA')

    @pytest.fixture
    def base():
        pytest.skip('sin PostgreSQL de prueba (SIDP_PRUEBA_BASE)')
con_base = pytest.mark.skipif(not os.environ.get('SIDP_PRUEBA_BASE'), reason='sin PostgreSQL de prueba (SIDP_PRUEBA_BASE)')

ACTA = generar.acta('JUAN', 'GARCIA', 'HERNANDEZ', '12/05/2020', 'HOMBRE', CURP_JUAN, 'PEDRO GARCIA RUIZ', 'ROSA HERNANDEZ DIAZ')
# Lo que "ve" el modelo en cada documento (con letra a mano: un apellido mal leído, un medicamento dudoso). En el certificado,
# el modelo se confunde con los antecedentes de la familia y "ve" hipertensión: al revisarlo dice que es de un familiar.
VISTO = {
    'Acta de nacimiento': {'lugar_nacimiento': ('Toluca, México', True), 'nombre_madre': ('Rosa Hernández Díaz', True),
                           'nombre_padre': ('PEDRO GARCIA RUIS', True)},
    'Certificado médico': {'alergias': ('Penicilina y polvo', True), 'padecimientos': ('Hipertensión', True),
                           'medicamentos': ('Salbutamol', False), 'discapacidad': ('Ninguna', True),
                           'telefonos_en_la_hoja': ('722 123 4567 (mamá)', True), 'correos_en_la_hoja': ('rosa.hdz@gmail.com', True),
                           'emergencia_en_la_hoja': ('Rosa Hernández Díaz, mamá', True),
                           'peso': ('22,5 kg', True), 'tipo_sangre': ('O positivo', True)},  # Peso y sangre: el modelo no los ve.
}
# La lectura completa de cada hoja (el certificado con su reverso). Trae el teléfono del médico, que no es de la familia.
TRANSCRITO = {
    'certificado_medico_1.pdf': [
        'CERTIFICADO MÉDICO\nNombre del paciente: Juan García Hernández\nAlergias: penicilina, polvo\nPadecimientos: ninguno\n'
        'Medicamentos: salbutamol\nDr. Ramón Pérez · Cédula 1234567 · Tel. 722 555 0000',
        'ANTECEDENTES HEREDOFAMILIARES:\nHipertensión [X] Diabetes [X]\nEnfermedad mental o nerviosa [X]\n'
        'En caso de emergencia avisar a: Rosa Hernández Díaz (mamá) Tel. 722 123 4567\nCorreo: rosa.hdz@gmail.com'],
}
REVISADO = {'HIPERTENSION': 'familiar'}  # Lo que contesta el modelo al revisar cada dato de salud (si no está: "lo_tiene").


CLASIFICADO = {'tipo': 'certificado_medico', 'titulo': 'Constancia de salud', 'motivo': 'trae el sello de un centro de salud'}


@pytest.fixture
def modelo_que_ve(monkeypatch):
    preguntas = []

    def preguntar(hojas, pregunta, esquema=None, nombre_modelo=None, url=None, espera=900, **opciones):
        preguntas.append(pregunta)
        if pregunta == vision.TRANSCRIBIR:  # La lectura completa de una hoja: "archivo#número de hoja".
            archivo, numero = hojas[0].decode().split('#')
            return TRANSCRITO.get(archivo, [''] * 9)[int(numero)], 20.0
        if pregunta.startswith('Esta es la hoja de salud'):
            valor = pregunta.split('«')[1].split('»')[0]
            return {'respuesta': REVISADO.get(comparador.normalizar(valor), 'lo_tiene'), 'dice': ''}, 5.0
        if pregunta == campos.PROPONER:
            return {'titulo': 'Ficha de inscripción', 'senas': ['FICHA DE INSCRIPCION', 'Ciclo escolar 2026-2027', 'Datos del alumno', 'ficha de inscripción'],
                    'campos': ['Nombre del alumno', 'Teléfono de la madre', 'Alergias', 'Teléfono de la madre']}, 2.0
        if pregunta.startswith('Esta hoja viene del expediente'):  # Una hoja que no se reconoció: qué documento parece.
            assert {'certificado_medico', 'boleta', 'otro'} <= set(esquema['properties']['tipo']['enum'])
            return dict(CLASIFICADO), 30.0
        documento = next(v for k, v in VISTO.items() if k in pregunta)
        return {c: {'valor': documento.get(c, ('', True))[0], 'seguro': documento.get(c, ('', True))[1]} for c in esquema['properties']}, 1.5

    monkeypatch.setattr(vision, 'elegir', lambda url=None, pedido=None: 'modelo-de-prueba')
    monkeypatch.setattr(vision, 'paginas', lambda contenido, nombre='': [f'{nombre}#{n}'.encode() for n in range(len(TRANSCRITO.get(nombre, [''])))])
    monkeypatch.setattr(vision, 'preguntar', preguntar)
    return preguntas


def pedido_de_juan():
    return {'curp': CURP_JUAN, 'expediente': '0001', 'documentos': [
        {'archivo': 'acta_alumno_1.pdf', 'tipo': 'acta_alumno', 'nombre': 'Acta de nacimiento del alumno',
         'contenido': base64.b64encode(pdf_de_hojas([ACTA.splitlines()])).decode()},
        {'archivo': 'certificado_medico_1.pdf', 'tipo': 'certificado_medico', 'nombre': 'Certificado médico',
         'contenido': base64.b64encode(pdf_de_hojas([generar.medico('JUAN GARCIA HERNANDEZ').splitlines()])).decode()},
    ]}


def datos_de(base, curp=CURP_JUAN):
    return {(categoria, f['campo'], f['valor']): (f['confianza'], f['estado'])
            for categoria, fs in base.expediente_alumno(curp)['datos'].items() for f in fs}


@con_base
def test_el_taller_saca_los_datos_y_los_guarda_por_categoria(base, modelo_que_ve, tmp_path):
    from validacion import base_cerebro
    base.cargar(base_cerebro.ciclo_de(HOY), alumnos())
    trabajos = taller.Taller(base, tmp_path / 'trabajos', hoy=HOY)
    with pytest.raises(LookupError):
        trabajos.encargar({**pedido_de_juan(), 'curp': generar.curp('XEXX010101HNEXXXA')})
    trabajo = trabajos.encargar(pedido_de_juan())
    assert sorted(p.name for p in (tmp_path / 'trabajos' / str(trabajo)).iterdir()) == ['acta_alumno_1.pdf', 'certificado_medico_1.pdf']
    cuenta = trabajos.una_vez()
    assert {k: cuenta[k] for k in ('esperados', 'alta', 'media', 'baja', 'vacios', 'utiles')} == \
        {'esperados': 14, 'alta': 6, 'media': 4, 'baja': 2, 'vacios': 3, 'utiles': 83}  # 10 de 12 datos salieron seguros.
    # Frente y vuelta del certificado se leyeron completos; cada dato de salud se revisó otra vez y la hipertensión (de la
    # familia) no se guardó. Los contactos salieron de la hoja aunque el certificado no tenga esos campos.
    assert {k: cuenta[k] for k in ('hojas_completas', 'revisados', 'descartados', 'contactos')} == \
        {'hojas_completas': 3, 'revisados': 4, 'descartados': 1, 'contactos': 3}
    assert not (tmp_path / 'trabajos' / str(trabajo)).exists()  # Los archivos del trabajo ya no hacen falta.
    assert trabajos.una_vez() is None
    assert datos_de(base) == {
        ('alumno', 'lugar_nacimiento', 'TOLUCA, MÉXICO'): ('media', 'vigente'),
        ('madre', 'nombre', 'ROSA HERNANDEZ DIAZ'): ('alta', 'vigente'),
        ('padre', 'nombre', 'PEDRO GARCIA RUIZ'): ('baja', 'por_confirmar'),
        ('salud', 'alergia', 'PENICILINA'): ('media', 'vigente'),
        ('salud', 'alergia', 'POLVO'): ('media', 'vigente'),
        ('salud', 'medicamento', 'SALBUTAMOL'): ('baja', 'por_confirmar'),
        ('salud', 'discapacidad', 'NINGUNA'): ('media', 'vigente'),  # Se ve que el certificado sí se leyó.
        ('contacto', 'telefono', '7221234567'): ('alta', 'vigente'),  # El del médico (722 555 0000) no.
        ('contacto', 'correo', 'rosa.hdz@gmail.com'): ('alta', 'vigente'),
        ('contacto', 'nombre_emergencia', 'ROSA HERNANDEZ DIAZ'): ('alta', 'vigente'),
    }  # Ni peso ni tipo de sangre: no son de lo que se saca. Ni la hipertensión de los antecedentes de la familia.
    contacto = {d['campo']: d['detalle'] for d in base.expediente_alumno(CURP_JUAN)['datos']['contacto']}
    assert contacto['telefono']['de'] == 'madre' and contacto['nombre_emergencia']['parentesco'] == 'MAMA'
    # El nombre, la CURP y la fecha ya estaban en la lista oficial: el acta los confirma y no se repiten.
    assert not any(c == 'alumno' and k in ('curp', 'nombres', 'fecha_nacimiento') for c, k, _ in datos_de(base))
    estado = trabajos.estado()
    assert estado['trabajos'] == {'hecho': 1} and estado['porcentaje_util'] == 83 and estado['por_confirmar'] == 2
    # Solo números en el resumen (nada de nombres).
    assert set(base._uno('SELECT resumen FROM sidp.trabajo WHERE id = %s', (trabajo,))) == \
        {'esperados', 'alta', 'media', 'baja', 'vacios', 'segundos_vision', 'modelo', 'utiles', 'hojas_completas', 'revisados', 'descartados',
         'contactos', 'duplicados'}


@con_base
def test_una_persona_confirma_corrige_o_agrega_y_la_lectura_no_pisa_lo_confirmado(base, modelo_que_ve, tmp_path):
    from validacion import base_cerebro
    base.cargar(base_cerebro.ciclo_de(HOY), alumnos())
    trabajos = taller.Taller(base, tmp_path / 'trabajos', hoy=HOY)
    trabajos.encargar(pedido_de_juan())
    trabajos.una_vez()
    dudosos = {d['campo']: d for d in base.por_confirmar()}
    assert set(dudosos) == {'nombre', 'medicamento'} and dudosos['nombre']['curp'] == CURP_JUAN and dudosos['nombre']['grado'] == 1
    base.confirmar_dato(dudosos['nombre']['id'], '', 'secretaria')  # Estaba bien.
    base.confirmar_dato(dudosos['medicamento']['id'], 'SALBUTAMOL INHALADO', 'secretaria')  # Lo corrige viendo la hoja.
    with pytest.raises(base_cerebro.ErrorBase):  # Lo que no pide el gobierno ni necesita la escuela no se guarda.
        base.agregar_dato(CURP_JUAN, 'salud', 'peso', '23', 'secretaria')
    base.agregar_dato(CURP_JUAN, 'contacto', 'Teléfono', '7221234567', 'secretaria', {'de': 'madre'})
    base.agregar_dato(CURP_JUAN, 'contacto', 'telefono', '7227654321', 'secretaria', {'de': 'padre'})
    datos = datos_de(base)
    assert datos[('padre', 'nombre', 'PEDRO GARCIA RUIZ')] == ('confirmado', 'vigente')
    assert datos[('salud', 'medicamento', 'SALBUTAMOL INHALADO')] == ('confirmado', 'vigente') and ('salud', 'medicamento', 'SALBUTAMOL') not in datos
    assert datos[('contacto', 'telefono', '7221234567')] == datos[('contacto', 'telefono', '7227654321')] == ('confirmado', 'vigente')
    assert not base.por_confirmar()
    # Se vuelve a escanear: lo confirmado no cambia y la lectura distinta queda por confirmar.
    VISTO['Acta de nacimiento']['lugar_nacimiento'] = ('Metepec, México', True)
    base.confirmar_dato(base._uno("SELECT id FROM sidp.dato WHERE campo = 'lugar_nacimiento' AND estado = 'vigente'"), '', 'secretaria')
    try:
        trabajos.encargar(pedido_de_juan())
        trabajos.una_vez()
    finally:
        VISTO['Acta de nacimiento']['lugar_nacimiento'] = ('Toluca, México', True)
    datos = datos_de(base)
    assert datos[('alumno', 'lugar_nacimiento', 'TOLUCA, MÉXICO')] == ('confirmado', 'vigente')
    assert datos[('alumno', 'lugar_nacimiento', 'METEPEC, MÉXICO')] == ('media', 'por_confirmar')
    assert datos[('salud', 'medicamento', 'SALBUTAMOL INHALADO')] == ('confirmado', 'vigente')
    assert ('salud', 'medicamento', 'SALBUTAMOL') not in datos  # Ya se corrigió una vez: no se vuelve a preguntar.
    with pytest.raises(Exception):
        base._consulta('DELETE FROM sidp.dato')  # En la base no se borra nada a mano.


@con_base
def test_volver_a_leer_un_expediente_reemplaza_lo_de_la_maquina_y_no_lo_de_las_personas(base, modelo_que_ve, tmp_path):
    from validacion import base_cerebro
    base.cargar(base_cerebro.ciclo_de(HOY), alumnos())
    trabajos = taller.Taller(base, tmp_path / 'trabajos', hoy=HOY)
    trabajos.encargar(pedido_de_juan())
    trabajos.una_vez()
    padre = next(d for d in base.por_confirmar() if d['campo'] == 'nombre')
    base.confirmar_dato(padre['id'], '', 'secretaria')
    antes = base._uno("SELECT count(*) FROM sidp.dato WHERE quien = 'cerebro' AND estado IN ('vigente', 'por_confirmar')")
    # La auditoría pide leer otra vez el mismo expediente; pedido dos veces, se lee una.
    primero = trabajos.encargar({**pedido_de_juan(), 'relectura': True})
    assert trabajos.encargar({**pedido_de_juan(), 'relectura': True}) == primero
    cuenta = trabajos.una_vez()
    assert trabajos.una_vez() is None
    assert cuenta['reemplazados'] == antes  # Todo lo que había leído la máquina de ese expediente.
    marcados = base._consulta("SELECT DISTINCT (detalle->>'relectura')::bigint FROM sidp.dato WHERE estado = 'reemplazado' AND quien = 'cerebro' "
                              "AND detalle ? 'relectura'")
    assert marcados == [[primero]]
    datos = datos_de(base)
    assert datos[('padre', 'nombre', 'PEDRO GARCIA RUIZ')] == ('confirmado', 'vigente')  # Lo de la persona no se tocó.
    assert datos[('salud', 'alergia', 'PENICILINA')] == ('media', 'vigente') and datos[('contacto', 'telefono', '7221234567')] == ('alta', 'vigente')
    vigentes = base._consulta("SELECT categoria, campo, sidp.normalizar(valor), count(*) FROM sidp.dato WHERE estado IN ('vigente', 'por_confirmar') "
                              "GROUP BY 1, 2, 3 HAVING count(*) > 1")
    assert vigentes == []  # Ningún dato repetido.


@con_base
def test_la_depuracion_marca_las_lecturas_repetidas_sin_borrar(base):
    from validacion import base_cerebro
    base.cargar(base_cerebro.ciclo_de(HOY), alumnos())
    alumno = id_de(base, CURP_JUAN)
    guardar = 'SELECT sidp.guardar_dato(%s, %s, %s, %s, CAST(%s AS jsonb), %s, %s, %s, %s, %s, %s)'
    for categoria, campo, valor, nivel, varios in (('salud', 'alergia', 'PENICILINA', 'media', True), ('salud', 'alergia', 'PENISILINA', 'baja', True),
                                                   ('salud', 'alergia', 'NINGUNA', 'media', True), ('contacto', 'telefono', '7221234567', 'alta', True),
                                                   ('contacto', 'telefono', '7221234561', 'baja', True)):
        base._uno(guardar, (alumno, categoria, campo, valor, '{}', nivel, 'vision', 'Certificado médico', None, 'cerebro', varios))
    base.agregar_dato(CURP_JUAN, 'salud', 'padecimiento', 'ASMA', 'secretaria')
    assert taller.depurar(base) == 3
    assert taller.depurar(base) == 0  # Ya no hay qué marcar.
    quedan = {(d[0], d[1]) for d in base._consulta("SELECT campo, valor FROM sidp.dato WHERE estado IN ('vigente', 'por_confirmar')")}
    assert quedan == {('alergia', 'PENICILINA'), ('telefono', '7221234567'), ('padecimiento', 'ASMA')}
    assert base._uno("SELECT count(*) FROM sidp.dato WHERE estado = 'reemplazado' AND detalle ? 'duplicado_de'") == 3  # Siguen ahí.


@con_base
def test_la_base_misma_no_deja_guardar_dos_veces_el_mismo_dato(base):
    from validacion import base_cerebro
    base.cargar(base_cerebro.ciclo_de(HOY), alumnos())
    alumno = id_de(base, CURP_JUAN)
    insertar = ("INSERT INTO sidp.dato (alumno_id, categoria, campo, valor, confianza, metodo, estado, quien) "
                "VALUES (%s, 'salud', 'alergia', %s, %s, 'vision', %s, 'cerebro') RETURNING id")
    # Una base de antes, sin la defensa: el mismo dato quedó tres veces (con y sin acento).
    base._consulta('DROP INDEX sidp.dato_sin_repetir')
    ids = [base._uno(insertar, (alumno, valor, nivel, estado)) for valor, nivel, estado in
           (('PENICILINA', 'baja', 'por_confirmar'), ('Penicilína', 'alta', 'vigente'), ('penicilina ', 'media', 'vigente'))]
    base.preparar()  # Al arrancar: lo repetido queda reemplazado (no se borra) y se queda el más seguro.
    filas = base._consulta('SELECT id, estado, detalle FROM sidp.dato WHERE id = ANY(%s) ORDER BY id', (ids,))
    assert [f[1] for f in filas] == ['reemplazado', 'vigente', 'reemplazado']
    assert filas[0][2]['duplicado_de'] == filas[2][2]['duplicado_de'] == ids[1]
    with pytest.raises(Exception):  # Ni metiéndolo a mano.
        base._uno(insertar, (alumno, 'PENICILINA', 'media', 'vigente'))
    guardar = 'SELECT sidp.guardar_dato(%s, %s, %s, %s, CAST(%s AS jsonb), %s, %s, %s, %s, %s, %s)'
    assert base._uno(guardar, (alumno, 'salud', 'alergia', 'PENICILINA', '{}', 'baja', 'vision', None, None, 'cerebro', True)) == ids[1]


@con_base
def test_la_base_avisa_de_alumnos_que_pueden_ser_el_mismo_nino(base):
    from validacion import base_cerebro
    base.cargar(base_cerebro.ciclo_de(HOY), alumnos())
    juan = id_de(base, CURP_JUAN)
    nuevo = ("INSERT INTO sidp.alumno (curp, nombres, primer_apellido, segundo_apellido, fecha_nacimiento) "
             "VALUES (%s, %s, %s, %s, %s) RETURNING id")
    sin_curp = base._uno(nuevo, (None, 'Juan', 'García', 'Hernández', None))  # El mismo nombre, cargado sin CURP.
    otra = CURP_JUAN[:10] + ('M' if CURP_JUAN[10] == 'H' else 'H') + CURP_JUAN[11:]
    curp_mal = base._uno(nuevo, (otra, 'Juana', 'Garza', 'Hernández', None))  # La CURP a una letra de la de Juan.
    gemela = base._uno(nuevo, (None, 'Juana', 'García', 'Hernández', None))  # Nombre parecido, pero es otra niña.
    repetidos = {(r['a']['id'], r['b']['id']): r['motivo'] for r in base.alumnos_repetidos()}
    assert repetidos[(juan, sin_curp)] == 'mismo_nombre' and repetidos[(juan, curp_mal)] == 'curp_casi_igual'
    assert all(r['a']['nombre'] and 'grado' in r['a'] for r in base.alumnos_repetidos())
    for par in [p for p in repetidos if gemela in p]:  # Una persona dice que son distintos: ya no se pregunta.
        assert base.marcar_distintos(par[1], par[0], 'directora') is True
    assert base.marcar_distintos(juan, gemela, 'directora') in (True, False)  # Repetirlo no rompe nada.
    assert not [r for r in base.alumnos_repetidos() if gemela in (r['a']['id'], r['b']['id'])]
    assert base._uno("SELECT count(*) FROM sidp.bitacora WHERE tabla = 'alumnos_distintos'") >= 1
    with pytest.raises(Exception):
        base.marcar_distintos(juan, juan, 'directora')
    base.borrar_alumno(CURP_JUAN, 'Se cargó dos veces', 'directora')  # Borrar a uno no se atora con lo anotado.


@con_base
def test_baja_y_borrado_de_un_alumno(base):
    from validacion import base_cerebro
    base.cargar(base_cerebro.ciclo_de(HOY), alumnos())
    otro = alumnos()[1]['curp']
    base.agregar_dato(otro, 'salud', 'alergia', 'POLVO', 'secretaria')
    with pytest.raises(base_cerebro.ErrorBase):
        base.dar_baja_alumno(otro, ' ', 'secretaria')
    base.dar_baja_alumno(otro, 'Cambio de domicilio', 'secretaria')
    assert base._uno('SELECT estado FROM sidp.alumno WHERE curp = %s', (otro,)) == 'baja'
    assert otro not in {a['curp'] for g in base.alumnos()['grupos'] for a in g['filas']}
    assert datos_de(base, otro)  # Su expediente se queda.
    # Borrar: se quita de todas las tablas y la bitácora dice quién, cuándo y por qué.
    with pytest.raises(Exception):
        base.borrar_alumno(CURP_JUAN, '', 'directora')
    base.agregar_dato(CURP_JUAN, 'contacto', 'telefono', '7221234567', 'secretaria')
    alumno = id_de(base, CURP_JUAN)
    base.borrar_alumno(CURP_JUAN, 'Se cargó dos veces', 'directora')
    assert base._uno('SELECT count(*) FROM sidp.alumno WHERE curp = %s', (CURP_JUAN,)) == 0
    assert base._uno('SELECT count(*) FROM sidp.dato WHERE alumno_id = %s', (alumno,)) == 0
    assert base._consulta("SELECT actor, despues->>'motivo' FROM sidp.bitacora WHERE accion = 'borrar'") == [['directora', 'Se cargó dos veces']]
    with pytest.raises(base_cerebro.ErrorBase):
        base.expediente_alumno(CURP_JUAN)
    with pytest.raises(Exception):  # Fuera de borrar_alumno sigue sin poderse borrar nada.
        base._consulta('DELETE FROM sidp.alumno WHERE curp = %s', (otro,))


@con_base
def test_del_escaneo_al_expediente_del_alumno_en_el_panel(base, modelo_que_ve, tmp_path):
    """El servidor guarda los PDF por alumno, la PC 4 saca los datos y el panel muestra todo junto."""
    trabajos = taller.Taller(base, tmp_path / 'trabajos', hoy=HOY)
    cerebro = cerebro_red.crear_servidor('clave-de-prueba', 0, url='http://127.0.0.1:9/api/chat', anfitrion='127.0.0.1', base=base,
                                         taller=trabajos)
    url_cerebro = iniciar(cerebro)
    app = aplicacion.Aplicacion(tmp_path / 'sistema', HOY, puerto=0, expedientes=tmp_path / 'Expedientes')
    servidor = aplicacion.crear_servidor(app, '127.0.0.1')
    url = iniciar(servidor)
    try:
        assert app.configurar(url_cerebro, 'clave-de-prueba') is None
        token = {'X-SIDP-Token': app.token}
        filas = [{k: a[k] for k in ('curp', 'primer_apellido', 'segundo_apellido', 'nombres')} for a in alumnos()]
        assert pedir(url, '/api/listas/guardar', {'grado': 1, 'letra': 'A', 'filas': filas}, token)[0] == 200
        leido = escanear(url, app, pdf_de_hojas([ACTA.splitlines(), generar.medico('JUAN GARCIA HERNANDEZ').splitlines()]))[1]
        paginas = {p['archivo']: p['tipo'] for p in leido['paginas']}
        assert sorted(paginas.values()) == ['acta_alumno', 'certificado_medico']
        codigo, guardado = pedir(url, '/api/escaneo/guardar', {'lote': leido['lote'], 'alumno': leido['propuesta']['alumno']}, token)
        assert codigo == 200, guardado
        # El servidor ya le encargó el expediente a la PC 4: no quedó nada pendiente de mandar.
        estado = pedir(url, '/api/trabajos', cabeceras=token)[1]
        assert estado['trabajos'] == {'pendiente': 1} and estado['pendientes_de_enviar'] == 0, estado
        assert trabajos.una_vez()['esperados'] == 14  # 10 del acta y 4 del certificado médico (solo lo necesario).
        estado = pedir(url, '/api/trabajos', cabeceras=token)[1]
        assert estado['trabajos'] == {'hecho': 1} and estado['porcentaje_util'] == 83
        codigo, expediente = pedir(url, f'/api/alumno?curp={CURP_JUAN}', cabeceras=token)
        assert codigo == 200, expediente
        assert expediente['alumno']['nombres'] == 'JUAN' and expediente['alumno']['grado'] == 1
        assert [d['valor'] for d in expediente['datos']['madre']] == ['ROSA HERNANDEZ DIAZ']
        assert {d['valor'] for d in expediente['datos']['salud'] if d['campo'] == 'alergia'} == {'PENICILINA', 'POLVO'}
        assert sorted(p['nombre'] for p in expediente['pdf']) == ['Acta de nacimiento del alumno', 'Certificado médico']
        # El PDF se abre en la ventana; nada fuera de la carpeta de expedientes.
        ruta = expediente['pdf'][0]['ruta']
        solicitud = urllib.request.Request(f'{url}/api/pdf?ruta={urllib.request.quote(ruta)}', headers=token)
        with urllib.request.urlopen(solicitud) as respuesta:
            assert respuesta.headers['Content-Type'] == 'application/pdf' and respuesta.read().startswith(b'%PDF')
        assert pedir(url, '/api/pdf?ruta=../sistema/plataforma_prueba.db', cabeceras=token)[0] == 400
        # Lo dudoso se confirma desde el panel.
        pendientes = pedir(url, '/api/por_confirmar', cabeceras=token)[1]
        dudosos = pendientes['datos']
        assert {d['campo'] for d in dudosos} == {'nombre', 'medicamento'} and pendientes['repetidos'] == []
        mismo = pedir(url, '/api/alumno/distintos', {'a': 1, 'b': 1}, token)  # Un alumno no se repite consigo mismo.
        assert mismo[0] != 200 and 'mismo alumno' in mismo[1]['error']
        assert pedir(url, '/api/dato/confirmar', {'id': dudosos[0]['id'], 'valor': ''}, token)[0] == 200
        assert pedir(url, '/api/dato/agregar', {'curp': CURP_JUAN, 'categoria': 'contacto', 'campo': 'telefono', 'valor': '7221234567',
                                                'de': 'madre'}, token)[0] == 200
        assert len(pedir(url, '/api/por_confirmar', cabeceras=token)[1]['datos']) == 1
        # El Excel por grado junta la lista de aquí con lo que la PC 4 sacó del escaneo.
        import io
        import openpyxl
        solicitud = urllib.request.Request(f'{url}/api/exportar?que=por_grado&formato=xlsx', headers=token)
        with urllib.request.urlopen(solicitud) as respuesta:
            libro = openpyxl.load_workbook(io.BytesIO(respuesta.read()))
        hoja = libro['1°']
        encabezado = [c.value for c in hoja[3]]
        juan = next(dict(zip(encabezado, [c.value for c in fila])) for fila in hoja.iter_rows(min_row=4) if fila[3].value == 'JUAN')
        assert juan['Alergias'] == 'PENICILINA, POLVO' and juan['CURP'] == CURP_JUAN and juan['Sexo'] == 'Hombre'
        assert libro['Resumen'].cell(5, 3).value == len(filas)  # Alumnos de 1°.
        # Baja: necesita motivo. Borrar: además hay que escribir BORRAR.
        assert pedir(url, '/api/alumno/baja', {'curp': CURP_JUAN}, token)[0] == 400
        assert pedir(url, '/api/alumno/borrar', {'curp': CURP_JUAN, 'motivo': 'error', 'confirmar': 'si'}, token)[0] == 400
        codigo, respuesta = pedir(url, '/api/alumno/baja', {'curp': CURP_JUAN, 'motivo': 'Se mudó'}, token)
        assert codigo == 200, respuesta
        assert base._uno('SELECT estado FROM sidp.alumno WHERE curp = %s', (CURP_JUAN,)) == 'baja'
        assert app.bandeja.registro.db.execute('SELECT estado FROM alumno WHERE curp = ?', (CURP_JUAN,)).fetchone() == ('baja',)
    finally:
        servidor.shutdown()
        cerebro.shutdown()


@con_base
def test_lo_escaneado_sin_la_pc_4_se_manda_cuando_vuelve(base, modelo_que_ve, tmp_path):
    trabajos = taller.Taller(base, tmp_path / 'trabajos', hoy=HOY)
    app = aplicacion.Aplicacion(tmp_path / 'sistema', HOY, puerto=0, expedientes=tmp_path / 'Expedientes')
    servidor = aplicacion.crear_servidor(app, '127.0.0.1')
    url = iniciar(servidor)
    cerebro = cerebro_red.crear_servidor('clave-de-prueba', 0, url='http://127.0.0.1:9/api/chat', anfitrion='127.0.0.1', base=base,
                                         taller=trabajos)
    url_cerebro = iniciar(cerebro)
    try:
        token = {'X-SIDP-Token': app.token}
        filas = [{k: a[k] for k in ('curp', 'primer_apellido', 'segundo_apellido', 'nombres')} for a in alumnos()]
        pedir(url, '/api/listas/guardar', {'grado': 1, 'letra': 'A', 'filas': filas}, token)
        leido = escanear(url, app, pdf_de_hojas([ACTA.splitlines()]))[1]
        assert pedir(url, '/api/escaneo/guardar', {'lote': leido['lote'], 'alumno': leido['propuesta']['alumno']}, token)[0] == 200
        assert [p.name for p in (tmp_path / 'sistema' / 'pendientes_pc4').iterdir()] == ['0001.json']
        # La PC 4 vuelve; ella todavía no tiene la lista de alumnos (409) y el servidor se la manda primero.
        assert app.configurar(url_cerebro, 'clave-de-prueba') is None
        assert app.bandeja.enviar_pendientes() == 1
        assert not list((tmp_path / 'sistema' / 'pendientes_pc4').iterdir())
        assert trabajos.una_vez()['esperados'] == 10
    finally:
        servidor.shutdown()
        cerebro.shutdown()


@con_base
def test_la_ia_propone_los_campos_de_una_plantilla_nueva(base, modelo_que_ve, tmp_path):
    cerebro = cerebro_red.crear_servidor('clave-de-prueba', 0, url='http://127.0.0.1:9/api/chat', anfitrion='127.0.0.1', base=base)
    url_cerebro = iniciar(cerebro)
    app = aplicacion.Aplicacion(tmp_path / 'sistema', HOY, puerto=0)
    servidor = aplicacion.crear_servidor(app, '127.0.0.1')
    url = iniciar(servidor)
    try:
        assert app.configurar(url_cerebro, 'clave-de-prueba') is None
        muestra = pdf_de_hojas([['FICHA DE INSCRIPCION']])
        cuerpo = (b'--limite\r\nContent-Disposition: form-data; name="archivo"; filename="ficha.pdf"\r\n'
                  b'Content-Type: application/pdf\r\n\r\n' + muestra + b'\r\n--limite--\r\n')
        codigo, respuesta = pedir(url, '/api/plantillas/campos', cuerpo, {'X-SIDP-Token': app.token}, tipo='multipart/form-data; boundary=limite')
        assert codigo == 200, respuesta
        assert [(c['etiqueta'], c['categoria'], c['campo']) for c in respuesta['campos']] == [
            ('Nombre del alumno', 'alumno', 'nombre'), ('Teléfono de la madre', 'contacto', 'telefono'), ('Alergias', 'salud', 'alergia')]
        # Las señas para reconocerlo al escanear: solo texto impreso fijo (sin números) y sin repetir.
        assert respuesta['titulo'] == 'Ficha de inscripción' and respuesta['senas'] == ['FICHA DE INSCRIPCION', 'Datos del alumno']
        # Sin taller (la PC 4 no tiene base o no arrancó) no se aceptan trabajos.
        assert pedir(url_cerebro, '/trabajo', pedido_de_juan(), {'X-SIDP-Clave': 'clave-de-prueba'})[0] == 503
    finally:
        servidor.shutdown()
        cerebro.shutdown()


@con_base
def test_la_ia_dice_que_documento_parece_una_hoja_que_no_se_reconocio(base, modelo_que_ve, tmp_path):
    cerebro = cerebro_red.crear_servidor('clave-de-prueba', 0, url='http://127.0.0.1:9/api/chat', anfitrion='127.0.0.1', base=base)
    url_cerebro = iniciar(cerebro)
    app = aplicacion.Aplicacion(tmp_path / 'sistema', HOY, puerto=0)
    servidor = aplicacion.crear_servidor(app, '127.0.0.1')
    url = iniciar(servidor)
    hoja_rara = ['HOJA DE REFERENCIA DEL CENTRO DE SALUD', 'Se atendió al menor en la unidad', 'Peso 24 kg Talla 1.20 m']
    try:
        assert app.configurar(url_cerebro, 'clave-de-prueba') is None
        leido = escanear(url, app, pdf_de_hojas([ACTA.splitlines(), hoja_rara, ['REVERSO SIN TEXTO FIJO PERO CON BASTANTES LETRAS ESCRITAS']]))[1]
        acta, rara, reverso = leido['paginas']
        assert acta['reconocido'] and 'ia' not in acta
        assert (rara['tipo'], rara['reconocido'], rara['ia']['titulo']) == ('certificado_medico', False, 'Constancia de salud')
        assert reverso['tipo'] == 'certificado_medico'  # Una hoja sin señas sigue a la anterior (su reverso), ya con lo que dijo la IA.
        assert sum(p.startswith('Esta hoja viene del expediente') for p in modelo_que_ve) == 2  # Solo las que no se reconocieron.
        # Si la IA no ve ninguna plantilla, lo dice con el nombre del documento y queda como "otro".
        CLASIFICADO.update(tipo='otro', titulo='Constancia de estudios', motivo='el título impreso')
        leido = escanear(url, app, pdf_de_hojas([hoja_rara]))[1]
        assert (leido['paginas'][0]['tipo'], leido['paginas'][0]['ia']['titulo']) == ('otro', 'Constancia de estudios')
    finally:
        CLASIFICADO.update(tipo='certificado_medico', titulo='Constancia de salud', motivo='trae el sello de un centro de salud')
        servidor.shutdown()
        cerebro.shutdown()


def test_las_boletas_se_separan_por_hoja_y_cada_una_lleva_su_grado(tmp_path):
    (tmp_path / 'boleta_1.pdf').write_bytes(pdf_de_hojas([['BOLETA PRIMER GRADO'], ['BOLETA SEGUNDO GRADO'], ['BOLETA TERCER GRADO']]))
    documentos = taller.separar_hojas(tmp_path, [{'archivo': 'boleta_1.pdf', 'tipo': 'boleta', 'nombre': 'Boleta', 'campos': []},
                                                 {'archivo': 'acta.pdf', 'tipo': 'acta_alumno', 'nombre': 'Acta', 'campos': []}])
    assert [(d['archivo'], d.get('hoja')) for d in documentos] == [('boleta_1_hoja1.pdf', 1), ('boleta_1_hoja2.pdf', 2),
                                                                    ('boleta_1_hoja3.pdf', 3), ('acta.pdf', None)]
    # El grado: el que dice la boleta ("SEGUNDO", "3°") o, si no se leyó, el que sale de su ciclo contra el que cursa ahora.
    assert taller.grado_de('SEGUNDO GRADO', '') == 2 and taller.grado_de('3°', None) == 3
    assert taller.grado_de('', '2023-2024', (4, '2026-2027')) == 1 and taller.grado_de('', '', (4, '2026-2027')) is None
    leido = [{'clave': 'grado_cursado', 'campo': 'grado_cursado', 'valor': 'Segundo'},
             {'clave': 'promedio', 'campo': 'promedio', 'categoria': 'escolar', 'valor': '8.5', 'confianza': 'alta'},
             {'clave': 'materias', 'campo': 'materia', 'categoria': 'escolar', 'valor': 'Español', 'confianza': 'alta'}]
    salida, grado = taller.por_grado(leido, (4, '2026-2027'))
    assert grado == 2 and [(r['campo'], r['varios'], r['detalle']['grado']) for r in salida] == [('promedio_2', False, 2), ('materia_2', True, 2)]
    sin_grado, _ = taller.por_grado(leido[1:])
    assert [(r['campo'], r['confianza']) for r in sin_grado] == [('promedio', 'baja'), ('materia', 'baja')]  # Por confirmar.
