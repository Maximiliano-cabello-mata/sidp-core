"""Los formatos de la supervisión: Metadato, base del 911 con fórmulas, condición y discapacidad y plantillas. Todo es inventado."""
import datetime
import io
import shutil
import subprocess

import openpyxl
import pytest

from pruebas.test_documentos import ESCUELA, HOY, dato, registro, registros
from validacion import documentos, formatos_oficiales

MADRE = 'RASE880521MMCMTL06'


def con_tutor():
    lista = registros()
    lista.append(registro('XEXD180101HMCXXXA4', 'JOSE', 'DIAZ', 1, 'A', datetime.date(2018, 1, 1), 'H',
                          {'tutor': [dato('nombre', 'MARIA ELENA RAMOS SOTO'), dato('parentesco', 'madre'), dato('curp', MADRE)],
                           'contacto': [dato('telefono', '5512345678'), dato('correo', 'familia@ejemplo.mx')],
                           'domicilio': [dato('domicilio', 'CALLE PINO 5 COL. CENTRO CP 56600 MUNICIPIO DE CHALCO')],
                           'salud': [dato('padecimiento', 'Asma'), dato('discapacidad', 'Baja visión'), dato('peso', '24.5'),
                                     dato('talla', '1.21'), dato('tipo_sangre', 'A+'), dato('usa_lentes', 'sí'), dato('vacunas', 'Completo'),
                                     dato('servicio_medico', 'IMSS Bienestar'), dato('fecha_certificado', '2026-08-30')]}))
    return lista


def libro(clave, lista=None):
    nombre, contenido, tipo = documentos.generar(clave, 'xlsx', lista or con_tutor(), ESCUELA, '2026-2027', HOY)
    assert nombre.endswith('.xlsx') and tipo == documentos.MIME['xlsx']
    return openpyxl.load_workbook(io.BytesIO(contenido)), contenido


def recalcular(contenido, tmp_path):
    """Los valores que da Excel al abrir: LibreOffice calcula las fórmulas (si no está instalado, la prueba se salta)."""
    soffice = shutil.which('soffice')
    if not soffice:
        pytest.skip('Sin LibreOffice para calcular las fórmulas')
    (tmp_path / 'entrada.xlsx').write_bytes(contenido)
    subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(tmp_path / 'calculado'),
                    str(tmp_path / 'entrada.xlsx')], check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(tmp_path / 'calculado' / 'entrada.xlsx', data_only=True)


def renglon(hoja, numero):
    return [c.value for c in hoja[numero]]


def test_el_metadato_trae_las_71_columnas_y_llena_lo_que_tiene_sidp():
    hoja = libro('metadato')[0]['METADATO']
    encabezados = renglon(hoja, 6)
    assert len(encabezados) == 71 and encabezados[:3] == ['NP', 'GRADO Y GRUPO', 'APELLIDO PATERNO'] and encabezados[-1] == 'OBSERVACIONES DEL TUTOR'
    columna = {texto: n for n, texto in enumerate(formatos_oficiales.METADATO_ALUMNO)}
    tutor = {texto: n + 40 for n, texto in enumerate(formatos_oficiales.METADATO_TUTOR)}
    jose = next(renglon(hoja, n) for n in range(7, 11) if renglon(hoja, n)[columna['NOMBRE (S)']] == 'JOSE')
    assert jose[columna['GÉNERO']] == 'HOMBRE' and jose[columna['LUGAR DE NACIMIENTO']] == 'ESTADO DE MÉXICO'
    assert jose[columna['EDAD']].startswith('=IF(H') and 'DATEDIF' in jose[columna['EDAD']]
    assert (jose[columna['CALLE']], jose[columna['NO. EXTERIOR']], jose[columna['C. P.']]) == ('CALLE PINO', '5', '56600')
    assert jose[columna['ENFERMEDAD CRÓNICA O PADECIMIENTO']] == 'ASMA' and jose[columna['PRESENTA ALGUNA DISCAPACIDAD']] == 'BAJA VISIÓN'
    # La salud general del certificado médico llena sus columnas.
    assert [jose[columna[c]] for c in ('TIPO DE SANGRE', 'PESO', 'TALLA', 'PIE PLANO, USA LENTES, APARATO AUDITIVO O PRÓTESIS',
                                       'CUENTA CON SERVICIO MÉDICO', 'INSTITUCIÓN', 'ENTREGÓ CERTIFICADO MÉDICO',
                                       'CUENTA CON ESQUEMA DE VACUNACIÓN COMPLETO')] == \
        ['A+', 24.5, 1.21, 'USA LENTES', 'SÍ', 'IMSS-BIENESTAR', 'SÍ', 'SÍ']
    pedro = next(renglon(hoja, n) for n in range(7, 11) if renglon(hoja, n)[columna['NOMBRE (S)']] == 'PEDRO')
    assert pedro[columna['TIPO DE SANGRE']] is None and pedro[columna['ENTREGÓ CERTIFICADO MÉDICO']] is None  # Sin certificado leído.
    # Las letras de la CURP de la madre dicen cuál es su apellido y cuál su nombre.
    assert [jose[tutor[c]] for c in ('APELLIDO PATERNO', 'APELLIDO MATERNO', 'NOMBRE (S)', 'PARENTESCO')] == ['RAMOS', 'SOTO', 'MARIA ELENA', 'MADRE']
    assert jose[tutor['GÉNERO']] == 'MUJER' and jose[tutor['FECHA DE NACIMIENTO']].date() == datetime.date(1988, 5, 21)
    pedro = next(renglon(hoja, n) for n in range(7, 11) if renglon(hoja, n)[columna['NOMBRE (S)']] == 'PEDRO')
    assert pedro[columna['SITUACIÓN DEL ALUMNO']] == 'REPETIDOR' and pedro[columna['HABLA ALGUNA LENGUA INDÍGENA']] == 'SÍ (MAZAHUA)'


def test_el_nombre_del_tutor_se_parte_en_apellidos_y_nombres():
    assert formatos_oficiales.partir_nombre('MARIA ELENA RAMOS SOTO', MADRE) == ('RAMOS', 'SOTO', 'MARIA ELENA')
    assert formatos_oficiales.partir_nombre('Juan Carlos de la Cruz Pérez') == ('DE LA CRUZ', 'PEREZ', 'JUAN CARLOS')
    assert formatos_oficiales.partir_nombre('ROSA DIAZ') == ('DIAZ', '', 'ROSA') and formatos_oficiales.partir_nombre('') == ('', '', '')


def test_la_base_del_911_se_calcula_sola_con_la_curp(tmp_path):
    hojas, contenido = libro('base_911')
    assert hojas.sheetnames == ['911', 'BASE']
    base = hojas['BASE']
    assert renglon(base, 5)[:12] == ['NP', 'CURP', 'PRIMER APELLIDO', 'SEGUNDO APELLIDO', 'NOMBRE(S)', 'GRADO', 'GRUPO', 'ESTATUS', 'SEXO',
                                     'FECHA DE NACIMIENTO', 'EDAD', 'RANGO 911']
    assert base['I6'].value.startswith('=IF(LEN(B6)<18') and 'MID(B6,11,1)' in base['I6'].value
    # Sobran renglones con fórmulas para los que lleguen después.
    assert base[f'I{6 + 4 + formatos_oficiales.FILAS_LIBRES - 1}'].value.startswith('=IF(LEN(')
    calculado = recalcular(contenido, tmp_path)
    filas = {(f[0], f[1], f[2]): f[3:] for f in calculado['911'].iter_rows(min_row=6, values_only=True) if f[0]}
    assert filas[('ESCUELA', 'TOTAL', None)][-1] == 4
    assert filas[('1°', 'HOMBRES', 'NUEVO INGRESO')][-1] == 2 and filas[('3°', 'HOMBRES', 'REPETIDORES')][-1] == 1
    edades = list(formatos_oficiales.indicadores.EDADES_911)
    assert filas[('1°', 'MUJERES', 'NUEVO INGRESO')][edades.index('7')] == 1  # Nació en marzo de 2019: 7 años al 1 de septiembre.
    assert [c.value for c in calculado['BASE'][6]][8:11][0] in ('H', 'M')


def test_el_censo_marca_con_x_cada_condicion_y_cuenta_el_total():
    hoja = libro('condicion')[0]['PRIMERO']
    titulos = renglon(hoja, 6)
    filas = {f[1]: f for f in hoja.iter_rows(min_row=7, values_only=True) if f[1]}
    assert filas['LOPEZ GIL ANA'][titulos.index('TDA')] == 'X' and filas['DIAZ GIL JOSE'][titulos.index('BAJA VISIÓN')] == 'X'
    assert filas['TOTAL'][titulos.index('TDA')].startswith('=COUNTIF(')


def test_las_plantillas_del_maestro_traen_a_los_alumnos_y_las_formulas(tmp_path):
    hojas, contenido = libro('calificaciones')
    assert hojas.sheetnames == ['CONCENTRADO', '1° A', '3° B']
    grupo = hojas['1° A']
    assert [grupo.cell(n, 4).value for n in (6, 7, 8)] == ['LUIS', 'JOSE', 'ANA']
    for fila, notas in zip((6, 7, 8), ([9, 8, 9, 10], [7, 7, 6, 7], [5, 6, 5, 6])):
        for columna, nota in zip('EFGH', notas):
            grupo[f'{columna}{fila}'] = nota
    salida = io.BytesIO()
    hojas.save(salida)
    calculado = recalcular(salida.getvalue(), tmp_path)
    assert [calculado['1° A'].cell(n, 11).value for n in (6, 7, 8)] == [9, 6.8, 5.5]
    concentrado = {f[0]: f[1:] for f in calculado['CONCENTRADO'].iter_rows(min_row=6, values_only=True) if f[0]}
    assert concentrado['1° A'][:5] == (3, 1, 0, 1, 1) and concentrado['1° A'][5] == 7.1
    asistencias = libro('asistencias')[0]['1° A']
    assert asistencias.cell(5, 5).value == 'SEPTIEMBRE' and asistencias.cell(7, 8).value.startswith('=IF(COUNT(E7:G7)')
    boleta = libro('boleta')[0]['BASE BOLETA']
    assert len(renglon(boleta, 5)) == 31 and renglon(boleta, 6)[3] == 'XEXB190710HDFXXXA2'


def test_un_grupo_sale_solo_con_sus_alumnos():
    nombre, contenido, _ = documentos.generar('calificaciones', 'xlsx', con_tutor(), ESCUELA, '2026-2027', HOY, grado=3, letra='B')
    assert openpyxl.load_workbook(io.BytesIO(contenido)).sheetnames == ['CONCENTRADO', '3° B'] and '3B' in nombre
