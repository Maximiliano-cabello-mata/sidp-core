"""El motor de documentos (911, constancias, listas) y los datos de la escuela. Todo es inventado."""
import datetime
import io
import json
import sqlite3
import urllib.request
import zipfile
from pathlib import Path

import openpyxl
import pytest

from pruebas.test_entorno_pruebas import entorno, pedir  # noqa: F401 (entorno es la prueba con el panel)
from validacion import documentos, escuela, estandar

HOY = datetime.date(2026, 10, 23)
DISENO = Path(__file__).resolve().parent.parent / 'database' / 'diseno_integral.sql'
ESCUELA = {'nombre': 'Escuela Primaria de Prueba', 'cct': '15EPR0000X', 'turno': 'matutino', 'municipio': 'Chalco',
           'director': 'PROFR. INVENTADO'}


def dato(campo, valor):
    return {'campo': campo, 'valor': valor, 'estado': 'vigente', 'detalle': {}}


def registro(curp, nombres, primer, grado, letra, nacimiento, sexo, datos=None):
    return estandar.registro({'curp': curp, 'nombres': nombres, 'primer_apellido': primer, 'segundo_apellido': 'GIL', 'grado': grado,
                              'letra': letra, 'fecha_nacimiento': nacimiento, 'sexo': sexo}, datos or {}, '2026-2027', HOY)


def registros():
    return [
        registro('XEXA190305MMCXXXA1', 'ANA', 'LOPEZ', 1, 'A', datetime.date(2019, 3, 5), 'M',
                 {'escolar': [dato('repite_grado', 'no'), dato('beca', 'federal'), dato('procedencia', 'misma')],
                  'salud': [dato('alergia', 'Polvo'), dato('discapacidad', 'TDAH')], 'alumno': [dato('lengua_indigena', 'no')]}),
        registro('XEXB190710HDFXXXA2', 'LUIS', 'CRUZ', 1, 'A', datetime.date(2019, 7, 10), 'H', {'escolar': [dato('repite_grado', 'no')]}),
        registro('XEXC170202HMCXXXA3', 'PEDRO', 'RUIZ', 3, 'B', datetime.date(2017, 2, 2), 'H',
                 {'escolar': [dato('repite_grado', 'sí')], 'alumno': [dato('lengua_indigena', 'sí'), dato('lengua', 'Mazahua')]}),
    ]


def test_el_911_cuenta_por_grado_sexo_y_edad_y_dice_si_cuadra():
    r = registros()
    c = documentos.cuadros_911(r)
    total = c['edad'][-1]
    assert total[:2] == ['Total', 'Total'] and total[-1] == 3 and sum(total[2:-1]) == 3
    primero = {f[1]: f[-1] for f in c['edad'] if f[0] == '1°'}
    assert primero == {'Hombres': 1, 'Mujeres': 1, 'Subtotal': 2}
    assert c['repetidores'][2] == ['3°', 1, 0, 1, 1, 0, 1, 0, 0] and c['repetidores'][-1][6] == 1  # Repite 1 de 3.
    assert c['discapacidad'] == [['TDAH', 0, 1] + [0] * 10 + [1]]
    assert c['lengua'][2][:4] == ['3°', 1, 0, 1] and c['lenguas'] == [('Mazahua', 1)]
    assert c['nacidos'][0][:4] == ['1°', 1, 0, 1]  # Nació en la Ciudad de México (DF en la CURP).
    assert ['Federal (por ejemplo, Rita Cetina)', 0, 1, 1] in c['becas']
    assert [g[1] for g in c['grupos']] == ['1° A', '3° B']
    revision = {f[0]: f[1] for f in documentos.revisar_911(r, ESCUELA, c)}
    assert revision['La tabla de edades suma lo mismo que la matrícula'] == 'Sí'
    assert revision['Todos tienen el dato de beca'] == 'No' and revision['Datos de la escuela completos'] == 'Sí'
    nombre, contenido, tipo = documentos.generar('911', 'xlsx', r, ESCUELA, '2026-2027', HOY)
    assert nombre == '911_inicio_2026-2027.xlsx' and tipo == documentos.MIME['xlsx']
    libro = openpyxl.load_workbook(io.BytesIO(contenido))
    assert libro.sheetnames == ['Escuela', 'Alumnos por edad', 'Repetidores', 'Grupos', 'Discapacidad', 'Lengua y afro',
                                'Lugar de nacimiento', 'Becas', 'Procedencia', 'Revisión', 'Faltan']
    faltan = [[c.value for c in fila] for fila in libro['Faltan'].iter_rows(min_row=5)]
    assert [f[1] for f in faltan] == ['CRUZ GIL LUIS', 'LOPEZ GIL ANA', 'RUIZ GIL PEDRO'] and 'beca' in faltan[0][3]
    assert faltan[1][3] == 'el dato de afromexicano'  # A Ana solo le falta eso.
    assert not any('peso' in str(c.value).lower() for hoja in libro for fila in hoja.iter_rows() for c in fila)  # Solo lo necesario.


def test_el_estado_de_salud_del_grupo_dice_como_esta_cada_alumno():
    r = registros()
    r[1] = registro('XEXB190710HDFXXXA2', 'LUIS', 'CRUZ', 1, 'A', datetime.date(2019, 7, 10), 'H',
                    {'salud': [dato('peso', '30 kg'), dato('talla', '118 cm'), dato('tipo_sangre', 'O+'), dato('agudeza_visual', '20/20'),
                               dato('salud_bucal', 'Caries'), dato('vacunas', 'Completo'), dato('educacion_fisica', 'Apto'),
                               dato('estado_general', 'Clínicamente sano'), dato('fecha_certificado', '2026-09-01')]})
    nombre, contenido, _ = documentos.generar('salud', 'xlsx', r, ESCUELA, '2026-2027', HOY, grado=1, letra='A')
    assert nombre == 'salud_1A_2026-10-23.xlsx'
    hoja = openpyxl.load_workbook(io.BytesIO(contenido))['1° A']
    assert hoja['A1'].value == 'Estado de salud · Grupo 1° A' and '2 alumnos: 1 con certificado médico leído y 2 con algo' in hoja['A2'].value
    encabezado = [c.value for c in hoja[4]]
    assert encabezado[:6] == ['No.', 'Nombre', 'Edad', 'Peso (kg)', 'Estatura (m)', 'IMC'] and encabezado[-1] == 'En general'
    luis = dict(zip(encabezado, [c.value for c in hoja[5]]))
    assert (luis['Nombre'], luis['Edad'], luis['Peso (kg)'], luis['Estatura (m)'], luis['IMC']) == ('CRUZ GIL LUIS', 7, 30, 1.18, 21.5)
    assert (luis['Estado nutricional'], luis['Sangre'], luis['Vista'], luis['Dientes'], luis['Vacunas']) == \
        ('Obesidad', 'O+', 'Normal', 'Necesita dentista', 'Completas')
    assert luis['En general'] == 'Sano según el médico. Requiere atención: obesidad para su edad, atención dental.'
    ana = dict(zip(encabezado, [c.value for c in hoja[6]]))
    assert ana['En general'] == 'Requiere atención: alergia a polvo, TDAH.' and ana['Peso (kg)'] is None
    nombre, contenido, _ = documentos.generar('salud', 'pdf', r, ESCUELA, '2026-2027', HOY)
    assert nombre == 'salud_escuela_2026-10-23.pdf' and contenido.startswith(b'%PDF')


def test_constancias_y_listas_salen_en_word_pdf_y_excel():
    r = registros()
    nombre, contenido, _ = documentos.generar('constancia', 'docx', r, ESCUELA, '2026-2027', HOY, curp='XEXA190305MMCXXXA1')
    assert nombre == 'constancia_LOPEZ_ANA.docx'
    with zipfile.ZipFile(io.BytesIO(contenido)) as archivo:
        texto = archivo.read('word/document.xml').decode()
        assert {'[Content_Types].xml', '_rels/.rels', 'word/document.xml'} <= set(archivo.namelist())
    assert 'la alumna' in texto and 'inscrita' in texto and 'primer grado, grupo A' in texto and 'XEXA190305MMCXXXA1' in texto
    assert '23 de octubre de 2026' in texto and 'PROFR. INVENTADO' in texto and 'director de la Escuela Primaria de Prueba' in texto
    # Todo el grupo: una hoja por alumno. Sin datos de la escuela, quedan espacios para llenar a mano.
    _, contenido, _ = documentos.generar('constancia', 'docx', r, {}, '2026-2027', HOY, grado=1, letra='A')
    texto = zipfile.ZipFile(io.BytesIO(contenido)).read('word/document.xml').decode()
    assert texto.count('w:type="page"') == 1 and texto.count('CONSTANCIA DE INSCRIPCIÓN') == 2 and documentos.EN_BLANCO in texto
    pypdf = pytest.importorskip('pypdf')
    pytest.importorskip('reportlab')
    _, contenido, tipo = documentos.generar('lista', 'pdf', r, ESCUELA, '2026-2027', HOY)
    paginas = pypdf.PdfReader(io.BytesIO(contenido)).pages
    assert tipo == 'application/pdf' and len(paginas) == 2  # Una por grupo.
    assert 'Lista del grupo 1° A' in paginas[0].extract_text() and 'XEXB190710HDFXXXA2' in paginas[0].extract_text()
    _, contenido, _ = documentos.generar('emergencia', 'xlsx', r, ESCUELA, '2026-2027', HOY, grado=1, letra='A')
    hoja = openpyxl.load_workbook(io.BytesIO(contenido))['1° A']
    assert [c.value for c in hoja[4]][2] == 'Alergias'
    assert [[c.value for c in fila][1:3] for fila in hoja.iter_rows(min_row=5)] == [['CRUZ GIL LUIS', None], ['LOPEZ GIL ANA', 'POLVO']]
    with pytest.raises(documentos.ErrorDocumento):
        documentos.generar('911', 'pdf', r, ESCUELA, '2026-2027', HOY)  # El 911 solo sale en Excel.
    with pytest.raises(documentos.ErrorDocumento):
        documentos.generar('lista', 'pdf', r, ESCUELA, '2026-2027', HOY, grado=6, letra='A')  # Grupo sin alumnos.


def test_los_datos_de_la_escuela_se_revisan_y_la_tabla_vieja_se_completa():
    db = sqlite3.connect(':memory:', isolation_level=None)
    db.execute("CREATE TABLE escuela (id INTEGER PRIMARY KEY CHECK (id = 1), cct TEXT NOT NULL, nombre TEXT NOT NULL, "
               "turno TEXT NOT NULL, zona_escolar TEXT, domicilio TEXT, version INTEGER NOT NULL DEFAULT 1) STRICT")  # La de antes.
    datos = escuela.Escuela(db)
    assert datos.datos()['completa'] is False and 'la clave CCT' in datos.datos()['faltan']
    with pytest.raises(escuela.ErrorEscuela):
        datos.guardar({**ESCUELA, 'cct': '15EPR000'})
    guardado = datos.guardar({**ESCUELA, 'cct': '15 epr-0000x', 'turno': 'Matutino'})
    assert (guardado['cct'], guardado['turno'], guardado['turno_texto'], guardado['completa']) == ('15EPR0000X', 'matutino', 'Matutino', True)
    assert datos.guardar({**ESCUELA, 'subdirector': 'PROFR. OTRO'})['subdirector'] == 'PROFR. OTRO'
    assert db.execute('SELECT version FROM escuela').fetchone() == (2,)
    nueva = sqlite3.connect(':memory:', isolation_level=None)
    nueva.executescript(DISENO.read_text(encoding='utf-8'))
    assert escuela.Escuela(nueva).guardar(ESCUELA)['director'] == 'PROFR. INVENTADO'  # El diseño ya trae las columnas.


def test_los_catalogos_del_gobierno_no_aceptan_otra_cosa():
    assert estandar.valor_de('escolar', 'repite_grado', 'Si') == 'sí' and estandar.valor_de('escolar', 'beca', 'Federal') == 'federal'
    assert estandar.valor_de('escolar', 'beca', 'De la propia escuela') == 'escuela'
    assert estandar.valor_de('alumno', 'lengua', ' Mazahua ') == 'Mazahua'  # Sin catálogo: el texto limpio.
    with pytest.raises(ValueError):
        estandar.valor_de('escolar', 'procedencia', 'de lejos')


def test_el_panel_arma_documentos_con_las_listas_de_aqui_si_no_esta_la_pc_4(entorno):  # noqa: F811
    app, url = entorno
    token = {'X-SIDP-Token': app.token}
    app.bandeja.cerebro = None  # La PC 4 apagada.
    assert pedir(url, '/api/listas/guardar', {'grado': 2, 'letra': 'A', 'filas': [
        {'nombres': 'LUIS', 'primer_apellido': 'HERNANDEZ', 'segundo_apellido': 'RAMOS', 'curp': 'HERL180314HMCRMSA9'}]}, token)[0] == 200
    codigo, r = pedir(url, '/api/escuela/guardar', {**ESCUELA, 'cct': 'mal'}, token)
    assert codigo == 400 and 'CCT' in r['error']
    assert pedir(url, '/api/escuela/guardar', ESCUELA, token)[1]['completa'] is True
    codigo, lista = pedir(url, '/api/documentos', cabeceras=token)
    assert codigo == 200 and lista['escuela']['completa'] and 'PC 4' in lista['aviso']
    assert {d['clave'] for d in lista['documentos']} == {'911', 'constancia', 'lista', 'emergencia', 'salud', 'vive_saludable', 'metadato', 'base_911',
                                                       'condicion', 'calificaciones', 'boleta', 'asistencias'}
    grupo = pedir(url, '/api/grupos', cabeceras=token)[1]['grupos'][0]['id']
    alumno = pedir(url, f'/api/grupo?id={grupo}', cabeceras=token)[1]['alumnos'][0]['id']
    solicitud = urllib.request.Request(url + '/api/documentos/generar', json.dumps({'clave': 'constancia', 'formato': 'docx', 'grupo_id': grupo,
                                       'alumno_id': alumno}).encode(), {'Content-Type': 'application/json', **token})
    with urllib.request.urlopen(solicitud, timeout=30) as respuesta:
        assert 'constancia_HERNANDEZ_LUIS.docx' in respuesta.headers['Content-Disposition']
        texto = zipfile.ZipFile(io.BytesIO(respuesta.read())).read('word/document.xml').decode()
    assert 'segundo grado, grupo A' in texto and 'el alumno' in texto and '15EPR0000X' in texto
    codigo, r = pedir(url, '/api/documentos/generar', {'clave': 'nada', 'formato': 'pdf'}, token)
    assert codigo == 400 and 'no existe' in r['error']
